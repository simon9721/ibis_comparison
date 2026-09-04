from __future__ import annotations

import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
LOCAL_DEPS = ROOT / ".codex_deps" / "presentation" / "python"
if LOCAL_DEPS.exists():
    sys.path.insert(0, str(LOCAL_DEPS))
sys.path.insert(0, str(ROOT))

from pptx.enum.text import PP_ALIGN

from tools.presentation_kit import EquationRenderer, GreenDeck


OUTPUT = ROOT / "results" / "presentation_toolkit_demo" / "green_template_toolkit_demo.pptx"


def build() -> Path:
    renderer = EquationRenderer(backend="auto")
    deck = GreenDeck(equation_renderer=renderer)
    deck.set_title_slide(
        "Reusable Scripted Presentation Kit",
        "Green lab template + cached equations\nDemonstration deck",
        notes=(
            "This deck is generated through the reusable presentation kit. The title slide and "
            "content layouts come from the checked-in green template."
        ),
    )

    slide = deck.add_slide(
        "1. One template, many projects",
        section="Workflow",
        notes=(
            "The reusable package separates visual conventions from project content. Future deck "
            "scripts import GreenDeck instead of importing an IBIS-specific generator."
        ),
    )
    deck.add_box(slide, "Project data", 0.7, 1.65, 2.1, 0.8, bold=True)
    deck.add_arrow(slide, 2.8, 2.05, 3.75, 2.05)
    deck.add_box(slide, "Reusable slide API", 3.75, 1.65, 2.4, 0.8, bold=True)
    deck.add_arrow(slide, 6.15, 2.05, 7.1, 2.05)
    deck.add_box(slide, "Green template", 7.1, 1.65, 2.1, 0.8, bold=True)
    deck.add_arrow(slide, 9.2, 2.05, 10.15, 2.05)
    deck.add_box(slide, "PPTX + notes", 10.15, 1.65, 2.1, 0.8, bold=True)
    deck.add_bullets(
        slide,
        [
            "Template masters and layouts stay authoritative",
            "Common slide components share spacing, fonts, and colors",
            "Project scripts contain only the technical story and evidence",
            "Speaker notes are generated together with the visible slide",
        ],
        1.0,
        3.25,
        10.8,
        2.3,
        size=18,
    )
    deck.add_takeaway(slide, "New projects reuse the visual system without copying an old project's implementation.")
    deck.add_source(slide, "Template: assets/presentation_templates/mst_green_16x9.pptx")

    slide = deck.add_slide(
        "2. Equations are rendered once and reused",
        section="Math",
        notes=(
            f"The active backend for this build is {renderer.selected_backend()}. Expressions are "
            "rendered to transparent high-resolution assets and cached by content hash."
        ),
    )
    deck.add_text(slide, "First-order state equation", 0.8, 1.15, 4.3, 0.35, size=19, bold=True, color=deck.theme.green)
    deck.add_equation(slide, r"\frac{dG}{dt}=\frac{G_{\mathrm{target}}-G}{\tau}", 0.85, 1.65, 5.2, 1.0, font_size_pt=34)
    deck.add_text(slide, "Closed-form response", 6.7, 1.15, 4.3, 0.35, size=19, bold=True, color=deck.theme.green)
    deck.add_equation(slide, r"G(t)=G_{\infty}+(G_0-G_{\infty})e^{-t/\tau}", 6.55, 1.65, 5.5, 1.0, font_size_pt=31)
    deck.add_text(slide, "Current scaling", 0.8, 3.2, 4.3, 0.35, size=19, bold=True, color=deck.theme.green)
    deck.add_equation(slide, r"I_{\mathrm{PU}}=K_u\,I_{\mathrm{PU,full}}(V)", 0.85, 3.72, 5.2, 0.9, font_size_pt=32)
    deck.add_code_box(
        slide,
        "deck.add_equation(\n"
        "    slide,\n"
        "    r'\\frac{dG}{dt}=\\frac{G_{target}-G}{\\tau}',\n"
        "    x=0.85, y=1.65, w=5.2, h=1.0,\n"
        ")",
        6.55,
        3.35,
        5.55,
        1.65,
        size=11.5,
    )
    deck.add_box(
        slide,
        f"Renderer used in this build: {renderer.selected_backend()}",
        3.65,
        5.45,
        5.6,
        0.62,
        fill=deck.theme.pale_green,
        line=deck.theme.green,
        bold=True,
    )
    deck.add_takeaway(slide, "PowerPoint receives a finished equation image; it does no interactive equation typesetting.")
    deck.add_source(slide, "Equation cache: .codex_deps/presentation/equation_cache")

    slide = deck.add_slide(
        "3. The API covers the recurring technical-slide patterns",
        section="Components",
        notes=(
            "This slide exercises the common primitives. More specialized scientific figures remain "
            "ordinary PNG assets produced by the project's analysis scripts."
        ),
    )
    deck.add_box(slide, "Process box", 0.7, 1.35, 2.25, 0.72, bold=True)
    deck.add_arrow(slide, 2.95, 1.71, 3.75, 1.71)
    deck.add_box(slide, "Next process", 3.75, 1.35, 2.25, 0.72, bold=True)
    deck.add_bullets(
        slide,
        ["Bullets", "Presenter notes", "Source footers", "Contained figures"],
        0.8,
        2.65,
        4.2,
        2.3,
        size=17,
    )
    deck.add_code_box(slide, "result = fit(data)\nassert result.pass_gate\nplot(result)", 5.3, 2.65, 3.2, 1.35, size=13)
    deck.add_box(
        slide,
        "Status\nPASS",
        9.2,
        2.65,
        2.1,
        1.35,
        fill=deck.theme.light_green,
        line=deck.theme.green,
        size=18,
        bold=True,
    )
    deck.add_text(
        slide,
        "The scientific plots remain generated by Matplotlib; the toolkit standardizes how they are framed and explained.",
        5.3,
        4.65,
        6.2,
        0.75,
        size=17,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    deck.add_takeaway(slide, "Use the toolkit for narrative consistency and project scripts for domain-specific evidence.")
    deck.add_source(slide, "API: tools/presentation_kit")

    return deck.save(OUTPUT)


if __name__ == "__main__":
    path = build()
    print(path)
