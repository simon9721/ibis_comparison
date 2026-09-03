from __future__ import annotations

import csv
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
LOCAL_DEPS = ROOT / ".codex_deps" / "presentation" / "python"
if LOCAL_DEPS.exists():
    sys.path.insert(0, str(LOCAL_DEPS))
sys.path.insert(0, str(ROOT))

import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from tools.presentation_kit import EquationRenderer, GreenDeck
from eye_diagram import parse_ngspice_raw


OUT = ROOT / "results" / "ibis_kukd_handwritten_notes_deck"
ASSETS = OUT / "generated_assets"
OUTPUT = OUT / "IBIS_KuKd_continuous_gate_state_from_notes_revised.pptx"
SOURCE_PAGES = OUT / "source_pages"
REAL_EXAMPLE_RAW = (
    ROOT
    / "results"
    / "io_buf_two_state_gate_model_2026-06-30"
    / "cases"
    / "short_pulse_1ns_high"
    / "ngspice_two_state_directional_residual"
    / "short_pulse_1ns_high_ngspice_two_state_directional_residual.raw"
)


def _save_figure(fig: plt.Figure, name: str) -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    path = ASSETS / name
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def make_problem_figure() -> Path:
    t = np.linspace(0.0, 3.0, 1000)
    reverse = 1.15
    tau_on = 0.75
    tau_off = 0.42
    rising = 1.0 - np.exp(-t / tau_on)
    legacy = rising.copy()
    mask = t >= reverse
    legacy[mask] = np.exp(-(t[mask] - reverse) / tau_off)
    state = rising.copy()
    g_rev = 1.0 - np.exp(-reverse / tau_on)
    state[mask] = g_rev * np.exp(-(t[mask] - reverse) / tau_off)

    fig, ax = plt.subplots(figsize=(8.6, 3.2))
    ax.plot(t, rising, color="#777777", lw=2, alpha=0.65, label="unfinished rising path")
    ax.plot(t, legacy, color="#ca3030", lw=3, label="table replay restart")
    ax.plot(t, state, color="#2b7a43", lw=3, label="continuous stored state")
    ax.axvline(reverse, color="#202020", ls="--", lw=1.5)
    ax.annotate(
        "reverse edge",
        xy=(reverse, 0.05),
        xytext=(reverse + 0.12, 0.18),
        fontsize=10,
        arrowprops={"arrowstyle": "->", "color": "#202020"},
    )
    ax.annotate(
        "unphysical jump",
        xy=(reverse + 0.015, np.exp(-0.015 / tau_off)),
        xytext=(1.65, 1.04),
        color="#ca3030",
        fontsize=10,
        arrowprops={"arrowstyle": "->", "color": "#ca3030"},
    )
    ax.set(xlabel="time after first edge", ylabel="normalized state", ylim=(-0.04, 1.12))
    ax.grid(alpha=0.22)
    ax.legend(loc="lower right", frameon=False, ncol=1)
    fig.tight_layout()
    return _save_figure(fig, "01_replay_discontinuity.png")


def make_tau_figure() -> Path:
    t = np.linspace(0.0, 4.0, 700)
    rise = 1.0 - np.exp(-t)
    fall = np.exp(-t)
    fig, ax = plt.subplots(figsize=(8.5, 3.1))
    ax.plot(t, rise, color="#2b7a43", lw=3, label=r"rising: $1-e^{-t/\tau}$")
    ax.plot(t, fall, color="#2467ad", lw=3, label=r"falling: $e^{-t/\tau}$")
    ax.axvline(1.0, color="#202020", ls="--", lw=1.4)
    ax.scatter([1.0, 1.0], [1.0 - np.exp(-1.0), np.exp(-1.0)], color=["#2b7a43", "#2467ad"], zorder=5)
    ax.annotate("63.2%", (1.0, 1.0 - np.exp(-1.0)), xytext=(1.18, 0.73), fontsize=11)
    ax.annotate("36.8%", (1.0, np.exp(-1.0)), xytext=(1.18, 0.24), fontsize=11)
    ax.set(xlabel=r"normalized time $t/\tau$", ylabel="normalized gate state", ylim=(-0.03, 1.03))
    ax.grid(alpha=0.22)
    ax.legend(frameon=False, loc="center right")
    fig.tight_layout()
    return _save_figure(fig, "02_tau_63_percent.png")


def make_mapping_figure() -> Path:
    t = np.linspace(0.0, 4.0, 500)
    g = 1.0 - np.exp(-t / 1.15)
    k = np.clip(g**1.65 - 0.035 * np.exp(-((g - 0.72) / 0.11) ** 2), 0.0, 1.0)
    fig, axes = plt.subplots(1, 3, figsize=(10.5, 3.2))
    axes[0].plot(t, k, color="#ca3030", lw=3)
    axes[0].set(title=r"original $K(t)$ table", xlabel="time", ylabel="coefficient")
    axes[1].plot(t, g, color="#2b7a43", lw=3)
    axes[1].set(title=r"reconstructed gate state $G(t)$", xlabel="time", ylabel="state")
    axes[2].plot(g, k, color="#7b3fa1", lw=3)
    axes[2].set(title=r"stored map $K=f(G)$", xlabel="gate state", ylabel="coefficient")
    for ax in axes:
        ax.grid(alpha=0.2)
        ax.set_ylim(-0.03, 1.03)
    fig.tight_layout(w_pad=2.0)
    return _save_figure(fig, "03_time_to_gate_mapping.png")


def load_real_example() -> dict[str, np.ndarray]:
    raw = parse_ngspice_raw(REAL_EXAMPLE_RAW)
    return {
        "time_ns": np.asarray(raw["time"], dtype=float) * 1e9,
        "input_v": np.asarray(raw["v(in_dig)"], dtype=float),
        "pad_v": np.asarray(raw["v(pad)"], dtype=float),
        "gup": np.asarray(raw["v(xdrv.gup)"], dtype=float),
        "gdn": np.asarray(raw["v(xdrv.gdn)"], dtype=float),
        "ku_gate": np.asarray(raw["v(xdrv.kugate)"], dtype=float),
        "kd_gate": np.asarray(raw["v(xdrv.kdgate)"], dtype=float),
        "ku": np.asarray(raw["v(xdrv.ku)"], dtype=float),
        "kd": np.asarray(raw["v(xdrv.kd)"], dtype=float),
    }


def input_edges(data: dict[str, np.ndarray]) -> tuple[float, float]:
    t = data["time_ns"]
    high = data["input_v"] >= 1.65
    changes = np.where(np.diff(high.astype(int)) != 0)[0]
    if len(changes) < 2:
        raise RuntimeError("The real short-pulse example does not contain two input edges")
    return float(t[changes[0] + 1]), float(t[changes[1] + 1])


def make_real_gate_state_figure(data: dict[str, np.ndarray]) -> Path:
    t = data["time_ns"]
    rise, reverse = input_edges(data)
    window = (t >= rise - 0.5) & (t <= 10.5)
    tw = t[window]
    gup = data["gup"][window]
    gdn = data["gdn"][window]
    input_v = data["input_v"][window]
    t_gup_turn = float(tw[np.argmax(gup)])
    t_gdn_turn = float(tw[np.argmin(gdn)])

    fig, axes = plt.subplots(3, 1, figsize=(10.2, 6.0), sharex=True)
    axes[0].plot(tw, input_v, color="#202020", lw=2.5)
    axes[0].set_ylabel("input (V)")
    axes[0].set_ylim(-0.15, 3.55)
    axes[1].plot(tw, gup, color="#2b7a43", lw=3, label="GUP")
    axes[1].scatter([t_gup_turn], [float(np.max(gup))], color="#2b7a43", s=35, zorder=5)
    axes[1].annotate(
        f"GUP reverses continuously\nstate = {np.max(gup):.3f}",
        (t_gup_turn, float(np.max(gup))),
        xytext=(t_gup_turn + 0.45, float(np.max(gup)) * 0.72),
        fontsize=10,
        arrowprops={"arrowstyle": "->", "color": "#2b7a43"},
    )
    axes[1].set_ylabel("GUP")
    axes[1].set_ylim(-0.008, 0.125)
    axes[2].plot(tw, gdn, color="#2467ad", lw=3, label="GDN")
    axes[2].scatter([t_gdn_turn], [float(np.min(gdn))], color="#2467ad", s=35, zorder=5)
    axes[2].annotate(
        f"GDN reverses continuously\nstate = {np.min(gdn):.4f}",
        (t_gdn_turn, float(np.min(gdn))),
        xytext=(t_gdn_turn - 2.0, 0.22),
        fontsize=10,
        arrowprops={"arrowstyle": "->", "color": "#2467ad"},
    )
    axes[2].set_ylabel("GDN")
    axes[2].set_xlabel("time (ns)")
    axes[2].set_ylim(-0.05, 1.08)
    for ax in axes:
        ax.axvline(rise, color="#777777", ls="--", lw=1.2)
        ax.axvline(reverse, color="#ca3030", ls="--", lw=1.2)
        ax.grid(alpha=0.20)
    axes[0].text(rise + 0.04, 3.15, "rise", color="#555555", fontsize=10)
    axes[0].text(reverse + 0.04, 3.15, "reverse", color="#ca3030", fontsize=10)
    fig.suptitle("Real cached ngspice example: 1 ns high pulse", fontsize=14, fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    return _save_figure(fig, "04_real_gup_gdn_reversal.png")


def make_real_mapped_waveform_figure(data: dict[str, np.ndarray]) -> Path:
    t = data["time_ns"]
    rise, reverse = input_edges(data)
    window = (t >= rise - 0.5) & (t <= 10.5)
    tw = t[window]
    gup = data["gup"][window]
    gdn = data["gdn"][window]
    ku_gate = data["ku_gate"][window]
    kd_gate = data["kd_gate"][window]
    ku_step_idx = int(np.argmax(np.abs(np.diff(ku_gate))))
    kd_step_idx = int(np.argmax(np.abs(np.diff(kd_gate))))
    ku_step = float(ku_gate[ku_step_idx + 1] - ku_gate[ku_step_idx])
    kd_step = float(kd_gate[kd_step_idx + 1] - kd_gate[kd_step_idx])

    fig, axes = plt.subplots(3, 1, figsize=(10.2, 6.1), sharex=True)
    axes[0].plot(tw, data["input_v"][window], color="#202020", lw=2.4)
    axes[0].set_ylabel("input (V)")
    axes[0].set_ylim(-0.15, 3.55)
    axes[1].plot(tw, gup, color="#2b7a43", lw=3, label="GUP state")
    axes[1].plot(tw, ku_gate, color="#7b3fa1", lw=2.6, label="Ku(GUP)")
    axes[1].set_ylabel("pullup")
    axes[1].set_ylim(-0.008, 0.125)
    axes[1].legend(loc="upper right", frameon=False, ncol=2)
    axes[1].annotate(
        f"direction-map change\nDelta Ku = {ku_step:+.3f}",
        (float(tw[ku_step_idx + 1]), float(ku_gate[ku_step_idx + 1])),
        xytext=(float(tw[ku_step_idx + 1]) + 0.5, 0.095),
        color="#7b3fa1",
        fontsize=9.5,
        arrowprops={"arrowstyle": "->", "color": "#7b3fa1"},
    )
    axes[2].plot(tw, gdn, color="#2467ad", lw=3, label="GDN state")
    axes[2].plot(tw, kd_gate, color="#e68613", lw=2.6, label="Kd(GDN)")
    axes[2].axhline(0.0, color="#555555", lw=0.8)
    axes[2].set_ylabel("pulldown")
    axes[2].set_xlabel("time (ns)")
    axes[2].set_ylim(-0.10, 1.10)
    axes[2].legend(loc="upper right", frameon=False, ncol=2)
    axes[2].annotate(
        f"direction-map change\nDelta Kd = {kd_step:+.3f}",
        (float(tw[kd_step_idx + 1]), float(kd_gate[kd_step_idx + 1])),
        xytext=(float(tw[kd_step_idx + 1]) + 0.55, 0.72),
        color="#e68613",
        fontsize=9.5,
        arrowprops={"arrowstyle": "->", "color": "#e68613"},
    )
    for ax in axes:
        ax.axvline(rise, color="#777777", ls="--", lw=1.2)
        ax.axvline(reverse, color="#ca3030", ls="--", lw=1.2)
        ax.grid(alpha=0.20)
    axes[0].text(rise + 0.04, 3.15, "rise", color="#555555", fontsize=10)
    axes[0].text(reverse + 0.04, 3.15, "reverse", color="#ca3030", fontsize=10)
    fig.suptitle("Real state-to-coefficient waveforms: states are continuous, direct maps are not perfect", fontsize=14, fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    return _save_figure(fig, "05_real_state_to_coefficient_reversal.png")


def make_appendix_source_images() -> list[Path]:
    """Create presentation-sized copies without modifying the source scans."""
    ASSETS.mkdir(parents=True, exist_ok=True)
    output: list[Path] = []
    for idx in range(1, 4):
        source = SOURCE_PAGES / f"page_{idx:02d}.png"
        target = ASSETS / f"source_page_{idx:02d}_deck.jpg"
        if source.exists():
            with Image.open(source) as image:
                image = image.convert("RGB")
                image.thumbnail((1200, 1600), Image.Resampling.LANCZOS)
                image.save(target, quality=86, optimize=True, progressive=True)
        output.append(target)
    return output


def add_line(slide, x1: float, y1: float, x2: float, y2: float, *, color: RGBColor, width: float = 1.8):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_rc_memory_diagram(deck: GreenDeck, slide, x: float, y: float, scale: float = 1.0) -> None:
    green = deck.theme.green
    dark = deck.theme.dark
    deck.add_box(slide, "Gtarget\n0 or 1", x, y + 0.45 * scale, 1.25 * scale, 0.75 * scale, size=14, bold=True)
    add_line(slide, x + 1.25 * scale, y + 0.82 * scale, x + 1.65 * scale, y + 0.82 * scale, color=dark)
    resistor = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x + 1.65 * scale),
        Inches(y + 0.60 * scale),
        Inches(1.05 * scale),
        Inches(0.44 * scale),
    )
    resistor.fill.solid()
    resistor.fill.fore_color.rgb = deck.theme.white
    resistor.line.color.rgb = green
    resistor.line.width = Pt(1.8)
    resistor.text_frame.text = "R"
    p = resistor.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.font.name = deck.theme.body_font
    p.font.size = Pt(16)
    p.font.bold = True
    add_line(slide, x + 2.70 * scale, y + 0.82 * scale, x + 3.25 * scale, y + 0.82 * scale, color=dark)
    node = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x + 3.18 * scale),
        Inches(y + 0.75 * scale),
        Inches(0.14 * scale),
        Inches(0.14 * scale),
    )
    node.fill.solid()
    node.fill.fore_color.rgb = green
    node.line.color.rgb = green
    add_line(slide, x + 3.25 * scale, y + 0.82 * scale, x + 3.25 * scale, y + 1.42 * scale, color=dark)
    add_line(slide, x + 2.90 * scale, y + 1.42 * scale, x + 3.60 * scale, y + 1.42 * scale, color=dark, width=2.4)
    add_line(slide, x + 2.90 * scale, y + 1.62 * scale, x + 3.60 * scale, y + 1.62 * scale, color=dark, width=2.4)
    add_line(slide, x + 3.25 * scale, y + 1.62 * scale, x + 3.25 * scale, y + 2.00 * scale, color=dark)
    add_line(slide, x + 2.95 * scale, y + 2.00 * scale, x + 3.55 * scale, y + 2.00 * scale, color=dark)
    add_line(slide, x + 3.05 * scale, y + 2.12 * scale, x + 3.45 * scale, y + 2.12 * scale, color=dark)
    add_line(slide, x + 3.14 * scale, y + 2.24 * scale, x + 3.36 * scale, y + 2.24 * scale, color=dark)
    deck.add_text(slide, "GUP or GDN", x + 3.40 * scale, y + 0.56 * scale, 1.55 * scale, 0.35 * scale, size=15, bold=True, color=green)
    deck.add_text(slide, "C stores the state", x + 3.65 * scale, y + 1.42 * scale, 1.7 * scale, 0.45 * scale, size=13, color=deck.theme.gray)


def add_flow(deck: GreenDeck, slide, labels: list[str], *, y: float, x0: float = 0.65, box_w: float = 2.05) -> None:
    gap = (12.0 - len(labels) * box_w) / max(1, len(labels) - 1)
    x = x0
    for idx, label in enumerate(labels):
        deck.add_box(slide, label, x, y, box_w, 0.78, size=14, bold=True)
        if idx < len(labels) - 1:
            deck.add_arrow(slide, x + box_w, y + 0.39, x + box_w + gap, y + 0.39)
        x += box_w + gap


def build() -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    problem_fig = make_problem_figure()
    tau_fig = make_tau_figure()
    mapping_fig = make_mapping_figure()
    real_example = load_real_example()
    real_gate_state_fig = make_real_gate_state_figure(real_example)
    real_mapped_waveform_fig = make_real_mapped_waveform_figure(real_example)
    appendix_pages = make_appendix_source_images()

    with (OUT / "real_short_pulse_1ns_high_waveforms.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["time_ns", "input_v", "gup", "gdn", "ku_of_gup", "kd_of_gdn", "final_ku", "final_kd", "pad_v"])
        writer.writerows(
            zip(
                real_example["time_ns"],
                real_example["input_v"],
                real_example["gup"],
                real_example["gdn"],
                real_example["ku_gate"],
                real_example["kd_gate"],
                real_example["ku"],
                real_example["kd"],
                real_example["pad_v"],
            )
        )

    renderer = EquationRenderer(backend="auto")
    deck = GreenDeck(equation_renderer=renderer)
    equations: list[tuple[int, str, str]] = []

    def eq(slide, tex: str, x: float, y: float, w: float, h: float, *, size: float = 30):
        _, asset = deck.add_equation(slide, tex, x, y, w, h, font_size_pt=size)
        equations.append((len(deck.prs.slides), tex, str(asset.path.relative_to(ROOT))))

    source = "Source: handwritten notes, IBIS kukd .pdf (2026-07-22)"

    deck.set_title_slide(
        "Continuous Gate-State Model for IBIS Ku/Kd",
        "From handwritten RC-state derivation to runtime ngspice behavior\nIBIS short-pulse study",
        notes=(
            "This deck transcribes the handwritten IBIS Ku/Kd notes into a structured explanation. "
            "The central idea is to replace elapsed-time table position with two continuous stored "
            "states: GUP for pullup drive and GDN for pulldown drive."
        ),
    )

    slide = deck.add_slide(
        "1. The original replay loses continuity at a reversal",
        section="Problem",
        notes=(
            "The legacy model detects an edge and replays a complete Ku/Kd transition table from its "
            "beginning. If the input reverses before that table finishes, switching to the opposite "
            "table at t=0 can assign a different coefficient immediately. The red conceptual curve "
            "shows this reset. The green curve shows the behavior we want: preserve the current state "
            "and only reverse its direction of motion."
        ),
    )
    deck.add_picture_contain(slide, problem_fig, 0.65, 1.15, 7.35, 4.65, border=False)
    eq(
        slide,
        r"K(t_{\mathrm{rev}}^- )\;\neq\;K_{\mathrm{opposite}}(0)",
        8.15,
        1.50,
        4.3,
        0.8,
        size=28,
    )
    deck.add_bullets(
        slide,
        [
            "A long pulse lets the first table finish, so restart looks acceptable.",
            "A short pulse reverses while Ku and Kd are still between endpoints.",
            "The missing quantity is a continuous memory of transition progress.",
        ],
        8.0,
        2.7,
        4.4,
        2.3,
        size=16,
    )
    deck.add_takeaway(slide, "The problem is not the table values themselves; it is how transition state is carried across a new edge.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "2. A capacitor voltage is a natural continuous memory",
        section="State",
        notes=(
            "The handwritten circuit uses an RC node as memory. The target is either zero or one. "
            "The capacitor voltage cannot jump, so it records how far the transition has progressed. "
            "In the generated model this stored voltage is called GUP or GDN rather than Vc."
        ),
    )
    add_rc_memory_diagram(deck, slide, 0.75, 1.45, 1.25)
    eq(slide, r"I_C=C\frac{dV_C}{dt}=\frac{V_{\mathrm{target}}-V_C}{R}", 6.45, 1.35, 5.8, 0.85, size=30)
    eq(slide, r"\frac{dV_C}{dt}=\frac{V_{\mathrm{target}}-V_C}{RC}", 6.45, 2.55, 5.8, 0.85, size=30)
    deck.add_box(
        slide,
        "Vc,PU becomes GUP\nVc,PD becomes GDN",
        7.30,
        4.05,
        4.0,
        0.95,
        fill=deck.theme.pale_green,
        bold=True,
        size=18,
    )
    deck.add_text(slide, "The solver stores capacitor charge between timesteps.", 1.0, 4.65, 4.8, 0.55, size=18, bold=True, color=deck.theme.green, align=PP_ALIGN.CENTER)
    deck.add_takeaway(slide, "GUP and GDN are not recomputed from elapsed time; they persist as capacitor voltages.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "3. The state follows a first-order RC equation",
        section="Equation",
        notes=(
            "Defining tau as RC gives the first-order state equation. The closed-form solution starts "
            "from the state value G0 present when the command changes. This initial value is why the "
            "model can reverse from the middle rather than restart at an endpoint."
        ),
    )
    eq(slide, r"\tau=RC", 0.85, 1.30, 2.7, 0.8, size=34)
    eq(slide, r"\frac{dG}{dt}=\frac{G_{\mathrm{target}}-G}{\tau}", 3.30, 1.15, 5.9, 1.05, size=36)
    eq(
        slide,
        r"G(t)=G_{\mathrm{target}}+\left(G_0-G_{\mathrm{target}}\right)e^{-(t-t_0)/\tau}",
        1.25,
        2.75,
        10.7,
        1.15,
        size=34,
    )
    deck.add_box(slide, "G0 is the state already stored at the command edge", 2.2, 4.45, 8.9, 0.72, bold=True, size=18)
    deck.add_bullets(
        slide,
        ["Changing the target changes the derivative.", "Changing direction selects a different tau.", "Neither action resets G itself."],
        2.35,
        5.35,
        8.5,
        0.95,
        size=16,
        spacing=3,
    )
    deck.add_takeaway(slide, "The state value is continuous; only its destination and rate change at an edge.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "4. Tau is the characteristic transition time",
        section="Calibration",
        notes=(
            "For a normalized first-order rising response, one tau reaches 63.2 percent of the total "
            "change. For a falling response, one tau leaves 36.8 percent. The handwritten derivation "
            "uses the time where the normalized original Ku/Kd transition reaches this level."
        ),
    )
    deck.add_picture_contain(slide, tau_fig, 0.65, 1.2, 7.1, 4.75, border=False)
    eq(slide, r"G_{\mathrm{rise}}(t)=1-e^{-t/\tau}", 7.85, 1.45, 4.55, 0.7, size=28)
    eq(slide, r"G_{\mathrm{fall}}(t)=e^{-t/\tau}", 7.85, 2.45, 4.55, 0.7, size=28)
    eq(slide, r"G_{\mathrm{rise}}(\tau)=1-e^{-1}\approx0.632", 7.65, 3.55, 4.95, 0.75, size=25)
    deck.add_box(slide, "Read tau from the 63.2% point of each normalized transition.", 8.0, 4.75, 4.25, 0.8, bold=True, size=16)
    deck.add_takeaway(slide, "Tau is the rate parameter that reproduces each original coefficient transition.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "5. Four physical directions require four time constants",
        section="Calibration",
        notes=(
            "Pullup and pulldown do not have to be symmetric. An input rise turns the pullup on and "
            "turns the pulldown off. An input fall turns the pullup off and the pulldown on. Each path "
            "therefore gets its own tau, extracted from the corresponding original coefficient table."
        ),
    )
    deck.add_text(slide, "Input rises", 0.75, 1.20, 2.3, 0.4, size=20, bold=True, color=deck.theme.green)
    deck.add_box(slide, "GUP: 0 -> 1\npullup ON", 0.75, 1.75, 2.55, 0.95, bold=True)
    deck.add_box(slide, "GDN: 1 -> 0\npulldown OFF", 3.65, 1.75, 2.55, 0.95, bold=True)
    eq(slide, r"\tau_{\mathrm{PU,on}}", 0.75, 3.0, 2.65, 0.7, size=27)
    eq(slide, r"\tau_{\mathrm{PD,off}}", 3.55, 3.0, 2.85, 0.7, size=27)
    deck.add_text(slide, "Input falls", 6.85, 1.20, 2.3, 0.4, size=20, bold=True, color=deck.theme.green)
    deck.add_box(slide, "GUP: 1 -> 0\npullup OFF", 6.85, 1.75, 2.55, 0.95, bold=True)
    deck.add_box(slide, "GDN: 0 -> 1\npulldown ON", 9.75, 1.75, 2.55, 0.95, bold=True)
    eq(slide, r"\tau_{\mathrm{PU,off}}", 6.75, 3.0, 2.85, 0.7, size=27)
    eq(slide, r"\tau_{\mathrm{PD,on}}", 9.65, 3.0, 2.85, 0.7, size=27)
    deck.add_bullets(
        slide,
        [
            "Each tau comes from the corresponding original Ku/Kd transition.",
            "No symmetry is assumed between pullup, pulldown, rising, and falling.",
            "All values are derived offline from the original IBIS/pybis coefficient tables.",
        ],
        1.15,
        4.45,
        11.0,
        1.35,
        size=17,
        spacing=5,
    )
    deck.add_takeaway(slide, "Direction-specific timing preserves real pullup/pulldown asymmetry.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "6. Offline: replace table time with gate state",
        section="Mapping",
        notes=(
            "For every sample in an original Ku or Kd table, reconstruct the gate-state value that "
            "would exist at the same time using the fitted tau. Pair that gate-state value "
            "with the original coefficient. The resulting point cloud becomes a PWL function K=f(G). "
            "The original coefficient information is preserved; only its independent variable changes "
            "from elapsed time to stored state."
        ),
    )
    deck.add_picture_contain(slide, mapping_fig, 0.45, 1.15, 8.15, 4.65, border=False)
    eq(slide, r"t_i\;\longrightarrow\;G_i=G(t_i)\;\longrightarrow\;(G_i,K_i)", 8.65, 1.55, 3.95, 0.85, size=25)
    eq(slide, r"K_u=f_{\mathrm{PU,dir}}(GUP)", 8.65, 2.75, 3.95, 0.65, size=25)
    eq(slide, r"K_d=f_{\mathrm{PD,dir}}(GDN)", 8.65, 3.70, 3.95, 0.65, size=25)
    deck.add_box(slide, "Four original tables -> four direction-aware state maps", 8.60, 4.75, 4.0, 0.8, bold=True, size=15)
    deck.add_takeaway(slide, "Ku/Kd tables are still used; they become functions of continuous state rather than elapsed time.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "7. Offline preparation and runtime simulation are separate",
        section="Workflow",
        notes=(
            "Everything in the top row happens once while generating the pybis SPICE model: extract "
            "coefficient tables, fit four taus, build the PWL state maps, and write the netlist. "
            "During simulation, the actual pulse history drives the targets. Ngspice integrates GUP "
            "and GDN and evaluates the precomputed maps at each timestep."
        ),
    )
    deck.add_text(slide, "OFFLINE - model generation", 0.65, 1.12, 4.0, 0.4, size=19, bold=True, color=deck.theme.green)
    add_flow(deck, slide, ["Extract Ku/Kd(t)", "Fit four taus", "Build K=f(G)", "Write .sp"], y=1.65)
    deck.add_text(slide, "RUNTIME - ngspice", 0.65, 3.25, 4.0, 0.4, size=19, bold=True, color=deck.theme.blue)
    add_flow(deck, slide, ["Detect edge", "Change targets", "Integrate GUP/GDN", "Evaluate Ku/Kd"], y=3.80)
    deck.add_box(slide, "Actual pulse width and interruption history exist only at runtime.", 3.2, 5.25, 6.9, 0.72, fill=deck.theme.pale_green, bold=True, size=17)
    deck.add_takeaway(slide, "Offline fitting supplies constants and maps; runtime integration supplies the actual state history.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "8. Runtime example: a normal rising input edge",
        section="Runtime",
        notes=(
            "Starting from stable low, GUP is zero and GDN is one. A rising command eventually changes "
            "the targets to one and zero. The states then "
            "move continuously with PU-on and PD-off time constants. Ku and Kd are read from the "
            "corresponding directional maps."
        ),
    )
    deck.add_box(slide, "Before edge\nGUP=0, GDN=1", 0.70, 1.35, 2.45, 0.85, bold=True)
    deck.add_arrow(slide, 3.15, 1.78, 4.00, 1.78)
    deck.add_box(slide, "Command targets\nGUPtarget=1, GDNtarget=0", 4.00, 1.35, 3.15, 0.85, bold=True)
    deck.add_arrow(slide, 7.15, 1.78, 8.00, 1.78)
    deck.add_box(slide, "Continuous motion\nGUP rises, GDN falls", 8.00, 1.35, 3.35, 0.85, bold=True)
    eq(slide, r"\frac{dGUP}{dt}=\frac{1-GUP}{\tau_{\mathrm{PU,on}}}", 0.95, 3.0, 5.2, 0.85, size=29)
    eq(slide, r"\frac{dGDN}{dt}=\frac{0-GDN}{\tau_{\mathrm{PD,off}}}", 6.55, 3.0, 5.2, 0.85, size=29)
    eq(slide, r"K_u=f_{\mathrm{PU,on}}(GUP),\qquad K_d=f_{\mathrm{PD,off}}(GDN)", 2.0, 4.55, 9.3, 0.9, size=29)
    deck.add_takeaway(slide, "For a complete transition, the new state model reconstructs the original Ku/Kd path.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "9. Runtime reversal: targets and rates switch, states do not",
        section="Runtime",
        notes=(
            "Suppose the input falls before the rising transition settles. At the reverse edge, the "
            "targets flip and the model selects PU-off and PD-on taus. GUP and GDN themselves remain "
            "exactly at the values accumulated up to that instant. Their derivatives reverse sign, "
            "which creates a corner in slope but no jump in state. The coefficients remain continuous "
            "if the directional maps agree at the switching state."
        ),
    )
    eq(slide, r"GUP(t_{\mathrm{rev}}^+)=GUP(t_{\mathrm{rev}}^-)", 0.85, 1.25, 5.5, 0.7, size=27)
    eq(slide, r"GDN(t_{\mathrm{rev}}^+)=GDN(t_{\mathrm{rev}}^-)", 6.75, 1.25, 5.5, 0.7, size=27)
    deck.add_box(slide, "falling edge", 0.75, 2.65, 1.65, 0.65, fill=deck.theme.light_red, line=deck.theme.red, bold=True)
    deck.add_arrow(slide, 2.40, 2.98, 3.35, 2.98, color=deck.theme.red)
    deck.add_box(slide, "targets flip\nGUPtarget: 1 -> 0\nGDNtarget: 0 -> 1", 3.35, 2.35, 2.70, 1.25, bold=True)
    deck.add_arrow(slide, 6.05, 2.98, 7.00, 2.98)
    deck.add_box(slide, "taus switch\nPU on -> PU off\nPD off -> PD on", 7.00, 2.35, 2.55, 1.25, bold=True)
    deck.add_arrow(slide, 9.55, 2.98, 10.50, 2.98)
    deck.add_box(slide, "integration continues\nfrom current GUP/GDN", 10.50, 2.45, 2.15, 1.05, bold=True, size=14)
    eq(slide, r"\frac{dGUP}{dt}=\frac{0-GUP}{\tau_{\mathrm{PU,off}}},\qquad\frac{dGDN}{dt}=\frac{1-GDN}{\tau_{\mathrm{PD,on}}}", 1.4, 4.45, 10.6, 0.95, size=29)
    deck.add_takeaway(slide, "A reversal changes the future trajectory, not the state already accumulated.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "10. Real example: GUP and GDN preserve state through reversal",
        section="Waveforms",
        notes=(
            "This is cached ngspice data from the real io_buf short_pulse_1ns_high case using the "
            "two_state_directional_residual model. The input rises near 5 ns and reverses near 6 ns. "
            "GUP and GDN are actual internal capacitor voltages. Each reaches its own turning point "
            "without a vertical jump. This slide focuses only on continuity and does not introduce "
            "additional timing mechanisms beyond the handwritten derivation."
        ),
    )
    deck.add_picture_contain(slide, real_gate_state_fig, 0.75, 1.05, 11.9, 5.15, border=False)
    deck.add_takeaway(slide, "The two states carry different histories and reverse independently, but neither state is reset.")
    deck.add_source(slide, "Cached data: io_buf_two_state_gate_model_2026-06-30 / short_pulse_1ns_high / directional_residual")

    slide = deck.add_slide(
        "11. Real example: state continuity does not guarantee Ku/Kd continuity",
        section="Waveforms",
        notes=(
            "The upper mapped waveform compares GUP with the direct pullup map output Ku(GUP). The "
            "lower waveform compares GDN with Kd(GDN). These are actual runtime signals, not analytic "
            "sketches. GUP and GDN are continuous, but selecting between directional maps creates sharp "
            "changes in the direct mapped coefficients: about -0.063 in Ku(GUP) and +0.069 in Kd(GDN) "
            "at their largest adjacent-sample changes. The final Ku/Kd state nodes smooth these changes, "
            "but this figure shows that the directional maps themselves still need alignment."
        ),
    )
    deck.add_picture_contain(slide, real_mapped_waveform_fig, 0.75, 1.05, 11.9, 5.15, border=False)
    deck.add_takeaway(slide, "Continuous GUP/GDN solve state reset; directional-map alignment remains necessary for continuous Ku/Kd.")
    deck.add_source(slide, "Underlying samples: results/ibis_kukd_handwritten_notes_deck/real_short_pulse_1ns_high_waveforms.csv")

    slide = deck.add_slide(
        "12. Ngspice stores GUP and GDN as capacitor voltages",
        section="Implementation",
        notes=(
            "The capacitor is the memory element. The behavioral current source produces the first-order "
            "derivative. With a one-picofarad numerical capacitor, ngspice integrates the state at every "
            "timestep. The sign shown follows the current orientation of the generated behavioral source. "
            "R_eff is implicit because tau divided by C is the equivalent resistance."
        ),
    )
    deck.add_code_box(
        slide,
        "* Continuous pullup gate state\n"
        "CGUP   GUP   0   1p   ic=0\n"
        "BGUP   GUP   0   I=-1p*(V(GUPTARGET)-V(GUP))/tau_selected\n\n"
        "* Continuous pulldown gate state\n"
        "CGDN   GDN   0   1p   ic=1\n"
        "BGDN   GDN   0   I=-1p*(V(GDNTARGET)-V(GDN))/tau_selected",
        0.65,
        1.25,
        7.25,
        3.3,
        size=12.3,
    )
    eq(slide, r"C\frac{dG}{dt}=\frac{G_{\mathrm{target}}-G}{R_{\mathrm{eff}}}", 8.05, 1.35, 4.55, 0.75, size=27)
    eq(slide, r"R_{\mathrm{eff}}=\frac{\tau}{C}", 8.25, 2.55, 4.1, 0.75, size=30)
    deck.add_box(slide, "C = 1 pF is a numerical state capacitor; tau sets the dynamics.", 8.05, 3.75, 4.45, 0.85, bold=True, size=15)
    deck.add_text(slide, "At each solver step: Gnew = Gold + (dG/dt) Delta t", 1.20, 5.15, 10.8, 0.45, size=19, bold=True, color=deck.theme.green, align=PP_ALIGN.CENTER)
    deck.add_takeaway(slide, "The previous capacitor voltage is the initial condition for the next timestep and for any reversal.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "13. The stored state selects Ku/Kd, which scale IBIS currents",
        section="Implementation",
        notes=(
            "GUP and GDN are hidden drive states, not replacements for Ku and Kd. The precomputed PWL "
            "maps convert gate state into the effective coefficients. Those coefficients retain their "
            "original IBIS role: scaling the full pullup and pulldown current-voltage tables."
        ),
    )
    add_flow(deck, slide, ["Input history", "GUP / GDN", "Ku / Kd maps", "IBIS V-I current"], y=1.35)
    eq(slide, r"K_u=f_{\mathrm{PU,dir}}(GUP),\qquad K_d=f_{\mathrm{PD,dir}}(GDN)", 1.2, 2.75, 10.9, 0.85, size=30)
    eq(slide, r"I_{\mathrm{PU}}(t)=K_u(t)\,I_{\mathrm{PU,full}}(V)", 0.85, 4.15, 5.7, 0.8, size=27)
    eq(slide, r"I_{\mathrm{PD}}(t)=K_d(t)\,I_{\mathrm{PD,full}}(V)", 6.75, 4.15, 5.7, 0.8, size=27)
    deck.add_code_box(slide, "B3 DIE PULLUP_REF I={V(Ku)*pwl(...)}\nB4 DIE PULLDOWN_REF I={V(Kd)*pwl(...)}", 3.25, 5.25, 6.75, 0.78, size=12)
    deck.add_takeaway(slide, "The innovation is state tracking; the established IBIS coefficient-scaled V-I structure remains intact.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "14. What the handwritten model establishes",
        section="Summary",
        notes=(
            "This is the complete conceptual chain in the notes. It explains how to eliminate elapsed-time "
            "restart while continuing to use the original Ku/Kd information. It does not by itself prove "
            "that one first-order state per network reproduces every interrupted transition. That must be "
            "checked against normal-table reconstruction, short pulses, and transistor-level pad behavior."
        ),
    )
    deck.add_box(slide, "1\nStore progress", 0.65, 1.35, 2.15, 1.0, bold=True, size=18)
    deck.add_box(slide, "2\nFit four paths", 3.05, 1.35, 2.15, 1.0, bold=True, size=18)
    deck.add_box(slide, "3\nMap K=f(G)", 5.45, 1.35, 2.15, 1.0, bold=True, size=18)
    deck.add_box(slide, "4\nIntegrate runtime", 7.85, 1.35, 2.15, 1.0, bold=True, size=18)
    deck.add_box(slide, "5\nScale IBIS V-I", 10.25, 1.35, 2.15, 1.0, bold=True, size=18)
    deck.add_text(slide, "Strengths", 0.85, 3.05, 2.2, 0.4, size=20, bold=True, color=deck.theme.green)
    deck.add_bullets(
        slide,
        ["Continuous mid-transition state", "No elapsed-time reset", "Preserves original coefficient tables", "Naturally handles arbitrary pulse width"],
        0.85,
        3.55,
        5.3,
        2.0,
        size=17,
    )
    deck.add_text(slide, "Validation still required", 6.75, 3.05, 3.2, 0.4, size=20, bold=True, color=deck.theme.red)
    deck.add_bullets(
        slide,
        ["Normal Ku/Kd reconstruction", "Directional-map continuity", "Short-high and short-low coefficients", "Pad comparison against native IBIS and transistor SPICE"],
        6.75,
        3.55,
        5.5,
        2.0,
        size=17,
    )
    deck.add_takeaway(slide, "Continuous state solves the restart mechanism; evidence decides whether the chosen state model is sufficient.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "15. Next steps recorded in the handwritten notes",
        section="Next Steps",
        notes=(
            "This slide restores the explicit next-step list from page three of the handwritten notes. "
            "The first item asks whether the capacitor state should have a narrower role: track only the "
            "correct mid-transition starting point, then let the original coefficient tables handle the "
            "rest. The remaining items call for investigating the unusual HSPICE behavior, checking "
            "termination capacitance sensitivity over roughly 1 fF to 10 pF, creating a transistor "
            "buffer in ADS, and consulting Jianchuan about the inverter model."
        ),
    )
    deck.add_box(
        slide,
        "Model scope question",
        0.70,
        1.25,
        2.55,
        0.62,
        fill=deck.theme.pale_green,
        bold=True,
        size=17,
    )
    deck.add_text(
        slide,
        "Should GUP/GDN only track the correct Ku/Kd starting point at a mid-transition reversal?",
        0.85,
        2.05,
        5.2,
        1.0,
        size=20,
        bold=True,
        color=deck.theme.green,
    )
    deck.add_box(
        slide,
        "Reference and circuit checks",
        6.65,
        1.25,
        3.15,
        0.62,
        fill=deck.theme.pale_green,
        bold=True,
        size=17,
    )
    deck.add_bullets(
        slide,
        [
            "Investigate why the HSPICE result looks unusual.",
            "Sweep termination capacitance from about 1 fF to 10 pF; account for driver strength.",
            "Use ADS to create a transistor-level SPICE buffer.",
            "Talk to Jianchuan about the inverter model.",
        ],
        6.65,
        2.05,
        5.75,
        2.85,
        size=17,
        spacing=7,
    )
    deck.add_box(
        slide,
        "These are open questions and validation tasks, not conclusions.",
        2.4,
        5.25,
        8.4,
        0.72,
        fill=deck.theme.light_red,
        line=deck.theme.red,
        bold=True,
        size=17,
    )
    deck.add_takeaway(slide, "The next experiment should test whether continuous state is best used as the full dynamics or only as the restart coordinate.")
    deck.add_source(slide, "Transcribed from handwritten source, page 3")

    slide = deck.add_slide(
        "Appendix: original handwritten source",
        section="Source",
        notes=(
            "These are the three original handwritten pages used for this deck. They are included so the "
            "transcription and notation can be checked directly."
        ),
    )
    for idx, page in enumerate(appendix_pages, start=1):
        if page.exists():
            deck.add_picture_contain(slide, page, 0.55 + (idx - 1) * 4.15, 1.15, 3.75, 5.35, border=True)
            deck.add_text(slide, f"Page {idx}", 1.65 + (idx - 1) * 4.15, 6.25, 1.6, 0.25, size=11, bold=True, align=PP_ALIGN.CENTER)
    deck.add_source(slide, r"Original PDF: \\minerfiles.mst.edu\dfs\users\sh3qm\Downloads\IBIS kukd .pdf")

    path = deck.save(OUTPUT)
    with (OUT / "equation_manifest.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["slide", "tex", "rendered_asset"])
        writer.writerows(equations)
    return path


if __name__ == "__main__":
    result = build()
    print(result)
