from __future__ import annotations

import csv
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
LOCAL_DEPS = ROOT / ".codex_deps" / "presentation" / "python"
if LOCAL_DEPS.exists():
    sys.path.insert(0, str(LOCAL_DEPS))
sys.path.insert(0, str(ROOT))

import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from eye_diagram import parse_ngspice_raw
from tools.presentation_kit import EquationRenderer, GreenDeck


OUT = ROOT / "results" / "ibis_kukd_handwritten_notes_deck"
ASSETS = OUT / "strict_vc_assets"
SOURCE_PAGES = OUT / "source_pages"
OUTPUT = OUT / "IBIS_KuKd_handwritten_notes_strict_Vc_corrected.pptx"
REAL_EXAMPLE_RAW = (
    ROOT
    / "results"
    / "io_buf_two_state_gate_model_2026-06-30"
    / "cases"
    / "short_pulse_1ns_high"
    / "ngspice_two_state_directional_residual"
    / "short_pulse_1ns_high_ngspice_two_state_directional_residual.raw"
)
MAPPING_EVIDENCE_CSV = (
    ROOT
    / "results"
    / "io_buf_two_state_gate_model_2026-06-30"
    / "gate_mapping_evidence"
    / "gate_mapping_time_data.csv"
)


def save_figure(fig: plt.Figure, name: str) -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    path = ASSETS / name
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def line(slide, x1: float, y1: float, x2: float, y2: float, color: RGBColor, width: float = 1.8):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    return shape


def add_rc_diagram(deck: GreenDeck, slide, x: float, y: float) -> None:
    green = deck.theme.green
    dark = deck.theme.dark
    deck.add_box(slide, "Vtarget\n0 or 1", x, y + 0.40, 1.35, 0.80, size=15, bold=True)
    line(slide, x + 1.35, y + 0.80, x + 1.75, y + 0.80, dark)
    resistor = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x + 1.75),
        Inches(y + 0.58),
        Inches(1.10),
        Inches(0.44),
    )
    resistor.fill.solid()
    resistor.fill.fore_color.rgb = deck.theme.white
    resistor.line.color.rgb = green
    resistor.line.width = Pt(1.8)
    resistor.text_frame.text = "R"
    p = resistor.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.font.name = deck.theme.body_font
    p.font.size = Pt(17)
    p.font.bold = True
    line(slide, x + 2.85, y + 0.80, x + 3.35, y + 0.80, dark)
    node = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x + 3.28),
        Inches(y + 0.73),
        Inches(0.14),
        Inches(0.14),
    )
    node.fill.solid()
    node.fill.fore_color.rgb = green
    node.line.color.rgb = green
    line(slide, x + 3.35, y + 0.80, x + 3.35, y + 1.42, dark)
    line(slide, x + 3.00, y + 1.42, x + 3.70, y + 1.42, dark, 2.4)
    line(slide, x + 3.00, y + 1.62, x + 3.70, y + 1.62, dark, 2.4)
    line(slide, x + 3.35, y + 1.62, x + 3.35, y + 2.02, dark)
    line(slide, x + 3.05, y + 2.02, x + 3.65, y + 2.02, dark)
    line(slide, x + 3.15, y + 2.14, x + 3.55, y + 2.14, dark)
    line(slide, x + 3.24, y + 2.26, x + 3.46, y + 2.26, dark)
    deck.add_text(slide, "Vc", x + 3.50, y + 0.55, 0.8, 0.35, size=18, bold=True, color=green)
    deck.add_text(slide, "C stores Vc", x + 3.78, y + 1.42, 1.55, 0.40, size=14, color=deck.theme.gray)


def make_problem_figure() -> Path:
    t = np.linspace(0.0, 3.0, 900)
    second_edge = 1.15
    rising = 1.0 - np.exp(-t / 0.75)
    restart = rising.copy()
    mask = t >= second_edge
    restart[mask] = np.exp(-(t[mask] - second_edge) / 0.42)
    vc = rising.copy()
    vc_at_edge = 1.0 - np.exp(-second_edge / 0.75)
    vc[mask] = vc_at_edge * np.exp(-(t[mask] - second_edge) / 0.42)

    fig, ax = plt.subplots(figsize=(8.8, 3.2))
    ax.plot(t, rising, color="#777777", lw=2.0, label="unfinished rising table")
    ax.plot(t, restart, color="#ca3030", lw=3.0, label="falling table restarted at its beginning")
    ax.plot(t, vc, color="#2b7a43", lw=3.0, label="Vc continues from its present value")
    ax.axvline(second_edge, color="#202020", ls="--", lw=1.4)
    ax.text(second_edge + 0.05, 0.06, "falling edge", fontsize=10)
    ax.annotate(
        "coefficient restart can jump",
        (second_edge + 0.015, 0.98),
        xytext=(1.62, 1.07),
        color="#ca3030",
        arrowprops={"arrowstyle": "->", "color": "#ca3030"},
        fontsize=10,
    )
    ax.set(xlabel="time", ylabel="normalized value", ylim=(-0.04, 1.13))
    ax.grid(alpha=0.20)
    ax.legend(frameon=False, loc="lower right")
    fig.tight_layout()
    return save_figure(fig, "01_mid_transition_problem.png")


def make_tau_figure() -> Path:
    t = np.linspace(0.0, 4.0, 700)
    rise = 1.0 - np.exp(-t)
    fall = np.exp(-t)
    fig, ax = plt.subplots(figsize=(8.6, 3.2))
    ax.plot(t, rise, color="#2b7a43", lw=3, label=r"rising Vc: $1-e^{-t/\tau}$")
    ax.plot(t, fall, color="#2467ad", lw=3, label=r"falling Vc: $e^{-t/\tau}$")
    ax.axvline(1.0, color="#202020", ls="--", lw=1.4)
    ax.scatter([1.0, 1.0], [1.0 - np.exp(-1.0), np.exp(-1.0)], color=["#2b7a43", "#2467ad"], zorder=5)
    ax.annotate("63.2%", (1.0, 1.0 - np.exp(-1.0)), xytext=(1.18, 0.73), fontsize=11)
    ax.annotate("36.8%", (1.0, np.exp(-1.0)), xytext=(1.18, 0.24), fontsize=11)
    ax.set(xlabel=r"normalized time $t/\tau$", ylabel="normalized Vc", ylim=(-0.03, 1.03))
    ax.grid(alpha=0.20)
    ax.legend(frameon=False, loc="center right")
    fig.tight_layout()
    return save_figure(fig, "02_tau_from_vc.png")


def make_mapping_figure() -> Path:
    with MAPPING_EVIDENCE_CSV.open(newline="", encoding="utf-8") as handle:
        rows = [
            row
            for row in csv.DictReader(handle)
            if row["process"] == "pu_on"
        ]
    rows.sort(key=lambda row: float(row["table_time_ns"]))
    t = np.asarray([float(row["table_time_ns"]) for row in rows])
    vc = np.asarray([float(row["gate_state"]) for row in rows])
    ku_original = np.asarray([float(row["original_coefficient"]) for row in rows])
    ku_mapped = np.asarray(
        [float(row["final_residual_corrected_coefficient"]) for row in rows]
    )
    error = ku_mapped - ku_original
    rmse = float(np.sqrt(np.mean(error**2)))
    max_error = float(np.max(np.abs(error)))
    marker_step = max(1, len(t) // 24)

    fig, axes = plt.subplots(1, 3, figsize=(11.2, 3.4))
    axes[0].plot(
        t,
        ku_original,
        color="#111111",
        lw=3.2,
        label=r"original $Ku_{\mathrm{rising}}(t)$",
    )
    axes[0].plot(
        t,
        ku_mapped,
        color="#ca3030",
        lw=1.8,
        ls="--",
        marker="o",
        ms=3.0,
        markevery=marker_step,
        markerfacecolor="white",
        label=r"mapped $Ku_{\mathrm{rising}}(Vc(t))$",
    )
    axes[0].set(
        title="Round trip on the same time axis",
        xlabel="table time (ns)",
        ylabel="Ku",
    )
    axes[0].legend(frameon=False, fontsize=8.0, loc="lower right")
    axes[0].text(
        0.04,
        0.93,
        f"RMSE {rmse:.5f}\nmax error {max_error:.5f}",
        transform=axes[0].transAxes,
        va="top",
        fontsize=8.5,
        color="#444444",
    )

    axes[1].plot(t, vc, color="#2b7a43", lw=3)
    axes[1].set(
        title=r"1. Build $Vc_{\mathrm{PU,rising}}(t)$",
        xlabel="table time (ns)",
        ylabel="Vc_PU",
    )

    axes[2].scatter(
        vc,
        ku_original,
        color="#111111",
        s=12,
        alpha=0.55,
        label="paired original samples",
        zorder=2,
    )
    axes[2].plot(
        vc,
        ku_mapped,
        color="#ca3030",
        lw=2.4,
        label=r"stored $Ku_{\mathrm{rising}}(Vc_{\mathrm{PU}})$ map",
        zorder=3,
    )
    axes[2].set(
        title="2. Re-index Ku by Vc_PU",
        xlabel="Vc_PU",
        ylabel="Ku",
    )
    axes[2].legend(frameon=False, fontsize=8.0, loc="lower right")
    for ax in axes:
        ax.grid(alpha=0.20)
        ax.set_ylim(-0.03, 1.03)
    fig.suptitle(
        "Actual io_buf mapping: the transfer curve looks different, but reconstructs the original time waveform",
        fontsize=13.0,
        fontweight="bold",
    )
    fig.tight_layout(rect=(0, 0, 1, 0.91), w_pad=1.6)
    return save_figure(fig, "03_ku_time_to_vc_mapping.png")


def make_conceptual_reversal_figure() -> Path:
    t = np.linspace(0.0, 4.0, 1000)
    edge = 1.5
    vc_pu = 1.0 - np.exp(-np.minimum(t, edge) / 1.15)
    pu_edge = float(1.0 - np.exp(-edge / 1.15))
    after = t >= edge
    vc_pu[after] = pu_edge * np.exp(-(t[after] - edge) / 0.55)
    vc_pd = np.exp(-np.minimum(t, edge) / 0.42)
    pd_edge = float(np.exp(-edge / 0.42))
    vc_pd[after] = 1.0 + (pd_edge - 1.0) * np.exp(-(t[after] - edge) / 0.48)
    ku = np.clip(vc_pu**1.45, 0, 1)
    kd = np.clip(vc_pd**1.10, 0, 1)

    fig, axes = plt.subplots(3, 1, figsize=(9.6, 5.5), sharex=True)
    input_v = np.where(t < edge, 1.0, 0.0)
    axes[0].plot(t, input_v, color="#202020", lw=2.5)
    axes[0].set_ylabel("input")
    axes[1].plot(t, vc_pu, color="#2b7a43", lw=3, label="Vc_PU")
    axes[1].plot(t, ku, color="#7b3fa1", lw=2.5, label="Ku(Vc_PU)")
    axes[1].set_ylabel("pullup")
    axes[1].legend(frameon=False, loc="upper right", ncol=2)
    axes[2].plot(t, vc_pd, color="#2467ad", lw=3, label="Vc_PD")
    axes[2].plot(t, kd, color="#e68613", lw=2.5, label="Kd(Vc_PD)")
    axes[2].set_ylabel("pulldown")
    axes[2].set_xlabel("time")
    axes[2].legend(frameon=False, loc="upper right", ncol=2)
    for ax in axes:
        ax.axvline(edge, color="#ca3030", ls="--", lw=1.3)
        ax.grid(alpha=0.20)
    axes[0].text(edge + 0.05, 0.83, "falling edge", color="#ca3030", fontsize=10)
    fig.suptitle("Handwritten-note behavior: Vc does not jump when the input changes direction", fontsize=14, fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    return save_figure(fig, "04_conceptual_vc_reversal.png")


def load_real_example() -> dict[str, np.ndarray]:
    raw = parse_ngspice_raw(REAL_EXAMPLE_RAW)
    return {
        "time_ns": np.asarray(raw["time"], dtype=float) * 1e9,
        "input_v": np.asarray(raw["v(in_dig)"], dtype=float),
        "pad_v": np.asarray(raw["v(pad)"], dtype=float),
        "vc_pu": np.asarray(raw["v(xdrv.gup)"], dtype=float),
        "vc_pd": np.asarray(raw["v(xdrv.gdn)"], dtype=float),
        "ku_vc": np.asarray(raw["v(xdrv.kugate)"], dtype=float),
        "kd_vc": np.asarray(raw["v(xdrv.kdgate)"], dtype=float),
        "ku": np.asarray(raw["v(xdrv.ku)"], dtype=float),
        "kd": np.asarray(raw["v(xdrv.kd)"], dtype=float),
    }


def find_edges(data: dict[str, np.ndarray]) -> tuple[float, float]:
    high = data["input_v"] >= 1.65
    idx = np.where(np.diff(high.astype(int)) != 0)[0]
    if len(idx) < 2:
        raise RuntimeError("Expected a rising and falling edge in the cached example")
    t = data["time_ns"]
    return float(t[idx[0] + 1]), float(t[idx[1] + 1])


def make_real_vc_figure(data: dict[str, np.ndarray]) -> Path:
    t = data["time_ns"]
    rise, fall = find_edges(data)
    keep = (t >= rise - 0.5) & (t <= 10.5)
    tw = t[keep]
    vc_pu = data["vc_pu"][keep]
    vc_pd = data["vc_pd"][keep]
    t_pu_turn = float(tw[np.argmax(vc_pu)])
    t_pd_turn = float(tw[np.argmin(vc_pd)])

    fig, axes = plt.subplots(3, 1, figsize=(10.2, 6.0), sharex=True)
    axes[0].plot(tw, data["input_v"][keep], color="#202020", lw=2.5)
    axes[0].set(ylabel="input (V)", ylim=(-0.15, 3.55))
    axes[1].plot(tw, vc_pu, color="#2b7a43", lw=3)
    axes[1].scatter([t_pu_turn], [float(np.max(vc_pu))], color="#2b7a43", s=35, zorder=5)
    axes[1].annotate(
        f"Vc_PU changes direction\nwithout a jump; Vc={np.max(vc_pu):.3f}",
        (t_pu_turn, float(np.max(vc_pu))),
        xytext=(t_pu_turn + 0.45, 0.073),
        fontsize=10,
        arrowprops={"arrowstyle": "->", "color": "#2b7a43"},
    )
    axes[1].set(ylabel="Vc_PU", ylim=(-0.008, 0.125))
    axes[2].plot(tw, vc_pd, color="#2467ad", lw=3)
    axes[2].scatter([t_pd_turn], [float(np.min(vc_pd))], color="#2467ad", s=35, zorder=5)
    axes[2].annotate(
        f"Vc_PD changes direction\nwithout a jump; Vc={np.min(vc_pd):.4f}",
        (t_pd_turn, float(np.min(vc_pd))),
        xytext=(t_pd_turn - 2.0, 0.22),
        fontsize=10,
        arrowprops={"arrowstyle": "->", "color": "#2467ad"},
    )
    axes[2].set(ylabel="Vc_PD", xlabel="time (ns)", ylim=(-0.05, 1.08))
    for ax in axes:
        ax.axvline(rise, color="#777777", ls="--", lw=1.2)
        ax.axvline(fall, color="#ca3030", ls="--", lw=1.2)
        ax.grid(alpha=0.20)
    axes[0].text(rise + 0.04, 3.15, "rising edge", color="#555555", fontsize=10)
    axes[0].text(fall + 0.04, 3.15, "falling edge", color="#ca3030", fontsize=10)
    fig.suptitle("Cached io_buf example: actual capacitor-voltage waveforms for a 1 ns high pulse", fontsize=14, fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    return save_figure(fig, "05_real_vc_pu_vc_pd.png")


def make_real_k_vc_figure(data: dict[str, np.ndarray]) -> Path:
    t = data["time_ns"]
    rise, fall = find_edges(data)
    keep = (t >= rise - 0.5) & (t <= 10.5)
    tw = t[keep]
    vc_pu = data["vc_pu"][keep]
    vc_pd = data["vc_pd"][keep]
    ku_vc = data["ku_vc"][keep]
    kd_vc = data["kd_vc"][keep]
    ku_idx = int(np.argmax(np.abs(np.diff(ku_vc))))
    kd_idx = int(np.argmax(np.abs(np.diff(kd_vc))))
    ku_step = float(ku_vc[ku_idx + 1] - ku_vc[ku_idx])
    kd_step = float(kd_vc[kd_idx + 1] - kd_vc[kd_idx])

    fig, axes = plt.subplots(3, 1, figsize=(10.2, 6.1), sharex=True)
    axes[0].plot(tw, data["input_v"][keep], color="#202020", lw=2.5)
    axes[0].set(ylabel="input (V)", ylim=(-0.15, 3.55))
    axes[1].plot(tw, vc_pu, color="#2b7a43", lw=3, label="Vc_PU")
    axes[1].plot(tw, ku_vc, color="#7b3fa1", lw=2.6, label="Ku(Vc_PU)")
    axes[1].annotate(
        f"rising-to-falling relationship change\nDelta Ku={ku_step:+.3f}",
        (float(tw[ku_idx + 1]), float(ku_vc[ku_idx + 1])),
        xytext=(float(tw[ku_idx + 1]) + 0.5, 0.095),
        fontsize=9.5,
        color="#7b3fa1",
        arrowprops={"arrowstyle": "->", "color": "#7b3fa1"},
    )
    axes[1].set(ylabel="pullup", ylim=(-0.008, 0.125))
    axes[1].legend(frameon=False, loc="upper right", ncol=2)
    axes[2].plot(tw, vc_pd, color="#2467ad", lw=3, label="Vc_PD")
    axes[2].plot(tw, kd_vc, color="#e68613", lw=2.6, label="Kd(Vc_PD)")
    axes[2].axhline(0.0, color="#555555", lw=0.8)
    axes[2].annotate(
        f"rising-to-falling relationship change\nDelta Kd={kd_step:+.3f}",
        (float(tw[kd_idx + 1]), float(kd_vc[kd_idx + 1])),
        xytext=(float(tw[kd_idx + 1]) + 0.55, 0.72),
        fontsize=9.5,
        color="#e68613",
        arrowprops={"arrowstyle": "->", "color": "#e68613"},
    )
    axes[2].set(ylabel="pulldown", xlabel="time (ns)", ylim=(-0.10, 1.10))
    axes[2].legend(frameon=False, loc="upper right", ncol=2)
    for ax in axes:
        ax.axvline(rise, color="#777777", ls="--", lw=1.2)
        ax.axvline(fall, color="#ca3030", ls="--", lw=1.2)
        ax.grid(alpha=0.20)
    axes[0].text(rise + 0.04, 3.15, "rising edge", color="#555555", fontsize=10)
    axes[0].text(fall + 0.04, 3.15, "falling edge", color="#ca3030", fontsize=10)
    fig.suptitle("Cached io_buf example: Ku(Vc_PU) and Kd(Vc_PD) during the same pulse", fontsize=14, fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    return save_figure(fig, "06_real_ku_vcpu_kd_vcpd.png")


def make_appendix_pages() -> list[Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    output: list[Path] = []
    for idx in range(1, 4):
        source = SOURCE_PAGES / f"page_{idx:02d}.png"
        target = ASSETS / f"source_page_{idx:02d}.jpg"
        if source.exists():
            with Image.open(source) as image:
                image = image.convert("RGB")
                image.thumbnail((1200, 1600), Image.Resampling.LANCZOS)
                image.save(target, quality=86, optimize=True, progressive=True)
        output.append(target)
    return output


def build() -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    problem_fig = make_problem_figure()
    tau_fig = make_tau_figure()
    mapping_fig = make_mapping_figure()
    conceptual_reversal_fig = make_conceptual_reversal_figure()
    real_data = load_real_example()
    real_vc_fig = make_real_vc_figure(real_data)
    real_k_vc_fig = make_real_k_vc_figure(real_data)
    appendix_pages = make_appendix_pages()

    with (OUT / "strict_vc_real_example_waveforms.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["time_ns", "input_v", "vc_pu", "vc_pd", "ku_of_vc_pu", "kd_of_vc_pd", "final_ku", "final_kd", "pad_v"])
        writer.writerows(
            zip(
                real_data["time_ns"],
                real_data["input_v"],
                real_data["vc_pu"],
                real_data["vc_pd"],
                real_data["ku_vc"],
                real_data["kd_vc"],
                real_data["ku"],
                real_data["kd"],
                real_data["pad_v"],
            )
        )

    renderer = EquationRenderer(backend="auto")
    deck = GreenDeck(equation_renderer=renderer)
    equations: list[tuple[int, str, str]] = []

    def eq(slide, tex: str, x: float, y: float, w: float, h: float, size: float = 30):
        _, asset = deck.add_equation(slide, tex, x, y, w, h, font_size_pt=size)
        equations.append((len(deck.prs.slides), tex, str(asset.path.relative_to(ROOT))))

    source = "Source: handwritten notes, IBIS kukd .pdf (2026-07-22)"

    deck.set_title_slide(
        "Continuous Vc Model for IBIS Ku/Kd",
        "Strict transcription of the handwritten derivation\nIBIS short-pulse study",
        notes=(
            "This version follows the handwritten terminology and sequence. Vc_PU and Vc_PD are "
            "capacitor voltages that track pullup and pulldown transition progress. No additional "
            "timing concepts are introduced."
        ),
    )

    slide = deck.add_slide(
        "1. Problem: changing tables in the middle is not continuous",
        section="Problem",
        notes=(
            "The notes begin with the mid-transition problem. A rising edge starts Ku_rising and "
            "Kd_rising. If a falling edge arrives before they settle, restarting Ku_falling and "
            "Kd_falling at the beginning can force new coefficient values rather than continue from "
            "the values already reached."
        ),
    )
    deck.add_picture_contain(slide, problem_fig, 0.65, 1.15, 7.25, 4.65, border=False)
    eq(
        slide,
        r"\left\{Ku_{\mathrm{rising}}(t),Kd_{\mathrm{rising}}(t)\right\}"
        r"\;\longrightarrow\;"
        r"\left\{Ku_{\mathrm{falling}}(0),Kd_{\mathrm{falling}}(0)\right\}",
        7.85,
        1.40,
        4.65,
        1.0,
        24,
    )
    deck.add_bullets(
        slide,
        [
            "The first transition has not reached its endpoint.",
            "The falling tables assume their own starting condition.",
            "A continuous quantity is needed to remember transition progress.",
        ],
        8.05,
        2.85,
        4.25,
        2.1,
        size=16,
    )
    deck.add_takeaway(slide, "The short-pulse problem is a mid-transition starting-point problem.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "2. Use capacitor voltage Vc to keep a continuous transition state",
        section="Vc",
        notes=(
            "The handwritten solution is to use a capacitor voltage. Capacitor voltage cannot jump "
            "instantaneously. Its present value therefore records how far the transition has progressed."
        ),
    )
    add_rc_diagram(deck, slide, 0.85, 1.50)
    deck.add_text(slide, "Vtarget = 0 or 1", 0.95, 4.55, 4.4, 0.45, size=20, bold=True, color=deck.theme.green, align=PP_ALIGN.CENTER)
    eq(slide, r"I_C=C\frac{dV_C}{dt}=\frac{V_{\mathrm{target}}-V_C}{R}", 6.15, 1.35, 6.0, 0.9, 30)
    eq(slide, r"\frac{dV_C}{dt}=\frac{V_{\mathrm{target}}-V_C}{RC}", 6.35, 2.65, 5.6, 0.9, 30)
    deck.add_box(slide, "Vc keeps the present transition position.", 7.0, 4.15, 4.3, 0.82, bold=True, size=18)
    deck.add_takeaway(slide, "The capacitor carries the previous value into the next solver timestep.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "3. Solve the first-order Vc equation",
        section="Equation",
        notes=(
            "The solution starts from V0, the capacitor voltage already present when the target changes. "
            "The same equation therefore covers a full transition and a change in the middle."
        ),
    )
    eq(slide, r"\tau=RC", 0.85, 1.25, 2.8, 0.8, 34)
    eq(slide, r"\frac{dV_C}{dt}=\frac{V_{\mathrm{target}}-V_C}{\tau}", 3.25, 1.10, 6.2, 1.0, 36)
    eq(
        slide,
        r"V_C(t)=V_{\mathrm{target}}+\left(V_0-V_{\mathrm{target}}\right)e^{-t/\tau}",
        1.25,
        2.65,
        10.8,
        1.15,
        35,
    )
    deck.add_box(slide, "V0 is the capacitor voltage already reached.", 2.4, 4.40, 8.5, 0.72, bold=True, size=18)
    deck.add_bullets(
        slide,
        ["Vtarget selects the destination.", "Tau selects how quickly Vc moves.", "Vc itself is not restarted."],
        2.55,
        5.25,
        8.2,
        1.0,
        size=17,
        spacing=3,
    )
    deck.add_takeaway(slide, "A new input edge changes the equation parameters, not the stored Vc value.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "4. Rising and falling are the two Vc solutions",
        section="Equation",
        notes=(
            "For rising, the target is one and the initial value is zero. For falling, the target is "
            "zero and the initial value is one. These are the two normalized capacitor-voltage curves "
            "used throughout the rest of the notes."
        ),
    )
    deck.add_text(slide, "Rising", 1.0, 1.20, 2.2, 0.45, size=22, bold=True, color=deck.theme.green, align=PP_ALIGN.CENTER)
    eq(slide, r"V_C(t)=1-e^{-t/\tau}", 0.75, 1.90, 5.0, 0.9, 34)
    deck.add_text(slide, "Vtarget = 1, V0 = 0", 1.10, 3.00, 4.3, 0.45, size=18, bold=True, align=PP_ALIGN.CENTER)
    deck.add_text(slide, "Falling", 7.35, 1.20, 2.2, 0.45, size=22, bold=True, color=deck.theme.blue, align=PP_ALIGN.CENTER)
    eq(slide, r"V_C(t)=e^{-t/\tau}", 6.55, 1.90, 5.0, 0.9, 34)
    deck.add_text(slide, "Vtarget = 0, V0 = 1", 6.90, 3.00, 4.3, 0.45, size=18, bold=True, align=PP_ALIGN.CENTER)
    deck.add_box(slide, "Both equations keep Vc continuous when the target changes.", 2.5, 4.40, 8.3, 0.80, bold=True, size=19)
    deck.add_takeaway(slide, "Rising moves Vc toward 1; falling moves Vc toward 0.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "5. Obtain tau from the original Ku/Kd tables",
        section="Tau",
        notes=(
            "At one tau, a rising first-order response reaches 63.2 percent. The handwritten method "
            "therefore estimates each tau from the time at which the corresponding normalized Ku or Kd "
            "table reaches the equivalent 63-percent transition point."
        ),
    )
    deck.add_picture_contain(slide, tau_fig, 0.65, 1.20, 7.1, 4.75, border=False)
    eq(slide, r"V_C(\tau)=1-e^{-1}\approx0.632", 7.70, 1.45, 4.8, 0.85, 28)
    deck.add_box(slide, "Example: use the time where Ku_rising reaches about 63% of its total change.", 7.85, 2.75, 4.3, 1.15, bold=True, size=16)
    deck.add_box(slide, "Repeat for the other three tables.", 8.25, 4.40, 3.5, 0.72, bold=True, size=17)
    deck.add_takeaway(slide, "The four tau values come directly from the original coefficient transitions.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "6. The result is four tau values",
        section="Tau",
        notes=(
            "The notes name the four constants by pullup or pulldown and by input transition direction: "
            "PU-rising, PU-falling, PD-rising, and PD-falling."
        ),
    )
    deck.add_text(slide, "Input rising", 0.75, 1.20, 2.4, 0.45, size=21, bold=True, color=deck.theme.green)
    deck.add_box(slide, "Vc_PU: 0 -> 1", 0.75, 1.85, 2.55, 0.80, bold=True, size=18)
    deck.add_box(slide, "Vc_PD: 1 -> 0", 3.65, 1.85, 2.55, 0.80, bold=True, size=18)
    eq(slide, r"\tau_{\mathrm{PU,rising}}", 0.75, 3.05, 2.55, 0.70, 28)
    eq(slide, r"\tau_{\mathrm{PD,rising}}", 3.65, 3.05, 2.55, 0.70, 28)
    deck.add_text(slide, "Input falling", 6.85, 1.20, 2.4, 0.45, size=21, bold=True, color=deck.theme.green)
    deck.add_box(slide, "Vc_PU: 1 -> 0", 6.85, 1.85, 2.55, 0.80, bold=True, size=18)
    deck.add_box(slide, "Vc_PD: 0 -> 1", 9.75, 1.85, 2.55, 0.80, bold=True, size=18)
    eq(slide, r"\tau_{\mathrm{PU,falling}}", 6.85, 3.05, 2.55, 0.70, 28)
    eq(slide, r"\tau_{\mathrm{PD,falling}}", 9.75, 3.05, 2.55, 0.70, 28)
    deck.add_box(slide, "No pullup/pulldown or rising/falling symmetry is assumed.", 2.4, 4.75, 8.4, 0.75, bold=True, size=18)
    deck.add_takeaway(slide, "Each original Ku/Kd table gets its corresponding Vc transition.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "7. Map Ku(t) and Kd(t) to Vc instead of time",
        section="Mapping",
        notes=(
            "The notes pair each original coefficient sample with the Vc value at the same table time. "
            "The left panel is the essential round-trip check: original Ku_rising of time and the mapped "
            "Ku_rising of Vc of time nearly overlap on one time axis. The right-hand transfer curve looks "
            "geometrically different because its horizontal axis is nonlinear Vc, not time; that difference "
            "does not by itself mean the coefficient waveform changed. Repeat the same process for "
            "Ku_falling, Kd_rising, and Kd_falling."
        ),
    )
    deck.add_picture_contain(slide, mapping_fig, 0.45, 1.15, 8.10, 4.65, border=False)
    eq(slide, r"t_i\longrightarrow Vc_i=Vc(t_i)\longrightarrow(Vc_i,K_i)", 8.55, 1.45, 4.0, 0.85, 24)
    eq(slide, r"Ku_{\mathrm{rising}}=f_{\mathrm{PU,rising}}(Vc_{\mathrm{PU}})", 8.55, 2.75, 4.0, 0.72, 22)
    deck.add_box(slide, "Create the same relationship for all four original tables.", 8.55, 4.15, 4.0, 0.95, bold=True, size=16)
    deck.add_takeaway(slide, "The coefficient tables are retained; their lookup coordinate changes from t to Vc.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "8. Before simulation, prepare four Ku/Kd versus Vc relationships",
        section="Preparation",
        notes=(
            "This slide makes explicit the result of the handwritten mapping step. There are four "
            "relationships: pullup rising and falling, plus pulldown rising and falling. These are "
            "prepared before simulation."
        ),
    )
    deck.add_box(slide, "Ku_rising(Vc_PU)", 0.75, 1.40, 2.55, 0.85, bold=True, size=18)
    deck.add_box(slide, "Ku_falling(Vc_PU)", 3.55, 1.40, 2.55, 0.85, bold=True, size=18)
    deck.add_box(slide, "Kd_rising(Vc_PD)", 6.35, 1.40, 2.55, 0.85, bold=True, size=18)
    deck.add_box(slide, "Kd_falling(Vc_PD)", 9.15, 1.40, 2.55, 0.85, bold=True, size=18)
    eq(slide, r"Ku(t)\;\longrightarrow\;Ku(Vc_{\mathrm{PU}})", 1.20, 3.05, 4.7, 0.85, 31)
    eq(slide, r"Kd(t)\;\longrightarrow\;Kd(Vc_{\mathrm{PD}})", 6.45, 3.05, 4.7, 0.85, 31)
    deck.add_box(slide, "At runtime, the current Vc selects the coefficient value.", 2.55, 4.70, 8.2, 0.78, bold=True, size=19)
    deck.add_takeaway(slide, "Time is used during preparation; Vc is used during simulation.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "9. Runtime example: stable low, then a rising edge",
        section="Runtime",
        notes=(
            "The handwritten runtime example begins from stable low. Vc_PU is zero and Vc_PD is one. "
            "On a rising input edge, the pullup target becomes one and the pulldown target becomes zero. "
            "The two rising-direction tau values govern the capacitor voltages."
        ),
    )
    deck.add_box(slide, "Stable low\nVc_PU=0\nVc_PD=1", 0.70, 1.25, 2.25, 1.20, bold=True, size=18)
    deck.add_arrow(slide, 2.95, 1.85, 3.85, 1.85)
    deck.add_box(slide, "Rising edge\nVtarget,PU=1\nVtarget,PD=0", 3.85, 1.25, 2.75, 1.20, bold=True, size=18)
    deck.add_arrow(slide, 6.60, 1.85, 7.50, 1.85)
    deck.add_box(slide, "Use\nTau_PU,rising\nTau_PD,rising", 7.50, 1.25, 2.75, 1.20, bold=True, size=18)
    eq(slide, r"Vc_{\mathrm{PU}}(t)=1-e^{-t/\tau_{\mathrm{PU,rising}}}", 0.75, 3.25, 5.65, 0.85, 28)
    eq(slide, r"Vc_{\mathrm{PD}}(t)=e^{-t/\tau_{\mathrm{PD,rising}}}", 6.65, 3.25, 5.65, 0.85, 28)
    eq(slide, r"Ku=Ku_{\mathrm{rising}}(Vc_{\mathrm{PU}}),\quad Kd=Kd_{\mathrm{rising}}(Vc_{\mathrm{PD}})", 1.55, 4.75, 10.1, 0.82, 27)
    deck.add_takeaway(slide, "The rising edge moves both capacitor voltages using their rising-direction relationships.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "10. A falling edge arrives before the rising transition finishes",
        section="Runtime",
        notes=(
            "At the falling edge, the targets flip and the model changes from rising-direction tau values "
            "to falling-direction tau values. Vc_PU and Vc_PD do not jump; they continue from the values "
            "already reached. The falling Ku/Kd versus Vc relationships are then used."
        ),
    )
    deck.add_picture_contain(slide, conceptual_reversal_fig, 0.65, 1.10, 7.20, 4.85, border=False)
    deck.add_box(slide, "Targets flip\nPU: 1 -> 0\nPD: 0 -> 1", 8.20, 1.35, 3.80, 1.05, bold=True, size=18)
    deck.add_box(slide, "Tau changes\nrising -> falling", 8.20, 2.80, 3.80, 0.85, bold=True, size=18)
    deck.add_box(slide, "Vc_PU and Vc_PD\ndo not jump", 8.20, 4.05, 3.80, 0.85, bold=True, size=18)
    deck.add_takeaway(slide, "The falling relationship starts from the present capacitor voltage, not from a settled endpoint.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "11. Cached io_buf example: actual Vc_PU and Vc_PD",
        section="Waveforms",
        notes=(
            "This slide uses cached ngspice data from the short_pulse_1ns_high case. The internal numerical "
            "capacitor voltages are relabeled in the handwritten notation as Vc_PU and Vc_PD. Both are "
            "continuous. Their measured turning points are shown without introducing another timing model."
        ),
    )
    deck.add_picture_contain(slide, real_vc_fig, 0.75, 1.05, 11.9, 5.15, border=False)
    deck.add_takeaway(slide, "The measured capacitor voltages carry different transition histories, but neither is reset.")
    deck.add_source(slide, "Cached data: io_buf short_pulse_1ns_high / two_state_directional_residual")

    slide = deck.add_slide(
        "12. Cached io_buf example: Ku(Vc_PU) and Kd(Vc_PD)",
        section="Waveforms",
        notes=(
            "This is the coefficient result of using the capacitor voltages as lookup coordinates. "
            "The capacitor voltages are continuous, but changing between the rising and falling "
            "relationships still produces sharp coefficient changes in the present implementation. "
            "The measured adjacent-sample changes are annotated. This directly motivates the first "
            "next-step question in the handwritten notes."
        ),
    )
    deck.add_picture_contain(slide, real_k_vc_fig, 0.75, 1.05, 11.9, 5.15, border=False)
    deck.add_takeaway(slide, "Continuous Vc fixes the starting coordinate; the rising/falling Ku/Kd relationships still need compatible values.")
    deck.add_source(slide, "Underlying samples: results/ibis_kukd_handwritten_notes_deck/strict_vc_real_example_waveforms.csv")

    slide = deck.add_slide(
        "13. Ku(Vc_PU) and Kd(Vc_PD) still control the IBIS V-I tables",
        section="IBIS",
        notes=(
            "The final handwritten statement is that Ku and Kd retain their normal role in the IBIS "
            "current calculation. The difference is only how their values are selected: from Vc_PU and "
            "Vc_PD rather than directly from elapsed time."
        ),
    )
    deck.add_box(slide, "Input history", 0.65, 1.35, 2.0, 0.78, bold=True)
    deck.add_arrow(slide, 2.65, 1.74, 3.50, 1.74)
    deck.add_box(slide, "Vc_PU / Vc_PD", 3.50, 1.35, 2.25, 0.78, bold=True)
    deck.add_arrow(slide, 5.75, 1.74, 6.60, 1.74)
    deck.add_box(slide, "Ku(Vc_PU) / Kd(Vc_PD)", 6.60, 1.35, 2.85, 0.78, bold=True)
    deck.add_arrow(slide, 9.45, 1.74, 10.30, 1.74)
    deck.add_box(slide, "IBIS V-I current", 10.30, 1.35, 2.25, 0.78, bold=True)
    eq(slide, r"I_{\mathrm{PU}}(t)=Ku(Vc_{\mathrm{PU}})\,I_{\mathrm{PU,full}}(V)", 0.75, 3.05, 5.8, 0.85, 27)
    eq(slide, r"I_{\mathrm{PD}}(t)=Kd(Vc_{\mathrm{PD}})\,I_{\mathrm{PD,full}}(V)", 6.70, 3.05, 5.8, 0.85, 27)
    deck.add_box(slide, "The original pullup and pulldown V-I tables are unchanged.", 2.2, 4.70, 8.9, 0.78, bold=True, size=19)
    deck.add_takeaway(slide, "Only the Ku/Kd lookup coordinate changes; the IBIS current structure remains the same.")
    deck.add_source(slide, source)

    slide = deck.add_slide(
        "14. SPICE-equivalent realization using the handwritten Vc notation",
        section="Implementation",
        notes=(
            "This slide expresses the capacitor-voltage equations in SPICE syntax while retaining the "
            "handwritten Vc names. Each capacitor stores one voltage. Each behavioral current source "
            "makes that voltage approach its target according to the selected tau."
        ),
    )
    deck.add_code_box(
        slide,
        "* Pullup capacitor voltage\n"
        "CVCPU  VCPU  0  1p  ic=0\n"
        "BVCPU  VCPU  0  I=-1p*(V(VCPUTARGET)-V(VCPU))/tau_selected\n\n"
        "* Pulldown capacitor voltage\n"
        "CVCPD  VCPD  0  1p  ic=1\n"
        "BVCPD  VCPD  0  I=-1p*(V(VCPDTARGET)-V(VCPD))/tau_selected",
        0.65,
        1.25,
        7.25,
        3.3,
        size=12.3,
    )
    eq(slide, r"\frac{dVc}{dt}=\frac{Vc_{\mathrm{target}}-Vc}{\tau}", 8.05, 1.55, 4.45, 0.85, 30)
    deck.add_box(slide, "CVCPU and CVCPD store the two Vc values.", 8.05, 3.05, 4.45, 0.82, bold=True, size=17)
    deck.add_box(slide, "The ngspice solver integrates Vc at every timestep.", 8.05, 4.25, 4.45, 0.82, bold=True, size=17)
    deck.add_takeaway(slide, "Vc is stored as capacitor voltage and carried continuously through a change in direction.")
    deck.add_source(slide, "Notation-matched SPICE realization of the handwritten equations")

    slide = deck.add_slide(
        "15. Next steps recorded in the handwritten notes",
        section="Next Steps",
        notes=(
            "The notes end with five next steps. First, test whether Vc should only locate the correct "
            "starting point for Ku/Kd when an edge arrives in the middle. Then investigate the unusual "
            "HSPICE result, check termination capacitance from about 1 fF to 10 pF with driver-strength "
            "dependence, create a SPICE buffer using ADS, and talk to Jianchuan about the inverter model."
        ),
    )
    deck.add_box(slide, "Main model question", 0.75, 1.20, 3.0, 0.65, bold=True, size=18)
    deck.add_text(
        slide,
        "Only use Vc to track the Ku/Kd starting point when an edge arrives mid-transition?",
        0.90,
        2.00,
        5.1,
        1.25,
        size=21,
        bold=True,
        color=deck.theme.green,
    )
    deck.add_box(slide, "Reference and circuit checks", 6.65, 1.20, 3.3, 0.65, bold=True, size=18)
    deck.add_bullets(
        slide,
        [
            "Investigate why the HSPICE result looks unusual.",
            "Check termination capacitance from about 1 fF to 10 pF; account for driver strength.",
            "Use ADS to create a SPICE buffer.",
            "Talk to Jianchuan about the inverter model.",
        ],
        6.65,
        2.00,
        5.75,
        3.0,
        size=17,
        spacing=7,
    )
    deck.add_box(slide, "These are open questions from the notes.", 2.70, 5.30, 7.9, 0.72, fill=deck.theme.light_red, line=deck.theme.red, bold=True, size=18)
    deck.add_takeaway(slide, "The next experiment should test Vc as a starting-point coordinate before expanding its role.")
    deck.add_source(slide, "Transcribed from handwritten source, page 3")

    slide = deck.add_slide(
        "Appendix: original handwritten source",
        section="Source",
        notes="These are the three handwritten pages used as the vocabulary and sequence for this strict version.",
    )
    for idx, page in enumerate(appendix_pages, start=1):
        if page.exists():
            deck.add_picture_contain(slide, page, 0.55 + (idx - 1) * 4.15, 1.15, 3.75, 5.35, border=True)
            deck.add_text(slide, f"Page {idx}", 1.65 + (idx - 1) * 4.15, 6.25, 1.6, 0.25, size=11, bold=True, align=PP_ALIGN.CENTER)
    deck.add_source(slide, r"Original PDF: \\minerfiles.mst.edu\dfs\users\sh3qm\Downloads\IBIS kukd .pdf")

    path = deck.save(OUTPUT)
    with (OUT / "strict_vc_equation_manifest.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["slide", "tex", "rendered_asset"])
        writer.writerows(equations)
    return path


if __name__ == "__main__":
    result = build()
    print(result)
