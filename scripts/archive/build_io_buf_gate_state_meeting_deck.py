from __future__ import annotations

import argparse
import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

import build_io_buf_two_state_gate_presentation as base


ROOT = Path(__file__).resolve().parents[2]
RESULT_ROOT = ROOT / "results" / "io_buf_two_state_gate_model_2026-06-30"
DEFAULT_TEMPLATE = Path(r"\\minerfiles.mst.edu\dfs\users\sh3qm\Downloads\0710_Simon_IBIS.pptx")
DEFAULT_OUTPUT = RESULT_ROOT / "presentation" / "0717_gate_state_short_pulse_meeting_deck.pptx"

EVIDENCE_DIR = RESULT_ROOT / "presentation_evidence_figures"
NATIVE_PLOT_DIR = RESULT_ROOT / "waveform_evidence_report" / "plots" / "directional_residual_vs_native_ibis"

NAVY = RGBColor(27, 72, 111)
LIGHT_BLUE = RGBColor(234, 244, 255)
PALE_ORANGE = RGBColor(255, 241, 228)
PURPLE = RGBColor(105, 73, 150)
LIGHT_PURPLE = RGBColor(241, 236, 248)


def remove_template_content_slides(prs: Presentation) -> None:
    """Keep the template title slide and remove its example content slides."""
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids)[1:]:
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def set_title_slide(prs: Presentation) -> None:
    slide = prs.slides[0]
    title = slide.shapes.title
    title.text = "Continuous Gate-State Model for Short-Pulse IBIS"
    paragraph = title.text_frame.paragraphs[0]
    paragraph.font.name = base.TITLE_FONT
    paragraph.font.size = Pt(35)
    paragraph.font.bold = True
    paragraph.font.color.rgb = base.BLACK

    subtitle = slide.placeholders[1]
    subtitle.text = "Simon Hwang\nAdvisor: Dr. Chulsoon Hwang\nDr. Zhiping Yang\n7/17/2026"
    for paragraph in subtitle.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.name = base.TITLE_FONT
        paragraph.font.size = Pt(16)
        paragraph.font.color.rgb = base.BLACK

    base.add_notes(
        slide,
        """
This presentation explains one idea: replace restart-based Ku/Kd playback with two continuous hidden gate states. The purpose is to handle a reverse input edge that arrives before the prior output transition has settled.

The deck starts with the failure mechanism in legacy pybis, then defines GUP and GDN, shows exactly how delay and tau are extracted from the four complete-edge coefficient traces, and walks through the generated ngspice implementation. The final slides show the measured evidence and the remaining limitation.

The method presented is two_state_directional_residual. It is the strongest structural method in the study, but it is still experimental. It passes the offline complete-edge reconstruction gate and produces a coefficient-correct short-low result. It does not yet preserve the normal path as well as legacy, and short-high Kd recovery remains open.
""",
    )


def add_section_tag(slide, text: str) -> None:
    base.add_text(slide, text.upper(), 10.45, 0.12, 2.15, 0.25, size=9.5, color=base.GRAY, bold=True, align=PP_ALIGN.RIGHT)


def slide_problem(prs: Presentation) -> None:
    slide = base.add_slide(prs, "1. The short-pulse problem: a second edge arrives too early")
    add_section_tag(slide, "Problem")

    base.add_text(slide, "Long pulse: the first transition finishes", 0.62, 1.0, 5.75, 0.35, size=19, bold=True, color=base.GREEN, align=PP_ALIGN.CENTER)
    base.add_text(slide, "Short pulse: the reverse edge interrupts it", 6.95, 1.0, 5.75, 0.35, size=19, bold=True, color=base.RED, align=PP_ALIGN.CENTER)

    # Long pulse timeline.
    base.add_arrow(slide, 0.9, 2.15, 6.0, 2.15, color=base.GRAY, width=1.8)
    base.add_text(slide, "time", 5.7, 2.28, 0.5, 0.25, size=10, color=base.GRAY)
    base.add_box(slide, "rise", 1.15, 1.8, 0.8, 0.55, fill=LIGHT_BLUE, line=base.BLUE, size=14, bold=True)
    base.add_arrow(slide, 1.95, 2.08, 3.05, 2.08, color=base.BLUE)
    base.add_box(slide, "Ku -> 1\nKd -> 0", 3.05, 1.63, 1.35, 0.9, fill=base.PALE_GREEN, line=base.GREEN, size=14, bold=True)
    base.add_arrow(slide, 4.4, 2.08, 5.15, 2.08, color=base.GREEN)
    base.add_box(slide, "settled high", 5.15, 1.8, 1.0, 0.55, fill=base.LIGHT_GRAY, line=base.GRAY, size=13, bold=True)

    # Short pulse timeline.
    base.add_arrow(slide, 7.2, 2.15, 12.3, 2.15, color=base.GRAY, width=1.8)
    base.add_box(slide, "rise", 7.45, 1.8, 0.8, 0.55, fill=LIGHT_BLUE, line=base.BLUE, size=14, bold=True)
    base.add_arrow(slide, 8.25, 2.08, 9.45, 2.08, color=base.BLUE)
    base.add_box(slide, "partial\nKu/Kd", 9.45, 1.63, 1.15, 0.9, fill=PALE_ORANGE, line=base.ORANGE, size=14, bold=True)
    base.add_arrow(slide, 10.6, 2.08, 11.05, 2.08, color=base.RED)
    base.add_box(slide, "fall", 11.05, 1.8, 0.8, 0.55, fill=base.LIGHT_RED, line=base.RED, size=14, bold=True)
    base.add_text(slide, "Previous state is unfinished", 9.15, 2.72, 3.1, 0.35, size=16, color=base.RED, bold=True, align=PP_ALIGN.CENTER)

    base.add_text(slide, "Legacy replay assumption", 0.75, 3.45, 3.0, 0.35, size=19, bold=True, color=base.RED)
    base.add_box(slide, "Every detected edge starts a complete Ku/Kd table at table time = 0.", 0.78, 3.92, 4.85, 1.05, fill=base.LIGHT_RED, line=base.RED, size=17, bold=True)
    base.add_text(slide, "What a real buffer needs", 6.55, 3.45, 3.0, 0.35, size=19, bold=True, color=base.GREEN)
    base.add_box(slide, "The new command must continue from the partially charged internal drive state.", 6.58, 3.92, 5.35, 1.05, fill=base.LIGHT_GREEN, line=base.GREEN, size=17, bold=True)

    base.add_box(slide, "Observable consequence", 0.78, 5.42, 2.0, 0.5, fill=base.LIGHT_GRAY, line=base.GRAY, bold=True)
    base.add_text(slide, "Legacy pybis can produce a nearly full Ku/Kd response from a pulse too short to create a full transistor-level output transition.", 3.05, 5.35, 9.0, 0.7, size=17, color=base.DARK, bold=True)
    base.add_takeaway(slide, "The problem is not the 1 ps edge rate; it is loss of unfinished switching history at the reverse edge.")
    base.add_source(slide, "Study stimulus: 3.3 V, 1 ps edges, 50 ohm || 2 pF; short-pulse cases reverse before coefficient settling.")
    base.add_notes(
        slide,
        """
Begin with the physical event. A normal input pulse remains high long enough for the pullup strength Ku to reach its high endpoint and for the pulldown strength Kd to reach its low endpoint. Once the output is settled, starting the falling transition from the normal falling table is reasonable.

The short-pulse case is different. The input rises and then falls before the internal drive strengths finish changing. At the reverse edge, Ku and Kd contain useful history: they tell us how far the pullup and pulldown processes have progressed. Legacy pybis does not preserve that history in the playback coordinate. It detects the new edge and begins the opposite complete-edge table at its first sample, which assumes the previous logic state had settled.

The desired behavior is not to make the input edge slower. The edge remains 1 ps throughout this phase. The desired behavior is to retain the partial internal state and reverse from there. This is analogous to a transistor predriver whose internal gate nodes have only partially charged or discharged when the command reverses.
""",
    )


def slide_legacy_ku_kd(prs: Presentation) -> None:
    slide = base.add_slide(prs, "2. Original Ku/Kd: effective current strengths played from tables")
    add_section_tag(slide, "Legacy model")

    base.add_box(slide, "Ku(t)\npullup strength", 0.62, 1.2, 2.15, 0.88, fill=LIGHT_BLUE, line=base.BLUE, size=17, bold=True)
    base.add_box(slide, "Kd(t)\npulldown strength", 0.62, 2.35, 2.15, 0.88, fill=PALE_ORANGE, line=base.ORANGE, size=17, bold=True)
    base.add_arrow(slide, 2.77, 1.64, 3.65, 1.64, color=base.BLUE)
    base.add_arrow(slide, 2.77, 2.79, 3.65, 2.79, color=base.ORANGE)
    base.add_box(slide, "Ku x pullup I-V", 3.65, 1.22, 2.6, 0.82, fill=base.LIGHT_GRAY, line=base.GRAY, size=16, bold=True)
    base.add_box(slide, "Kd x pulldown I-V", 3.65, 2.37, 2.6, 0.82, fill=base.LIGHT_GRAY, line=base.GRAY, size=16, bold=True)
    base.add_arrow(slide, 6.25, 2.22, 7.15, 2.22, color=base.GREEN)
    base.add_box(slide, "PAD current\nand voltage", 7.15, 1.72, 2.2, 1.0, fill=base.PALE_GREEN, line=base.GREEN, size=18, bold=True)

    base.add_text(slide, "Meaning", 9.85, 1.1, 2.2, 0.35, size=18, bold=True, color=base.GREEN)
    base.add_bullets(
        slide,
        [
            "Near 1: network fully enabled",
            "Near 0: network disabled",
            "Small overshoot is allowed",
            "Ku/Kd are model coefficients, not transistor gate voltages",
        ],
        9.82,
        1.48,
        2.7,
        1.9,
        size=14.5,
        spacing=4,
    )

    base.add_text(slide, "Legacy event-driven replay", 0.72, 3.75, 3.5, 0.35, size=19, bold=True, color=base.RED)
    base.add_code_box(
        slide,
        "* simplified legacy pattern\n"
        "elapsed = time - last_edge_time\n"
        "if rising:  Ku,Kd = rising_tables(elapsed)\n"
        "if falling: Ku,Kd = falling_tables(elapsed)",
        0.75,
        4.15,
        5.2,
        1.45,
        size=13.5,
    )
    base.add_text(slide, "At a reverse edge", 6.55, 3.75, 2.6, 0.35, size=19, bold=True, color=base.RED)
    base.add_box(slide, "elapsed resets to 0", 6.62, 4.25, 2.25, 0.65, fill=base.LIGHT_RED, line=base.RED, size=16, bold=True)
    base.add_arrow(slide, 8.87, 4.58, 9.55, 4.58, color=base.RED)
    base.add_box(slide, "opposite table starts from its settled endpoint", 9.55, 4.08, 2.65, 1.0, fill=base.LIGHT_RED, line=base.RED, size=15, bold=True)

    base.add_text(slide, "Actual current scaling in the generated subcircuit:", 6.62, 5.35, 5.25, 0.28, size=13, color=base.GRAY, bold=True)
    base.add_code_box(slide, "B3 DIE PULLUP_REF   I={V(Ku)*pwl(...)}\nB4 DIE PULLDOWN_REF I={V(Kd)*pwl(...)}", 6.62, 5.65, 5.55, 0.72, size=11.5)
    base.add_takeaway(slide, "Ku/Kd are the final drive multipliers. The weakness is how their time history is generated, not how they scale the IBIS I-V tables.")
    base.add_source(slide, "Exact current-source pattern: generated driver_OutputInput_Typical.sub. Legacy replay shown in simplified form.")
    base.add_notes(
        slide,
        """
Ku and Kd are effective switching coefficients. Ku multiplies the pullup I-V table and Kd multiplies the pulldown I-V table. They are convenient model variables: around one means the corresponding network is fully enabled, around zero means disabled. The original coefficient extraction may contain small overshoot or undershoot, so these are not hard-clipped Boolean states.

It is important not to confuse Ku and Kd with physical transistor gate voltages. They are output-strength multipliers derived from IBIS waveform data. The generated subcircuit uses V(Ku) and V(Kd) directly in behavioral current sources B3 and B4.

Legacy InputDriven logic uses the most recent input edge to choose a rising or falling coefficient table and evaluates that table using elapsed time from the edge. When a reverse edge arrives, the elapsed coordinate restarts. That works after a settled transition because the opposite table's first sample represents the correct endpoint. It fails during an interrupted transition because the table's first sample does not equal the actual partial state.
""",
    )


def slide_core_model(prs: Presentation) -> None:
    slide = base.add_slide(prs, "3. Core innovation: store two continuous hidden gate states")
    add_section_tag(slide, "Proposed solution")

    base.add_text(slide, "Do not replay Ku/Kd directly", 0.62, 1.0, 4.2, 0.38, size=20, bold=True, color=base.RED, align=PP_ALIGN.CENTER)
    base.add_box(slide, "edge", 0.78, 1.75, 1.0, 0.62, fill=base.LIGHT_GRAY, line=base.GRAY, bold=True)
    base.add_arrow(slide, 1.78, 2.06, 2.55, 2.06, color=base.RED)
    base.add_box(slide, "start Ku/Kd table\nat t = 0", 2.55, 1.55, 2.2, 1.0, fill=base.LIGHT_RED, line=base.RED, size=16, bold=True)
    base.add_text(slide, "history is discarded", 1.1, 2.8, 3.5, 0.35, size=16, color=base.RED, bold=True, align=PP_ALIGN.CENTER)

    base.add_text(slide, "Store internal progress, then derive Ku/Kd", 5.5, 1.0, 6.95, 0.38, size=20, bold=True, color=base.GREEN, align=PP_ALIGN.CENTER)
    stages = [
        ("edge", 5.65, 1.0, base.LIGHT_GRAY, base.GRAY),
        ("delayed\ntargets", 7.0, 1.35, base.LIGHT_GREEN, base.GREEN),
        ("GUP / GDN\nstored states", 8.75, 1.65, LIGHT_PURPLE, PURPLE),
        ("directional\nPWL maps", 10.8, 1.55, base.LIGHT_GREEN, base.GREEN),
    ]
    for idx, (label, x, w, fill, line) in enumerate(stages):
        base.add_box(slide, label, x, 1.55, w, 1.0, fill=fill, line=line, size=15, bold=True)
        if idx < len(stages) - 1:
            base.add_arrow(slide, x + w, 2.05, stages[idx + 1][1], 2.05, color=base.GREEN)
    base.add_box(slide, "Ku / Kd", 11.0, 3.0, 1.35, 0.62, fill=LIGHT_BLUE, line=base.BLUE, size=16, bold=True)
    base.add_arrow(slide, 11.58, 2.55, 11.58, 3.0, color=base.GREEN)

    base.add_text(slide, "What the states mean", 0.72, 3.65, 3.2, 0.38, size=19, bold=True, color=base.GREEN)
    base.add_box(slide, "GUP\npartial pullup gate-drive progress\n0 = low-state endpoint\n1 = high-state endpoint", 0.78, 4.1, 3.1, 1.45, fill=LIGHT_BLUE, line=base.BLUE, size=15.5, bold=True)
    base.add_box(slide, "GDN\npartial pulldown gate-drive progress\n1 = low-state endpoint\n0 = high-state endpoint", 4.18, 4.1, 3.1, 1.45, fill=PALE_ORANGE, line=base.ORANGE, size=15.5, bold=True)
    base.add_box(slide, "Ku/Kd\neffective output-network strengths\ncomputed from GUP/GDN\nthrough fitted maps", 7.58, 4.1, 3.1, 1.45, fill=base.PALE_GREEN, line=base.GREEN, size=15.5, bold=True)
    base.add_text(slide, "G is memory; K is output strength", 10.95, 4.42, 1.6, 0.8, size=15, color=PURPLE, bold=True, align=PP_ALIGN.CENTER)

    base.add_takeaway(slide, "At a reverse edge, only the targets change. GUP and GDN continue from their current capacitor voltages.")
    base.add_source(slide, "Primary experimental mode: InputDrivenTwoStateGateDirectionalResidualFull.")
    base.add_notes(
        slide,
        """
The core change is a level of indirection. Instead of treating Ku and Kd as the states that are replayed from complete-edge tables, the model introduces two hidden continuous variables. GUP represents pullup predriver or gate-drive progress. GDN represents pulldown progress. These are normalized internal coordinates, not literal measured MOS gate voltages.

The input edge first creates delayed command targets. GUP and GDN then move continuously toward those targets. Direction-specific maps convert the hidden states to the final Ku and Kd coefficients. The maps are necessary because the relationship between hidden progress and effective output strength is not the same during turn-on and turn-off.

This answers the short-pulse problem directly. Suppose a falling edge arrives while GUP is 0.31 and GDN is 0.42. The new edge changes the desired target directions, but the stored states remain 0.31 and 0.42 at that instant. Their derivatives change; their values do not jump. Ku and Kd are recomputed from that continuous history.
""",
    )


def slide_gate_derivation(prs: Presentation) -> None:
    slide = base.add_slide(prs, "4. Derive GUP/GDN from the four normal Ku/Kd transitions")
    add_section_tag(slide, "Model derivation")

    base.add_text(slide, "Complete rising input", 0.6, 0.98, 5.9, 0.35, size=19, bold=True, color=base.GREEN, align=PP_ALIGN.CENTER)
    base.add_text(slide, "Complete falling input", 6.82, 0.98, 5.9, 0.35, size=19, bold=True, color=base.GREEN, align=PP_ALIGN.CENTER)

    cards = [
        ("Ku rise", "pullup ON", 0.72, 1.55, LIGHT_BLUE, base.BLUE, "GUP: 0 -> 1"),
        ("Kd fall", "pulldown OFF", 3.72, 1.55, PALE_ORANGE, base.ORANGE, "GDN: 1 -> 0"),
        ("Ku fall", "pullup OFF", 6.95, 1.55, LIGHT_BLUE, base.BLUE, "GUP: 1 -> 0"),
        ("Kd rise", "pulldown ON", 9.95, 1.55, PALE_ORANGE, base.ORANGE, "GDN: 0 -> 1"),
    ]
    for title, process, x, y, fill, line, state in cards:
        base.add_box(slide, f"{title}\n{process}\n{state}", x, y, 2.65, 1.15, fill=fill, line=line, size=15, bold=True)

    base.add_text(slide, "1. Normalize each transition", 0.72, 3.12, 4.8, 0.35, size=18, bold=True, color=base.GREEN)
    base.add_code_box(slide, "progress(t) = (K(t) - K_start) / (K_end - K_start)", 0.75, 3.55, 5.6, 0.7, size=13)
    base.add_text(slide, "2. Fit a delayed hidden state", 6.7, 3.12, 5.0, 0.35, size=18, bold=True, color=base.GREEN)
    base.add_code_box(slide, "G(t) = G_start + (G_end-G_start)\n       * [1-exp(-(t-delay)/tau)]", 6.72, 3.55, 5.55, 0.92, size=12.3)

    base.add_text(slide, "3. Pair reconstructed state with the original coefficient", 0.72, 4.75, 5.6, 0.35, size=18, bold=True, color=base.GREEN)
    base.add_box(slide, "(GUP, Ku) samples -> Ku_on(GUP) and Ku_off(GUP)", 0.75, 5.15, 5.6, 0.62, fill=LIGHT_BLUE, line=base.BLUE, size=15, bold=True)
    base.add_box(slide, "(GDN, Kd) samples -> Kd_off(GDN) and Kd_on(GDN)", 6.72, 5.15, 5.55, 0.62, fill=PALE_ORANGE, line=base.ORANGE, size=15, bold=True)

    base.add_text(slide, "Why four paths?", 0.75, 6.05, 1.6, 0.28, size=14, color=base.RED, bold=True)
    base.add_text(slide, "Pullup/pulldown and on/off transitions have different onset and shape; one shared state law cannot reproduce all four.", 2.25, 6.0, 9.9, 0.45, size=15.5, color=base.DARK, bold=True)
    base.add_takeaway(slide, "The model is trained only from the original IBIS-derived complete-edge Ku/Kd tables, never from HSPICE waveforms.")
    base.add_source(slide, "Endpoints and progress are inferred from io_buf rising/falling coefficient tables.")
    base.add_notes(
        slide,
        """
There are four training traces. During a complete rising input transition, Ku turns on and Kd turns off. During a complete falling transition, Ku turns off and Kd turns on. These four paths are fitted separately because they have different onset delays, time constants, and coefficient shapes.

For each trace, the code infers the start and end coefficient values and normalizes the transition to a zero-to-one progress coordinate. It then fits a delayed first-order state trajectory. For GUP, the rising Ku path moves GUP from zero to one and falling Ku moves it from one to zero. For GDN, the rising-input Kd-off path moves GDN from one to zero, while the falling-input Kd-on path moves it from zero to one.

After reconstructing the state trajectory, the code pairs each hidden-state sample with the original coefficient sample. Those pairs form four PWL maps: Ku-on, Ku-off, Kd-off, and Kd-on. This preserves the original nonlinear coefficient shape while giving the runtime simulator a continuous hidden coordinate that survives a reverse edge.
""",
    )


def slide_tau_delay(prs: Presentation) -> None:
    slide = base.add_slide(prs, "5. Delay and tau: when motion begins, and how fast it proceeds")
    add_section_tag(slide, "Parameter extraction")

    base.add_text(slide, "Normalized progress of one complete transition", 0.68, 0.98, 5.7, 0.35, size=19, bold=True, color=base.GREEN)
    # Axes and illustrative exponential.
    base.add_arrow(slide, 0.95, 4.25, 6.05, 4.25, color=base.GRAY, width=1.6)
    base.add_arrow(slide, 0.95, 4.25, 0.95, 1.55, color=base.GRAY, width=1.6)
    curve = []
    for index in range(45):
        x = index / 44
        y = 0 if x < 0.18 else 1 - math.exp(-(x - 0.18) / 0.22)
        curve.append((x, min(y, 0.97)))
    base.add_curve(slide, curve, 1.05, 1.65, 4.85, 2.5, color=PURPLE, width=3)
    for frac, label, color in [(0.05, "t05", base.BLUE), (0.63, "t63", base.GREEN), (0.90, "t90", base.ORANGE)]:
        y = 4.15 - frac * 2.5
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(y), Inches(5.1), Inches(0.01))
        line.fill.solid()
        line.fill.fore_color.rgb = color
        line.line.color.rgb = color
        base.add_text(slide, f"{int(frac*100)}%", 0.3, y - 0.12, 0.55, 0.25, size=10, color=color, bold=True, align=PP_ALIGN.RIGHT)
        base.add_text(slide, label, 5.85, y - 0.13, 0.45, 0.25, size=10, color=color, bold=True)

    base.add_text(slide, "Extraction rule", 0.72, 4.72, 2.0, 0.35, size=18, bold=True, color=base.GREEN)
    base.add_code_box(slide, "delay = t05\ntau = max(20 ps, t63-delay, (t90-delay)/ln(10))", 0.75, 5.08, 5.6, 0.95, size=13)

    base.add_text(slide, "Measured io_buf parameters", 6.72, 0.98, 5.1, 0.35, size=19, bold=True, color=base.GREEN)
    rows = [
        ("Pullup ON", "1.302 ns", "1.284 ns", LIGHT_BLUE, base.BLUE),
        ("Pullup OFF", "0.440 ns", "0.363 ns", LIGHT_BLUE, base.BLUE),
        ("Pulldown OFF", "1.443 ns", "0.237 ns", PALE_ORANGE, base.ORANGE),
        ("Pulldown ON", "2.480 ns", "0.275 ns", PALE_ORANGE, base.ORANGE),
    ]
    base.add_text(slide, "Path", 6.9, 1.55, 2.1, 0.3, size=13, color=base.GRAY, bold=True)
    base.add_text(slide, "onset delay", 9.05, 1.55, 1.35, 0.3, size=13, color=base.GRAY, bold=True, align=PP_ALIGN.CENTER)
    base.add_text(slide, "tau", 10.78, 1.55, 1.2, 0.3, size=13, color=base.GRAY, bold=True, align=PP_ALIGN.CENTER)
    y = 1.92
    for label, delay, tau, fill, line in rows:
        base.add_box(slide, label, 6.85, y, 2.1, 0.62, fill=fill, line=line, size=14.5, bold=True)
        base.add_box(slide, delay, 9.05, y, 1.45, 0.62, fill=base.LIGHT_GRAY, line=base.MID_GRAY, size=14, bold=True)
        base.add_box(slide, tau, 10.68, y, 1.45, 0.62, fill=base.LIGHT_GRAY, line=base.MID_GRAY, size=14, bold=True)
        y += 0.8

    base.add_box(slide, "delay", 6.85, 5.35, 1.15, 0.55, fill=base.LIGHT_GRAY, line=base.GRAY, size=15, bold=True)
    base.add_text(slide, "transport time before a state responds", 8.15, 5.4, 3.85, 0.3, size=15)
    base.add_box(slide, "tau", 6.85, 6.0, 1.15, 0.55, fill=base.LIGHT_GRAY, line=base.GRAY, size=15, bold=True)
    base.add_text(slide, "exponential charge/discharge speed after onset", 8.15, 6.05, 4.2, 0.3, size=15)
    base.add_takeaway(slide, "Delay and tau are fitted independently for PU-on, PU-off, PD-off, and PD-on.")
    base.add_source(slide, "Actual selected tau uses t63-t05; the 90% term is a safeguard against an unrealistically fast tail.")
    base.add_notes(
        slide,
        """
Delay and tau describe different parts of the response. Delay is the onset time: the command has arrived at the input, but the corresponding hidden gate state has not yet started moving. Tau is the rate of exponential motion after that onset.

The implementation normalizes each complete-edge coefficient trace and measures its 5%, 63%, and 90% crossing times. The 5% crossing is used as the delay. A first-order system reaches approximately 63% one tau after it begins, so t63 minus delay estimates tau. The 90% crossing provides a tail safeguard through division by ln(10). A 20 ps floor prevents a numerically singular time constant.

The actual io_buf values are strongly asymmetric. Pullup turn-on begins at about 1.302 ns and is slow, with tau about 1.284 ns. Pullup turn-off begins much earlier and is faster. Pulldown turn-off begins around 1.443 ns, while pulldown turn-on has a long 2.480 ns onset delay. This asymmetry is exactly why one generic state or one generic replay clock is inadequate.
""",
    )


def slide_command_caps(prs: Presentation) -> None:
    slide = base.add_slide(prs, "6. Runtime part 1: delayed command memory and analog gate memory")
    add_section_tag(slide, "Ngspice implementation")

    base.add_text(slide, "A. Command capacitor: requested endpoint", 0.58, 1.02, 6.1, 0.32, size=17.5, bold=True, color=base.GREEN)
    base.add_code_box(
        slide,
        "TPUONP  RISEEDGE 0 PUONP  0 Z0=50 Td=1.302n\n"
        "TPUOFFP FALLEDGE 0 PUOFFP 0 Z0=50 Td=0.440n\n"
        "CGUPCMD GUPCMD 0 {gate_c} ic=0\n"
        "BGUPCMDON  GUPCMD 0 I=-{gate_c}*V(PUONP)/edge_delay\n"
        "BGUPCMDOFF GUPCMD 0 I= {gate_c}*V(PUOFFP)/edge_delay\n"
        "BGUPTARGET GUPTARGET 0 V=min(max(V(GUPCMD),0),1)",
        0.62,
        1.36,
        6.08,
        2.25,
        size=10.7,
    )
    base.add_box(slide, "CGUPCMD answers:\nWhere is the pullup commanded to go?", 0.82, 3.88, 5.65, 0.78, fill=base.LIGHT_GREEN, line=base.GREEN, size=16, bold=True)

    base.add_text(slide, "B. Gate capacitor: actual partial state", 6.85, 1.02, 5.9, 0.32, size=17.5, bold=True, color=PURPLE)
    base.add_code_box(
        slide,
        "BGUP GUP 0 I=-{gate_c}*(V(GUPTARGET)-V(GUP))\n"
        "+ / ((V(GUPTARGET)>V(GUP)) ? 1.284n : 0.363n)\n"
        "CGUP GUP 0 {gate_c} ic=0\n"
        "RGUP GUP 0 1e12\n\n"
        "dGUP/dt = (GUPTARGET-GUP)/selected_tau",
        6.9,
        1.36,
        5.75,
        2.25,
        size=10.8,
    )
    base.add_box(slide, "CGUP answers:\nHow far has the pullup actually moved?", 7.1, 3.88, 5.35, 0.78, fill=LIGHT_PURPLE, line=PURPLE, size=16, bold=True)

    base.add_text(slide, "Example: state at the reverse edge", 0.65, 5.05, 4.8, 0.35, size=18, bold=True, color=base.GREEN)
    base.add_box(slide, "before reverse\nGUPTARGET = 1\nGUP = 0.31", 0.72, 5.45, 2.25, 0.9, fill=LIGHT_BLUE, line=base.BLUE, size=15, bold=True)
    base.add_arrow(slide, 2.97, 5.9, 3.8, 5.9, color=base.RED)
    base.add_box(slide, "fall command", 3.8, 5.57, 1.75, 0.65, fill=base.LIGHT_RED, line=base.RED, size=15, bold=True)
    base.add_arrow(slide, 5.55, 5.9, 6.4, 5.9, color=base.GREEN)
    base.add_box(slide, "after reverse\nGUPTARGET = 0\nGUP = 0.31", 6.4, 5.45, 2.25, 0.9, fill=base.PALE_GREEN, line=base.GREEN, size=15, bold=True)
    base.add_text(slide, "Value stays continuous; only its slope changes.", 9.0, 5.62, 3.2, 0.55, size=16, color=PURPLE, bold=True, align=PP_ALIGN.CENTER)

    base.add_takeaway(slide, "Command caps remember the delayed target; gate caps remember partial analog progress. They are intentionally separate.")
    base.add_source(slide, "Exact structure from generated driver_OutputInput_Typical.sub; gate_c=1 pF is numerical state scaling, not a physical extracted capacitance.")
    base.add_notes(
        slide,
        """
This slide distinguishes the two kinds of capacitor, which was unclear in the prior deck.

The command capacitor, such as CGUPCMD, stores the persistent requested endpoint after delayed edge pulses arrive. A rise event is transported through TPUONP. Its short pulse injects charge into GUPCMD and moves the command toward one. A delayed fall pulse removes charge and moves the command toward zero. BGUPTARGET clamps the command to the valid zero-to-one range. The same structure exists for the pulldown command.

The gate-state capacitor CGUP is different. Its voltage V(GUP) is the actual continuous hidden state. BGUP supplies exactly the current needed to implement dGUP/dt=(target-GUP)/tau. Because current and capacitance obey I=C dV/dt, the chosen gate_c cancels algebraically. The 1 pF value is a numerical scaling convenience, not a claim that the transistor gate capacitance is 1 pF.

When the target reverses, GUPTARGET may change quickly, but V(GUP) cannot jump. The source selects the off tau and GUP decays from its current value. GDN uses the same structure with different on and off taus and an initial state of one.
""",
    )


def slide_mapping(prs: Presentation) -> None:
    slide = base.add_slide(prs, "7. Runtime part 2: map hidden state into effective Ku/Kd")
    add_section_tag(slide, "Ngspice implementation")

    base.add_text(slide, "GUP/GDN are progress coordinates; Ku/Kd are current-strength coefficients", 0.62, 0.96, 12.0, 0.38, size=19, bold=True, color=base.GREEN, align=PP_ALIGN.CENTER)

    base.add_box(slide, "GUP = 0.40", 0.75, 1.62, 1.55, 0.62, fill=LIGHT_PURPLE, line=PURPLE, size=16, bold=True)
    base.add_arrow(slide, 2.3, 1.93, 3.15, 1.93, color=base.BLUE)
    base.add_box(slide, "turning ON map\nKu about 0.23", 3.15, 1.42, 2.25, 1.0, fill=LIGHT_BLUE, line=base.BLUE, size=15, bold=True)
    base.add_arrow(slide, 2.3, 2.75, 3.15, 2.75, color=base.RED)
    base.add_box(slide, "turning OFF map\nKu about 0.43", 3.15, 2.25, 2.25, 1.0, fill=base.LIGHT_RED, line=base.RED, size=15, bold=True)

    base.add_box(slide, "GDN = 0.50", 6.85, 1.62, 1.55, 0.62, fill=LIGHT_PURPLE, line=PURPLE, size=16, bold=True)
    base.add_arrow(slide, 8.4, 1.93, 9.25, 1.93, color=base.ORANGE)
    base.add_box(slide, "turning OFF map\nKd about 0.56", 9.25, 1.42, 2.25, 1.0, fill=PALE_ORANGE, line=base.ORANGE, size=15, bold=True)
    base.add_arrow(slide, 8.4, 2.75, 9.25, 2.75, color=base.GREEN)
    base.add_box(slide, "turning ON map\nKd about 0.43", 9.25, 2.25, 2.25, 1.0, fill=base.PALE_GREEN, line=base.GREEN, size=15, bold=True)

    base.add_text(slide, "Direction selector and final coefficient nodes", 0.72, 3.62, 5.5, 0.35, size=18.5, bold=True, color=base.GREEN)
    base.add_code_box(
        slide,
        "BKUGATE_BASE ... V=(GUPTARGET>=GUP)\n"
        "+ ? V(KUGATE_ON) : V(KUGATE_OFF)\n"
        "BKDGATE_BASE ... V=(GDNTARGET>=GDN)\n"
        "+ ? V(KDGATE_ON) : V(KDGATE_OFF)\n\n"
        "B44 Ku 0 I=-C*(KUTARGET-Ku)/5p\n"
        "B45 Kd 0 I=-C*(KDTARGET-Kd)/5p",
        0.75,
        4.02,
        5.75,
        1.95,
        size=10.8,
    )

    base.add_text(slide, "Then the ordinary IBIS output equations stay unchanged", 6.85, 3.62, 5.5, 0.35, size=18.5, bold=True, color=base.GREEN)
    base.add_code_box(
        slide,
        "B3 DIE PULLUP_REF\n"
        "+ I={V(Ku)*pwl(V(DIE,VSS), ...)}\n\n"
        "B4 DIE PULLDOWN_REF\n"
        "+ I={V(Kd)*pwl(V(DIE,VSS), ...)}",
        6.9,
        4.02,
        5.45,
        1.42,
        size=11.5,
    )
    base.add_box(slide, "Small 5 ps coefficient capacitors smooth map-branch changes; they do not create the main switching history.", 6.9, 5.65, 5.45, 0.7, fill=base.LIGHT_GRAY, line=base.GRAY, size=14.5, bold=True)

    base.add_takeaway(slide, "The same G value can yield a different K value by direction, so one shared G-to-K map is not sufficient.")
    base.add_source(slide, "Illustrative values are read from the generated direction-specific PWL maps for io_buf.")
    base.add_notes(
        slide,
        """
GUP and GDN are not replacements with the same meaning as Ku and Kd. G is a hidden progress coordinate designed to evolve continuously. K is the effective output-network strength used by the IBIS current equations.

The relationship is direction dependent. In the fitted io_buf maps, a GUP value around 0.40 corresponds to Ku around 0.23 while turning on but around 0.43 while turning off. Similarly, GDN around 0.50 maps to different Kd values depending on whether the pulldown is turning off or on. This hysteresis-like directional dependence is empirical evidence that a single static map cannot represent both paths.

At runtime, the behavioral selector compares each target with its current state. If target is greater than state, it selects the on map; otherwise it selects the off map. The resulting KUGATE and KDGATE values feed small capacitor-backed Ku and Kd nodes with a 5 ps smoothing constant. Those final coefficients multiply the original IBIS pullup and pulldown I-V tables exactly as before.
""",
    )


def slide_full_syntax(prs: Presentation) -> None:
    slide = base.add_slide(prs, "8. Everything together: the generated ngspice signal path")
    add_section_tag(slide, "Implementation summary")

    stages = [
        ("1\ninput edge", 0.4, 1.35, base.LIGHT_GRAY, base.GRAY),
        ("2\ndelayed event", 2.0, 1.55, base.LIGHT_GREEN, base.GREEN),
        ("3\ncommand cap", 3.8, 1.55, base.LIGHT_GREEN, base.GREEN),
        ("4\ngate-state cap", 5.6, 1.65, LIGHT_PURPLE, PURPLE),
        ("5\ndirectional map", 7.5, 1.7, LIGHT_BLUE, base.BLUE),
        ("6\nKu/Kd", 9.45, 1.2, PALE_ORANGE, base.ORANGE),
        ("7\nIBIS I-V", 10.9, 1.5, base.LIGHT_GRAY, base.GRAY),
    ]
    for idx, (label, x, w, fill, line) in enumerate(stages):
        base.add_box(slide, label, x, 1.08, w, 0.85, fill=fill, line=line, size=13.5, bold=True)
        if idx < len(stages) - 1:
            base.add_arrow(slide, x + w, 1.51, stages[idx + 1][1], 1.51, color=base.GREEN, width=1.7)

    base.add_code_box(
        slide,
        "* 1-3: edge -> delayed command -> target\n"
        "TPUONP RISEEDGE 0 PUONP 0 Z0=50 Td=1.3019n\n"
        "CGUPCMD GUPCMD 0 1p ic=0\n"
        "BGUPCMDON GUPCMD 0 I=-1p*V(PUONP)/edge_delay\n"
        "BGUPTARGET GUPTARGET 0 V=min(max(V(GUPCMD),0),1)\n\n"
        "* 4: continuous hidden state\n"
        "BGUP GUP 0 I=-1p*(V(GUPTARGET)-V(GUP))/tau_selected\n"
        "CGUP GUP 0 1p ic=0",
        0.55,
        2.35,
        6.0,
        3.25,
        size=10.3,
    )
    base.add_code_box(
        slide,
        "* 5-6: state -> directional coefficient\n"
        "BKUGATE_BASE KUGATE_BASE 0 V=\n"
        "+ (GUPTARGET>=GUP) ? KUGATE_ON : KUGATE_OFF\n"
        "B44 Ku 0 I=-1p*(V(KUTARGET)-V(Ku))/5p\n"
        "Cku Ku 0 1p ic=0.00194\n\n"
        "* 7: coefficient -> IBIS current\n"
        "B3 DIE PULLUP_REF I={V(Ku)*pwl(...)}",
        6.78,
        2.35,
        5.95,
        3.25,
        size=10.3,
    )

    base.add_box(slide, "Parallel pulldown path", 0.62, 5.93, 2.2, 0.48, fill=PALE_ORANGE, line=base.ORANGE, size=14, bold=True)
    base.add_text(slide, "uses PDOFFP/PDONP, GDNCMD, GDN, Kd-on/off maps, and the PD-specific taus.", 3.0, 5.96, 8.7, 0.35, size=15.5, bold=True)
    base.add_takeaway(slide, "Every box corresponds to a real node or element in the generated .sub file; no external runtime controller is required.")
    base.add_source(slide, "Condensed exact syntax from cases/short_pulse_1ns_high/ngspice_two_state_directional_residual/driver_OutputInput_Typical.sub")
    base.add_notes(
        slide,
        """
This is the complete runtime chain in one slide. First, the input detector creates a rise or fall event pulse. Second, transmission-line elements implement transport delay for each physical path. Third, the delayed pulse charges or discharges a command capacitor, which stores the persistent desired endpoint. Fourth, a behavioral current source and gate-state capacitor implement the continuous first-order state equation. Fifth, direction-specific PWL maps convert the hidden state to an effective coefficient. Sixth, a small coefficient capacitor produces continuous final Ku or Kd. Seventh, that coefficient scales the existing IBIS I-V source.

The code shown is condensed but uses the actual generated ngspice element types and numerical values. The pulldown path is structurally parallel but uses separate delay, tau, state, and map values.

There is no Python process controlling the simulation after netlist generation. pybis2spice writes this structure into the subcircuit, and ngspice solves the event lines, capacitor states, behavioral sources, and output currents together in the transient analysis.
""",
    )


def slide_real_example(prs: Presentation) -> None:
    slide = base.add_slide(prs, "9. Real 1 ns short-high example: follow every event")
    add_section_tag(slide, "Worked example")

    plot = NATIVE_PLOT_DIR / "short_pulse_1ns_high_directional_residual_vs_native_ibis.png"
    base.add_picture_contain(slide, plot, 5.35, 1.0, 7.45, 5.55, border=True)

    base.add_text(slide, "Input and delayed internal events", 0.55, 1.0, 4.5, 0.35, size=18.5, bold=True, color=base.GREEN, align=PP_ALIGN.CENTER)
    events = [
        ("5.000 ns", "input rises", base.BLUE),
        ("6.000 ns", "input falls", base.RED),
        ("6.302 ns", "PU-on arrives", base.BLUE),
        ("6.440 ns", "PU-off arrives", base.RED),
        ("6.443 ns", "PD-off arrives", base.ORANGE),
        ("8.480 ns", "PD-on arrives", base.GREEN),
    ]
    y = 1.55
    for time_text, label, color in events:
        base.add_box(slide, time_text, 0.7, y, 1.25, 0.48, fill=base.LIGHT_GRAY, line=color, size=13, bold=True)
        base.add_text(slide, label, 2.08, y + 0.06, 2.65, 0.3, size=14.5, color=color, bold=True)
        y += 0.63

    base.add_box(slide, "GUP target is high for only about 0.138 ns", 0.72, 5.5, 4.15, 0.55, fill=LIGHT_BLUE, line=base.BLUE, size=15, bold=True)
    base.add_box(slide, "Result: Ku makes a small partial bump, not a full turn-on", 0.72, 6.08, 4.15, 0.55, fill=base.PALE_GREEN, line=base.GREEN, size=14.5, bold=True)

    base.add_takeaway(slide, "The gate-state model gets the partial Ku behavior; the long PD-on delay exposes the remaining short-high Kd recovery problem.")
    base.add_source(slide, "Cached waveform evidence; black = HSPICE native IBIS, red = two_state_directional_residual.")
    base.add_notes(
        slide,
        """
Walk through the event times. The input rises at 5.000 ns and falls at 6.000 ns, so the command pulse is only 1 ns wide. The four internal paths have different transport delays. Pullup-on arrives at about 6.302 ns, after the external input has already fallen. Pullup-off arrives at about 6.440 ns. Therefore GUPTARGET is high for only about 0.138 ns. GUP can charge only a little, and the direction-specific map produces a small Ku bump. This is the intended partial-response behavior.

Pulldown-off arrives at about 6.443 ns, while the delayed pulldown-on command does not arrive until about 8.480 ns because the complete-edge PD-on delay is 2.480 ns. This creates the central remaining limitation. Native HSPICE IBIS shows a different Kd hold and recovery sequence. The current two-state model preserves continuity, but its Kd recovery staging is too late.

The plot therefore contains both progress and an open issue. Ku and the tiny pad excursion are close, while Kd diverges. This is why the study requires coefficient-first validation rather than declaring success from pad voltage alone.
""",
    )


def slide_reconstruction(prs: Presentation) -> None:
    slide = base.add_slide(prs, "10. First evidence gate: can the model reproduce normal Ku/Kd?")
    add_section_tag(slide, "Validation")
    image = EVIDENCE_DIR / "reconstruction_gate_evidence.png"
    base.add_picture_contain(slide, image, 0.55, 0.92, 9.15, 5.55, border=True)

    base.add_text(slide, "Model evolution", 9.95, 1.0, 2.3, 0.35, size=19, bold=True, color=base.GREEN)
    base.add_metric_card(slide, "Single map", "max error 0.20467", "FAIL", 9.95, 1.5, 2.55, base.RED, value_size=13.5)
    base.add_metric_card(slide, "+ directional maps", "max error 0.07389", "FAIL", 9.95, 2.95, 2.55, base.RED, value_size=13.5)
    base.add_metric_card(slide, "+ Kd rate residual", "max error 0.04869", "PASS", 9.95, 4.4, 2.55, base.GREEN, value_size=13.5)

    base.add_takeaway(slide, "Only the directional-map plus residual model passes the predefined complete-edge reconstruction gate.")
    base.add_source(slide, "Offline cached reconstruction data; black = original IBIS-derived coefficient table, red = compact-model reconstruction.")
    base.add_notes(
        slide,
        """
Before making any short-pulse claim, the compact structure must reproduce the original complete-edge coefficient tables used to derive it. This is the reconstruction gate.

The left column uses a single map and has worst maximum error about 0.20467, so it fails. The middle column adds direction-specific on and off maps. This greatly improves shape and reduces worst maximum error to about 0.07389, but it still fails the predefined limit. The largest remaining defect is the negative Kd excursion, which a static map does not capture.

The right column adds a small Kd dynamic residual linked to state rate. Its worst table RMSE is 0.019939 and worst maximum error is 0.048685, so it passes. This is why two_state_directional_residual is called the best structural method so far. The residual is a refinement after the gate-state core, not the core idea itself.

These curves are offline table reconstructions, not transient HSPICE waveforms. HSPICE is not used to fit the parameters.
""",
    )


def slide_results(prs: Presentation) -> None:
    slide = base.add_slide(prs, "11. Measured result: one quadrant works; short-high Kd remains open")
    add_section_tag(slide, "Validation")
    image = EVIDENCE_DIR / "three_quadrants_one_open.png"
    base.add_picture_contain(slide, image, 0.52, 0.88, 9.75, 5.75, border=True)

    base.add_box(slide, "Short-low 2 ns", 10.48, 1.05, 2.05, 0.52, fill=base.LIGHT_GREEN, line=base.GREEN, size=16, bold=True)
    base.add_bullets(
        slide,
        [
            "Pad, Ku, and Kd improve together",
            "Pad RMSE: 23.271 mV",
            "Ku RMSE: 0.02141",
            "Kd RMSE: 0.01319",
            "Status: GOOD",
        ],
        10.42,
        1.7,
        2.25,
        1.65,
        size=13.2,
        spacing=2,
    )
    base.add_box(slide, "Short-high 1 ns", 10.48, 3.55, 2.05, 0.52, fill=base.LIGHT_RED, line=base.RED, size=16, bold=True)
    base.add_bullets(
        slide,
        [
            "Ku is partial and close",
            "Pad swing is tiny",
            "Kd recovery diverges",
            "Kd RMSE: 0.48754",
            "Status: OPEN",
        ],
        10.42,
        4.2,
        2.25,
        1.55,
        size=13.2,
        spacing=2,
    )
    base.add_box(slide, "Normal edge: WARN\n18.739 mV vs 5.289 mV legacy", 10.4, 5.72, 2.3, 0.72, fill=PALE_ORANGE, line=base.ORANGE, size=11.5, bold=True)

    base.add_takeaway(slide, "Success requires pad, Ku, and Kd together. A small pad error cannot hide a wrong coefficient trajectory.")
    base.add_source(slide, "Cached HSPICE/native-IBIS and ngspice directional-residual results; transistor reference appears only on pad panels.")
    base.add_notes(
        slide,
        """
The left column is the strongest success case: a 2 ns short-low pulse. The pad waveform and both coefficients improve together. This is right for the right reasons and demonstrates that continuous directional gate states can solve an interrupted-transition quadrant.

The right column is the 1 ns short-high case. Ku remains partial and close to native IBIS. The pad output is tiny, so its absolute error is also small. Kd, however, follows a clearly different hold and recovery trajectory, with RMSE about 0.48754. That case remains open.

The normal complete-edge control must also be stated. Directional-residual passes the offline reconstruction gate, but its long-pulse pad RMSE is 18.739 mV compared with 5.289 mV for legacy pybis. Therefore the model is not ready to replace legacy globally. A production version will likely need to preserve the legacy normal path while invoking continuous state handling only when interruption is detected and validated.
""",
    )


def slide_status(prs: Presentation) -> None:
    slide = base.add_slide(prs, "12. Current conclusion and focused next step")
    add_section_tag(slide, "Conclusion")

    base.add_text(slide, "What is proven", 0.72, 1.0, 3.0, 0.38, size=20, bold=True, color=base.GREEN)
    base.add_box(slide, "Continuous GUP/GDN states remove the forced table restart.", 0.78, 1.48, 3.8, 0.9, fill=base.LIGHT_GREEN, line=base.GREEN, size=16, bold=True)
    base.add_box(slide, "Four directional delay/tau paths are derived from IBIS tables only.", 0.78, 2.58, 3.8, 0.9, fill=base.LIGHT_GREEN, line=base.GREEN, size=16, bold=True)
    base.add_box(slide, "Directional maps plus Kd residual pass normal table reconstruction.", 0.78, 3.68, 3.8, 0.9, fill=base.LIGHT_GREEN, line=base.GREEN, size=16, bold=True)
    base.add_box(slide, "A short-low case is coefficient-correct, not just pad-correct.", 0.78, 4.78, 3.8, 0.9, fill=base.LIGHT_GREEN, line=base.GREEN, size=16, bold=True)

    base.add_text(slide, "What is not proven", 4.95, 1.0, 3.0, 0.38, size=20, bold=True, color=base.RED)
    base.add_box(slide, "Normal long-pulse behavior is still worse than legacy.", 5.0, 1.48, 3.3, 0.9, fill=base.LIGHT_RED, line=base.RED, size=16, bold=True)
    base.add_box(slide, "Short-high Kd hold/recovery is not represented correctly.", 5.0, 2.58, 3.3, 0.9, fill=base.LIGHT_RED, line=base.RED, size=16, bold=True)
    base.add_box(slide, "The method is validated only on io_buf so far.", 5.0, 3.68, 3.3, 0.9, fill=base.LIGHT_RED, line=base.RED, size=16, bold=True)
    base.add_box(slide, "It is experimental, not the pybis default.", 5.0, 4.78, 3.3, 0.9, fill=base.LIGHT_RED, line=base.RED, size=16, bold=True)

    base.add_text(slide, "Next candidate", 8.7, 1.0, 3.0, 0.38, size=20, bold=True, color=PURPLE)
    base.add_box(slide, "Keep", 8.78, 1.48, 1.0, 0.52, fill=LIGHT_PURPLE, line=PURPLE, size=15, bold=True)
    base.add_text(slide, "GUP/GDN, directional maps, reconstruction gate", 9.95, 1.47, 2.45, 0.7, size=14.5, bold=True)
    base.add_box(slide, "Add", 8.78, 2.45, 1.0, 0.52, fill=LIGHT_PURPLE, line=PURPLE, size=15, bold=True)
    base.add_text(slide, "one causal pending-command or recovery-phase state for Kd", 9.95, 2.4, 2.45, 0.85, size=14.5, bold=True)
    base.add_box(slide, "Protect", 8.78, 3.55, 1.0, 0.52, fill=LIGHT_PURPLE, line=PURPLE, size=15, bold=True)
    base.add_text(slide, "legacy-equivalent normal behavior outside interruption", 9.95, 3.5, 2.45, 0.85, size=14.5, bold=True)
    base.add_box(slide, "Validate", 8.78, 4.68, 1.0, 0.52, fill=LIGHT_PURPLE, line=PURPLE, size=15, bold=True)
    base.add_text(slide, "both pulse directions, widths, repeated toggles, loads, and other buffers", 9.95, 4.6, 2.45, 1.05, size=14.5, bold=True)

    base.add_box(slide, "Best structural method so far", 8.78, 5.95, 2.2, 0.52, fill=base.PALE_GREEN, line=base.GREEN, size=14.5, bold=True)
    base.add_text(slide, "two_state_directional_residual", 11.05, 5.96, 1.55, 0.45, size=13.5, color=base.GREEN, bold=True, align=PP_ALIGN.CENTER)

    base.add_takeaway(slide, "The core innovation is validated as a structure; the remaining work is narrow and measurable: short-high Kd recovery plus normal-path preservation.")
    base.add_source(slide, "Conclusion applies to the current io_buf study and declared 50 ohm || 2 pF validation scope.")
    base.add_notes(
        slide,
        """
Close with a precise claim. The two-state directional gate architecture is real progress. It removes forced replay restart, derives all core parameters from the original coefficient tables, passes complete-edge reconstruction after the directional and rate refinements, and produces at least one coefficient-correct interrupted result.

It is not yet a general replacement. Normal long-pulse behavior regresses relative to legacy, short-high Kd recovery remains wrong, and the evidence is currently concentrated on io_buf.

The next step should not be another wholesale model rewrite. Keep the validated GUP/GDN state equations, four directional paths, and maps. Add the minimum causal state needed to represent a pending pulldown-off command or recovery phase during short-high interruption. Outside a proven interruption window, preserve legacy-equivalent normal behavior. Then validate across pulse widths, both directions, repeated toggles, loads, and additional IBIS buffers.
""",
    )


def slide_appendix_residual(prs: Presentation) -> None:
    slide = base.add_slide(prs, "Appendix: why the Kd rate residual was added")
    add_section_tag(slide, "Optional detail")
    image = RESULT_ROOT / "fit_diagnostics" / "directional_maps_and_residual.png"
    base.add_picture_contain(slide, image, 0.55, 0.95, 7.5, 5.6, border=True)

    base.add_text(slide, "Observed limitation", 8.35, 1.05, 3.5, 0.35, size=19, bold=True, color=base.RED)
    base.add_box(slide, "Static directional maps reproduce the main Kd shape but miss its small negative excursion.", 8.42, 1.48, 4.05, 0.95, fill=base.LIGHT_RED, line=base.RED, size=16, bold=True)
    base.add_text(slide, "Measured refinement", 8.35, 2.75, 3.5, 0.35, size=19, bold=True, color=base.GREEN)
    base.add_code_box(slide, "Kd = Kd_map(GDN)\n   + residual_table\n   + 0.000238691 ns * dGDN/dt", 8.42, 3.15, 4.05, 1.2, size=13)
    base.add_box(slide, "This refinement lowers worst reconstruction max error from 0.07389 to 0.04869.", 8.42, 4.72, 4.05, 0.82, fill=base.LIGHT_GREEN, line=base.GREEN, size=15.5, bold=True)
    base.add_box(slide, "It improves normal-table fidelity; it does not solve short-high Kd recovery staging.", 8.42, 5.75, 4.05, 0.72, fill=PALE_ORANGE, line=base.ORANGE, size=14.5, bold=True)
    base.add_takeaway(slide, "Residual is a measured correction to the gate-state model, not a substitute for state memory.")
    base.add_source(slide, "Cached fit diagnostic and generated BKDRES/BGDNRATE implementation.")
    base.add_notes(
        slide,
        """
Use this slide only if the audience asks how the final reconstruction gate was passed. Direction-specific static maps removed most of the error but missed a small negative Kd excursion. That feature depends on how quickly GDN is moving, not only on its instantaneous value.

The implemented correction is the directional Kd map plus a fitted residual table and a small term proportional to dGDN/dt. The fitted gain is about 0.000238691 ns. This reduces worst maximum complete-edge reconstruction error from about 0.07389 to 0.04869 and passes the offline gate.

The residual does not solve the short-high recovery problem. That problem is a larger timing and staging disagreement after the reverse edge. The appendix therefore keeps the distinction clear: GUP/GDN provide history, directional maps provide the main static relationship, the residual restores a dynamic detail, and a future recovery-phase state is still needed for the open quadrant.
""",
    )


def slide_appendix_references(prs: Presentation) -> None:
    slide = base.add_slide(prs, "Appendix: reference hierarchy and testbench scope")
    add_section_tag(slide, "Optional detail")

    base.add_box(slide, "Input\n3.3 V, 1 ps", 0.65, 1.35, 1.6, 0.82, fill=base.LIGHT_GRAY, line=base.GRAY, size=16, bold=True)
    base.add_arrow(slide, 2.25, 1.76, 3.0, 1.76)
    base.add_box(slide, "io_buf driver", 3.0, 1.35, 2.0, 0.82, fill=base.LIGHT_GREEN, line=base.GREEN, size=16, bold=True)
    base.add_arrow(slide, 5.0, 1.76, 5.75, 1.76)
    base.add_box(slide, "PAD", 5.75, 1.45, 1.25, 0.62, fill=base.LIGHT_GRAY, line=base.GRAY, size=16, bold=True)
    base.add_arrow(slide, 7.0, 1.76, 7.75, 1.76)
    base.add_box(slide, "50 ohm || 2 pF", 7.75, 1.35, 2.1, 0.82, fill=base.LIGHT_GRAY, line=base.GRAY, size=16, bold=True)

    base.add_box(slide, "HSPICE native IBIS", 0.78, 3.0, 3.55, 0.75, fill=RGBColor(245, 245, 245), line=base.BLACK, size=17, bold=True)
    base.add_bullets(slide, ["Pad + Ku + Kd", "Coefficient-playback reference", "Cached; validation only"], 0.88, 3.92, 3.3, 1.25, size=15)
    base.add_box(slide, "HSPICE transistor io_buf.sp", 4.78, 3.0, 3.55, 0.75, fill=base.LIGHT_GRAY, line=base.GRAY, size=17, bold=True)
    base.add_bullets(slide, ["Pad only", "Circuit-level reference", "No physical Ku/Kd observables"], 4.88, 3.92, 3.3, 1.25, size=15)
    base.add_box(slide, "ngspice + pybis", 8.78, 3.0, 3.55, 0.75, fill=base.LIGHT_GREEN, line=base.GREEN, size=17, bold=True)
    base.add_bullets(slide, ["Pad + Ku + Kd + G states", "Candidate under test", "Derived from IBIS tables only"], 8.88, 3.92, 3.3, 1.25, size=15)

    base.add_box(slide, "Important", 0.82, 5.72, 1.35, 0.55, fill=base.LIGHT_RED, line=base.RED, size=15, bold=True)
    base.add_text(slide, "Native IBIS and transistor io_buf.sp are different references. Ku/Kd agreement means playback consistency; transistor agreement is a separate pad-level check.", 2.35, 5.68, 9.85, 0.7, size=16, color=base.DARK, bold=True)
    base.add_takeaway(slide, "HSPICE is never used to tune the generated model; it is an external audit of cached results.")
    base.add_source(slide, "Canonical study: results/io_buf_two_state_gate_model_2026-06-30")
    base.add_notes(
        slide,
        """
The testbench holds edge rate, voltage, and load constant so the algorithm is the variable. The input is 3.3 V with 1 ps edges, and the output drives 50 ohms in parallel with 2 pF.

There are two HSPICE references with different meanings. HSPICE native IBIS exposes pad voltage and internal Ku/Kd playback coefficients, so it is the coefficient reference. HSPICE transistor-level io_buf.sp exposes pad voltage only; it has no Ku or Kd because those are IBIS model coefficients, not physical transistor nodes.

The two references can disagree at the pad. Therefore a claim must say which reference is being matched. The ngspice gate-state model is fitted only from IBIS-derived coefficient tables. Cached HSPICE results are used afterward for validation and are not an input to parameter extraction.
""",
    )


def build_deck(template: Path, output: Path) -> None:
    prs = Presentation(str(template))
    remove_template_content_slides(prs)
    set_title_slide(prs)

    slide_problem(prs)
    slide_legacy_ku_kd(prs)
    slide_core_model(prs)
    slide_gate_derivation(prs)
    slide_tau_delay(prs)
    slide_command_caps(prs)
    slide_mapping(prs)
    slide_full_syntax(prs)
    slide_real_example(prs)
    slide_reconstruction(prs)
    slide_results(prs)
    slide_status(prs)
    slide_appendix_residual(prs)
    slide_appendix_references(prs)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build a concise gate-state meeting deck.")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build_deck(args.template, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
