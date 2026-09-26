#!/usr/bin/env python3
"""Build the 2026-09-17 deck: inside the buffer, what it taught us, and the open question.

This deck deliberately stops before the solution. It shows the internal-stage study on
three buffers, what was learned, the verification that driving Ku/Kd from the transistor's
own gate node works, why the shipped model's gate is not that node, the question that
leaves, and the experiments that did not answer it. Track 1, track 2 and the four-number fit
are not on these slides.

Style follows the 0904 and 0911 decks:

* Simulation figures only; a number is listed, not charted. One relaxation, agreed for this
  deck: one slide of schematic, drawn as native shapes, for the stage structure.
* Two or three lines state the point; the rest of the slide is the waveform.
* The wrong turns stay on the slides with their corrections.

One layout rule learned on this deck's first render: the 0904 script placed every figure at
a fixed y under a fixed 0.75 in per bullet, and any bullet that wrapped put the figure
straight through the text. Here `points()` returns the y it actually used and figures are
placed from that.

    py -3.14 scripts/build_0917_deck_figures.py     # first
    py -3.14 scripts/build_0917_deck.py
    powershell -ExecutionPolicy Bypass -File scripts/render_deck_slides.ps1 `
        -Deck results/meeting_deck_2026-09-17/inside_the_buffer_2026-09-17.pptx
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
for _p in (ROOT, ROOT / ".codex_deps" / "presentation" / "python"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

from pptx.dml.color import RGBColor  # noqa: E402

from tools.presentation_kit import GreenDeck  # noqa: E402

R = ROOT / "results"
OUT = R / "meeting_deck_2026-09-17" / "inside_the_buffer_2026-09-17.pptx"
F = R / "meeting_deck_2026-09-17" / "figures"
F11 = R / "meeting_deck_2026-09-11" / "figures"
T2 = R / "track2_train_check_2026-09-13" / "figs"

FIG = {
    "recap_pad": F / "recap_pad.png", "recap_kukd": F / "recap_kukd.png",
    "ku_vs_gate": F / "ku_vs_gate.png",
    "gate_vs_real_three": F / "gate_vs_real_three.png",
    "replay_ex2": F / "replay_ex2.png", "replay_iobuf": F / "replay_iobuf.png",
    "cascade_ex2": F / "cascade_ex2.png", "slew_ex2": F / "slew_ex2.png",
    "reversal_rules": F / "reversal_rules.png", "value_match_split": F / "value_match_split.png",
    "ccomp_ex2": F / "ccomp_ex2.png", "cascade_pair": F / "cascade_pair.png",
    "superpose_ex2": F / "superpose_ex2.png",
    "static_map": F / "static_map_ex2.png",
    "inv_why": F / "inv_why.png", "replay_inv_fail": F / "replay_inv_fail.png",
    "replay_inv_fixed": F / "replay_inv_fixed.png", "iobuf_bump": F / "iobuf_bump.png",
    "step_ex2": F / "gate_step_ex2.png", "step_inv": F / "gate_step_inv_chain.png",
    "step_io": F / "gate_step_io_buf.png",
    # slide-sized figures from the 0911 deck and the studies, reused as they are.
    # The stressed walk only: the full-swing walks were shown on 09-11, and stacking
    # both on one slide halved each into a 2.5 in box with 5 pt axis text.
    "walk_ex2": F / "walk_ex2.png",
    "walk_inv": F / "walk_inv_chain.png",
    "walk_io": F / "walk_io_buf.png",
    "inv_maps": R / "gate_cascade_prototype_2026-09-09" / "inv_chain_c0.6" / "gate_replay_silicon_full" / "maps.png",
    "pu_off": R / "pu_off_conflict_2026-09-07" / "pu_off_conflict.png",
}

GREEN = RGBColor(43, 122, 67)
PALE = RGBColor(241, 248, 242)
INK = RGBColor(25, 31, 28)
GATE_FILL = RGBColor(255, 236, 214)
GATE_LINE = RGBColor(192, 86, 33)

LINE = 0.30      # inches per line at 17 pt: measured 1.15 in for four one-line bullets
CHARS = 110      # characters per line across 12.1 in at 17 pt; measured ~118, kept conservative
BOTTOM = 7.05    # the usable bottom of a slide


def points(deck, slide, items, y=1.28, size=17.0):
    """The lines that state the point. Returns the y just below them.

    Sized from the render, not from a guess: the first pass used 0.40 in per line and 100
    characters, which put every figure under a band of white and opened a 1.6 in gap on the
    closing slide. Under-estimating is the worse failure - text through a figure - so the
    character count stays a little conservative."""
    lines = sum(max(1, -(-len(t) // CHARS)) for t in items)
    h = LINE * lines + 0.04 * len(items)
    deck.add_bullets(slide, items, 0.7, y, 12.1, h, size=size, spacing=3)
    return y + h + 0.14


def wide(y):
    return dict(x=0.55, y=y, w=12.2, h=BOTTOM - y)


def halves(y):
    h = BOTTOM - 0.1 - y
    return dict(x=0.45, y=y, w=6.1, h=h), dict(x=6.75, y=y, w=6.1, h=h)


def code(deck, slide, text, y, *, size=14.5, w=12.0, x=0.7):
    """A monospaced block sized to its own line count, not to a guess."""
    n = text.count("\n") + 1
    # The kit pads the box itself; the first render left ~0.4 in empty under every table.
    h = 0.062 * size * n / 2.6 + 0.08
    deck.add_code_box(slide, text, x, y, w, h, size=size)
    return y + h + 0.15


def chain(deck, slide, y, label, nodes, gate=None, x0=1.9, bw=1.05, gap=0.32, bh=0.5):
    """One row of boxes joined by lines. `gate` names the box to highlight."""
    deck.add_text(slide, label, 0.45, y - 0.02, 1.4, 0.55, size=15, bold=True, color=INK)
    x = x0
    for i, n in enumerate(nodes):
        is_gate = n == gate
        deck.add_box(slide, n, x, y, bw, bh, size=13,
                     fill=(GATE_FILL if is_gate else PALE),
                     line=(GATE_LINE if is_gate else GREEN), bold=is_gate)
        if i:
            deck.add_arrow(slide, x - gap + 0.02, y + bh / 2, x - 0.02, y + bh / 2, width=1.8)
        x += bw + gap
    return x


def main() -> int:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    missing = [k for k, v in FIG.items() if not v.exists()]
    if missing:
        print("missing figures: " + ", ".join(missing))
        return 1

    d = GreenDeck()
    d.set_title_slide(
        "Inside the Buffer:\nWhat the Gate Tells Us",
        "Simon Hwang\nDr. Chulsoon Hwang\nDr. Zhiping Yang\n9/17/2026")

    # ---------------------------------------------------------------- contents
    s = d.add_slide("Table of contents")
    points(d, s, [
        "The failure, in one slide",
        "ex2, start to finish: the stages, the gate, the output stage, the real gate put in",
        "inv_chain: the same four steps",
        "io_buf: the same four steps",
        "Three buffers, side by side",
        "The question this leaves",
        "What we tried after looking inside",
    ], y=1.5)

    # ------------------------------------------------------------------- recap
    s = d.add_slide("The failure, in one slide", section="Recap")
    y = points(d, s, [
        "A pulse cut short reverses before the pad has finished. Through the edge all "
        "three agree; after it, both IBIS models overshoot.",
        "ex2 at 810 ps: native +70.7 %, ours +66.8 %. Eleven of twelve buffers fail this "
        "way, +34…+71 % at 50 % depth. io_buf alone stays within ±3 %.",
    ])
    left, right = halves(y)
    d.add_picture_contain(s, FIG["recap_pad"], **left)
    d.add_picture_contain(s, FIG["recap_kukd"], **right)
    d.add_notes(s, "Same case as the film. Our model is the shipped gate-state build, "
                   "InputDrivenTwoStateGateDelayCommandFull, at the file's declared C_comp of "
                   "5 pF, like native: +66.8 %. The ex2 section (slides 5-8) uses the measured "
                   "1.7 pF, where the same model is +73.7 %. The Ku panel shows why: "
                   "both models keep Ku rising after the reversal, as if the transition "
                   "were still completing. io_buf being within 3 % on the peak is the "
                   "one anomaly in the twelve, and it is picked up again in section 3.")

    # ============================================================ what we probed
    s = d.add_slide("What we probed", section="Inside the buffer")
    points(d, s, [
        "Transistor-level HSPICE, 3.3 V, 50 Ω and 2 pF on the pad, 50 ps edges. Every "
        "internal node of the predriver recorded, on a full swing and on five pulse widths.",
        "ex2 first, start to finish. Then the same four steps on the other two.",
    ])
    chain(d, s, 2.75, "ex2", ["in", "n2", "n3", "n4", "pad"], gate="n4")
    chain(d, s, 3.7, "inv_chain", ["in"] + [f"vout{k}" for k in range(1, 8)] + ["pad"],
          gate="vout7", bw=0.98, gap=0.22)
    chain(d, s, 4.65, "io_buf, pull-down", ["in", "n1", "nand_n3", "n3", "pad"], gate="n3")
    chain(d, s, 5.6, "io_buf, pull-up", ["in", "n2", "pad"], gate="n2")
    d.add_text(s, "highlighted: the last predriver node - the gate that drives the output stage",
               0.7, 6.4, 12.0, 0.45, size=14, color=GATE_LINE)
    d.add_notes(s, "The one schematic in the deck, by agreement. The four steps, the same on "
                   "every buffer: 1 the pulse down the chain, 2 the real gate against ours, "
                   "3 the real gate put into our model, 4 what that buffer adds. "
                   "Runs: results/predriver_stages_2026-09-09.")

    # ============================================================ ex2
    s = d.add_slide("ex2, step 1: the pulse down the chain", section="ex2")
    y = points(d, s, [
        "At full swing every node is later and softer than the one before: n2 at 248 ps, "
        "n3 at 730, n4 at 1096, pad at 1400.",
        "Cut the pulse to 810 ps and the last node n4 (orange, the gate of the output stage) "
        "stops at 0.76 and turns round. The pad reaches half its swing.",
    ])
    d.add_picture_contain(s, FIG["walk_ex2"], **wide(y))
    d.add_notes(s, "The recap's case, 810 ps. Orange is the same node that is highlighted on "
                   "the schematic. n4's 10-90 rise is 658 ps, so an 810 ps pulse catches it "
                   "part of the way up. n4's own step responses, superposed, predict 0.93: "
                   "each stage answers a short pulse with less than the sum of its steps.")

    s = d.add_slide("ex2, step 2: the real gate against ours", section="ex2")
    y = points(d, s, [
        "Our model has a gate node too. It runs on a schedule: a fixed delay after each input "
        "edge, then a fixed ramp. So it turns round 189 ps after the real gate does.",
        "What the pad sees is Ku, below. In those 189 ps ours keeps climbing to 0.96 while "
        "the transistor's is already falling. That is the overshoot.",
    ])
    d.add_picture_contain(s, FIG["step_ex2"], **wide(y))
    d.add_notes(s, "Compare WHEN the gates turn round, not how high they get. Our Ku curve is "
                   "built by pairing the file's Ku(t) with our own gate's full-swing ramp, so "
                   "on an uninterrupted edge curve(gate) reproduces the file's Ku(t) exactly, "
                   "whatever level our gate sits at: a lower gate simply gets a steeper curve. "
                   "189 ps is 29 % of this gate's 658 ps rise. The transistor's Ku is masked "
                   "from 20 ps before to 90 ps after each edge, where the solve is "
                   "differentiation noise.")

    s = d.add_slide("ex2, step 3: the output stage only follows the gate", section="ex2")
    y = points(d, s, [
        "Solve Ku from two fixture runs (a), record the gate at the same instants (b), plot "
        "one against the other (c): time drops out and one curve is left.",
        "So the buffer is two things in series: a predriver that moves the gate, and an output "
        "stage that turns 'how open' into pad current through one fixed curve.",
    ])
    d.add_picture_contain(s, FIG["static_map"], **wide(y))
    d.add_notes(s, "If the curve is really fixed, then the right gate plus this curve must "
                   "give the right pad under any pulse. That is the next slide. On ex2 the "
                   "curve is single-valued only once C_comp in the solve is the measured "
                   "1.7 pF rather than the declared 5 (results/gate_physics_2026-09-08).")

    s = d.add_slide("ex2, step 4: put the real gate in", section="ex2")
    y = points(d, s, [
        "The test: replace our gate with the transistor's probed n4. Everything else "
        "unchanged, including the file's own Ku curve.",
        "810 ps: +73.7 % becomes -5.5 %. Across five widths -3.3...-5.8 %, and the lag drops "
        "from 149...300 ps to 18...24. On ex2 the whole error was the gate's motion.",
    ])
    d.add_picture_contain(s, FIG["replay_ex2"], **wide(y))
    d.add_notes(s, "Shipped across the five widths: +3.7 / +13.3 / +26.4 / +45.8 / +73.7 %. "
                   "Real gate: -3.3 / -4.7 / -5.8 / -5.4 / -5.5 %. "
                   "Source: results/gate_cascade_prototype_2026-09-09/ex2_c1.7/gate_replay.")

    # ============================================================ inv_chain
    s = d.add_slide("inv_chain, step 1: the pulse down the chain", section="inv_chain")
    y = points(d, s, [
        "Seven fast inverters, about 40 ps each. Cut the pulse to 104 ps and the chain still "
        "passes it: the gate vout7 (orange) reaches 0.88.",
        "The pad reaches only 0.49. On this buffer the pulse is lost at the output stage, not "
        "in the predriver.",
    ])
    d.add_picture_contain(s, FIG["walk_inv"], **wide(y))
    d.add_notes(s, "Every inverter sits within superposition of its own step responses. The "
                   "gate crosses 10-90 % in 52 ps, against ex2's 658.")

    s = d.add_slide("inv_chain, step 2: the real gate against ours", section="inv_chain")
    y = points(d, s, [
        "Same schedule, same symptom: ours turns round 29 ps late. Small - but more than half "
        "of this gate's 52 ps rise time.",
        "Ku runs on to 0.97 where the transistor's stops near 0.85 and is already falling: "
        "+31 % on the pad.",
    ])
    d.add_picture_contain(s, FIG["step_inv"], **wide(y))
    d.add_notes(s, "29 ps is 56 % of the gate's rise; ex2's 189 ps was 29 % of its own.")

    s = d.add_slide("inv_chain, step 3: put the real gate in", section="inv_chain")
    y = points(d, s, [
        "Same test as ex2. Here it fails: +31.1 % becomes -68.8 %. With the right gate, the "
        "pad is worse.",
        "So on inv_chain the gate is not the whole story. The next slide is why.",
    ])
    d.add_picture_contain(s, FIG["replay_inv_fail"], **wide(y))
    d.add_notes(s, "Five widths: -13.6 / -29.9 / -41.3 / -53.8 / -68.8 %. Our shipped model "
                   "hides this defect: its late gate and the curve problem partly cancel.")

    s = d.add_slide("inv_chain, step 4: why - one gate value, two different Ku", section="inv_chain")
    y = points(d, s, [
        "Our model keeps one Ku curve for a rising gate and one for a falling gate. A full "
        "swing only ever uses one at a time, and both versions agree (left).",
        "A short pulse must jump from one to the other at the top. From the file, the same "
        "gate value means Ku 0.64 rising and 0.16 falling, so Ku collapses. On the transistor "
        "the two are nearly one curve: 0.92 and 0.76.",
    ])
    d.add_picture_contain(s, FIG["inv_why"], **wide(y))
    d.add_notes(s, "Both runs carry the same real gate. Teal turns it into Ku with curves "
                   "built from the file; green with curves measured on the transistor. The "
                   "jump is at 0.326 ns, gate 0.83 and still rising: the direction selector "
                   "leads the gate by a few ps. Ku then peaks at 0.48 against 0.83. At full "
                   "swing the two Ku agree to 7 ps. The file's tables come from complete "
                   "transitions only, so nothing in them ties the rising curve to the falling "
                   "one part-way up. This is step 3 of ex2 failing: on the transistor the "
                   "output stage IS one curve; what the file implies is two.")

    s = d.add_slide("inv_chain, step 4 continued: with the transistor's own curves", section="inv_chain")
    y = points(d, s, [
        "Keep the real gate, and take the Ku curves from the transistor instead of the file: "
        "-68.8 % becomes +9.4 %. Five widths: +2.0...+9.4 %.",
        "inv_chain needs both: the gate's motion, and a Ku curve that is the same going up "
        "and coming down.",
    ])
    d.add_picture_contain(s, FIG["replay_inv_fixed"], **wide(y))
    d.add_notes(s, "Measured curves, five widths: +2.0 / +2.8 / +4.6 / +7.0 / +9.4 %. The "
                   "measured curve is evidence of what the transistor knows; how a converter "
                   "would get it is part of the question, not of this deck.")

    # ============================================================ io_buf
    s = d.add_slide("io_buf, step 1: two paths to the pad", section="io_buf")
    y = points(d, s, [
        "io_buf drives its two output transistors separately. Pull-up: the PMOS gate n2 is a "
        "slow 2.6 ns ramp, and the 2090 ps pulse cuts it short.",
        "Pull-down: the NMOS gate n3, through a NAND and an inverter. It turns the pull-down "
        "back on 1.9 ns after the reversal - the small bump on the pad. It returns in step 4.",
    ])
    d.add_picture_contain(s, FIG["walk_io"], **wide(y))
    d.add_notes(s, "Orange in each panel is that path's last node, the gate of its output "
                   "transistor. Every node is drawn 0 = at rest, 1 = fully switched, so on the "
                   "pull-down panel n3 at 1 means the pull-down is OFF. io_buf is linear from "
                   "input to pad. This section is at 2090 ps because that is where the gate "
                   "test of step 3 ran.")

    s = d.add_slide("io_buf, step 2: the real gates against ours", section="io_buf")
    y = points(d, s, [
        "Pull-up (left): our gate starts 1 ns late and sits lower all the way - and Ku is "
        "right anyway. Our Ku curve is built against our own gate, so a lower gate just gets "
        "a steeper curve.",
        "What matters is when it turns round: 51 ps late, on a gate that takes 2.7 ns to "
        "rise - 2 %. ex2 was 29 %, inv_chain 56 %. That is why io_buf's peak is within 3 %.",
    ])
    d.add_picture_contain(s, FIG["step_io"], **wide(y))
    d.add_notes(s, "This is the answer to 'our GUP/GDN only seems to work on io_buf': it is "
                   "the same schedule as on the other two, and the schedule's lateness is "
                   "negligible against a 2.7 ns ramp. Right-hand panels: the pull-down gate "
                   "and Kd. Look at 3.5-4.5 ns: the real gate eases back on from 3.6 ns, ours "
                   "snaps on at 3.95. That is the bump of step 4. On the right the probe is "
                   "drawn 1 = pull-down on, to match the model's node.")

    s = d.add_slide("io_buf, step 3: put the real gates in", section="io_buf")
    y = points(d, s, [
        "Put the real gates in: the peak barely moves, +1.6 to +1.4 %, because it was already "
        "right. What improves is the timing: lag 67 ps to 37.",
        "The box marks the bump from step 1. Next slide.",
    ])
    d.add_picture_contain(s, FIG["replay_iobuf"], **wide(y))
    d.add_notes(s, "Both gates replaced: n2 for the pull-up, n3 for the pull-down. The test "
                   "ran at two of io_buf's five widths; ngspice stalls on the deeper three.")

    s = d.add_slide("io_buf, step 4: the bump - the schedule, seen on the pad", section="io_buf")
    y = points(d, s, [
        "The bump is the pull-down turning back on. On the transistor it comes 202 ps earlier "
        "after a shorter pulse: the predriver got less far, so it has less far to come back.",
        "Ours comes back 2.02 ns after every reversal, whatever the pulse was. It is the same "
        "defect as ex2's - a schedule where the buffer has a state - where the peak cannot see it.",
    ])
    d.add_picture_contain(s, FIG["iobuf_bump"], **wide(y))
    d.add_notes(s, "Why show it: io_buf's peak being right could be read as 'our gate works "
                   "here'. It does not; the schedule's error has moved to the one event on "
                   "this buffer that is timed from the predriver's state. Mechanism: at the "
                   "shorter pulse n1 is caught at 0.854 instead of 0.978, decays from lower "
                   "and crosses the NAND sooner. Five parameters and two command forms were "
                   "tried and none moved it.")

    # ============================================================ side by side
    s = d.add_slide("Three buffers, side by side", section="What we learn")
    y = code(d, s, """                               ex2 (810 ps)       inv_chain (104 ps)       io_buf (2090 ps)

where the pulse is lost        in the predriver   at the output stage      nowhere: a slow ramp cut short
our gate turns round late by   189 ps             29 ps                    51 ps
  ... as a share of its rise   29 %               56 %                     2 %
our model, pad peak            +73.7 %            +31.1 %                  +1.6 %
real gate put in               -5.5 %             -68.8 %                  +1.4 %, lag 67 -> 37 ps
what is still missing          nothing            one Ku curve, not two    when the pull-down comes back""",
             1.45, size=13, x=0.5, w=12.3)
    d.add_text(s, "Our gate runs on a schedule; the buffer's gate is a state. How much that "
                  "costs depends on how fast the gate is.",
               0.7, y + 0.2, 12.0, 0.9, size=18, bold=True)
    d.add_text(s, "\"Stressed\" covers three mechanisms, so any one correction that treats "
                  "them alike fits one buffer and misses two.",
               0.7, y + 1.15, 12.0, 0.9, size=17)
    d.add_notes(s, "Every number on this slide is on one of the slides before it, at the "
                   "same pulse width.")

    # ============================================================ 4. question
    s = d.add_slide("The question", section="The question")
    d.add_text(s, "The output stage is a fixed map of the gate.\n"
                  "The gate comes from the predriver.\n"
                  "Driven by the real gate, the model is right.",
               0.7, 1.55, 12.0, 1.9, size=22)
    d.add_text(s, "The IBIS file has no gate, and its V-T tables pin Ku(t) only at full swing.\n"
                  "Native fails for the same reason with no gate at all: its Ku is a clock.",
               0.7, 3.55, 12.0, 1.3, size=18)
    d.add_text(s, "What information would reproduce the gate's motion - for any input, "
                  "including one cut short - without the transistor?",
               0.7, 5.05, 12.0, 1.2, size=22, bold=True, color=GATE_LINE)
    d.add_notes(s, "Silicon data settled the map-gate pair for our three buffers, but a "
                   "converter does not have silicon. The two halves of the question: "
                   "how the gate moves, and how the pad follows it. The next section is what "
                   "we tried after we started looking inside, in the order we tried it: from "
                   "the first probe of the real gate (09-08) to a linear predriver built from "
                   "the file (09-10). What came after that - a non-linear stage fitted at full "
                   "swing - is the next deck. The methods from before the probe (reversal-entry "
                   "rules, timing parameters, pu_off) are in backup.")

    # ============================================================ 5. tried after looking inside
    s = d.add_slide("1. Correct C_comp to the measured value", section="Tried after looking inside")
    y = points(d, s, [
        "The first thing the probed gate gave us: Ku against the real gate only forms one "
        "curve at one C_comp - 1.5-1.75 pF on ex2, at all five widths. The file declares 5.0.",
        "Put the measured 1.7 pF into our model and the overshoot grows (858 ps, the width the "
        "study ran). Native gets worse too. The declared 5 pF was hiding part of the gate error.",
    ])
    d.add_picture_contain(s, FIG["ccomp_ex2"], **wide(y))
    d.add_notes(s, "results/ex2_ccomp_correction_2026-09-08. Peak / lag / RMSE at 858 ps: "
                   "ours +244 mV / +236 ps / 336 mV at 5.0 pF, +289 / +268 / 382 at 1.7; native "
                   "+263 -> +342 mV. At 975 ps: ours +20 -> +52 mV, native +18 -> +75. The "
                   "full-swing control does not move. Lesson kept since: a single-parameter "
                   "improvement on this model has to be read against a compensating error "
                   "having moved. Fix the gate first.")

    s = d.add_slide("2. Slow the gate", section="Tried after looking inside")
    y = points(d, s, [
        "Slow our gate's ramp to match how slowly the probed gate moves. ex2 at 810 ps: +66.6 "
        "to +48.7 % - a quarter off, not a fix. The gate-ramp form does no better, 70.9 to 66.9 %.",
        "At 810 ps our pad is already 46 % too tall at +600 ps, before our model's turn-off "
        "even begins at +733: nothing done to the ramp can act before the schedule does.",
    ])
    d.add_picture_contain(s, FIG["slew_ex2"], **wide(y))
    d.add_notes(s, "The figure is the slew500ps build from gate_cascade_prototype_2026-09-09, "
                   "the one slowed build with raw runs on disk at this width. The gate-ramp "
                   "prototype (results/gate_ramp_prototype_2026-09-09, k-scaled ramp with the "
                   "map re-derived) is the source of the 70.9 to 66.9 % figure and of the "
                   "quote; it did fix io_buf's pedestal, +63…+69 ps to +10 / +9 / +2 / −4 / "
                   "−20, and took inv_base8 from 65 to 49 %. io_buf and the other eleven "
                   "want opposite fall treatments, single map against dual; the cost on "
                   "io_buf's peak was −3.5 to −10.5 %.")

    s = d.add_slide("3. An RC cascade in place of the schedule", section="Tried after looking inside")
    y = points(d, s, [
        "The probe showed ex2's gate is slow and smooth, so replace our fixed delay with 5 RC "
        "stages - a gate that moves by itself. ex2 at 810 ps: -2.6 %, where our model built "
        "alongside it is +66.6 %. It looked like the answer.",
        "On an inverter chain the same idea is dead: a 102 ps pulse cannot get through 4 RC "
        "stages (-98 %); real inverters regenerate. Delay plus one 60 ps stage only halves "
        "it, +65 to +22 %.",
    ])
    d.add_picture_contain(s, FIG["cascade_pair"], **wide(y))
    d.add_notes(s, "results/gate_cascade_prototype_2026-09-09. inv_base8 is the eight-inverter "
                   "variant of the family; inv_chain itself was not run as a cascade. Right "
                   "panel at its 50 % depth (102 ps). Cascades of 1-4 stages all give -82...-99 "
                   "%. On the ex2 family N = 5 gives -9...+8 %. The next slide shows why the ex2 "
                   "result is a fit: no linear structure reproduces ex2's gate. io_buf: neutral. "
                   "Our model built in this folder is +66.6 % on ex2 at 810 ps, not the +73.7 % of "
                   "the ex2 section: this folder is at the declared C_comp of 5 pF, the ex2 "
                   "section at the measured 1.7 pF. The recap (slide 3) is 5 pF too, +66.8 %.")

    s = d.add_slide("4. Derive the gate from the file, then superpose", section="Tried after looking inside")
    y = points(d, s, [
        "Assume the output stage's curve has a MOSFET shape, ((g - vt)/(1 - vt))^α, and invert "
        "the file's Ku(t) through it: the gate comes out of the file, no transistor (left). Below the threshold (shaded) the file is blind.",
        "Then treat the predriver as linear: a cut pulse is the rising step plus the falling "
        "step. ex2: 0.99 where the real gate stops at 0.76 (right). inv_chain: 0.99 against "
        "0.88. io_buf: within 0.035.",
    ])
    d.add_picture_contain(s, FIG["superpose_ex2"], **wide(y))
    d.add_notes(s, "results/physics_map_gate_2026-09-10. The three measured curves share the "
                   "shape: vt 0.57 / 0.49 / 0.50, alpha 0.64 / 0.60 / 0.78, rms 0.068 / 0.070 / "
                   "0.014 (ex2 / inv_chain / io_buf). Through it the file gives inv_chain's gate "
                   "to 10 ps; ex2's to within about 60 ps at the 50 % point (38 ps in the study's "
                   "normalisation), its top end stretched (10-90 986 against 662 ps). So the tables can give the gate at full swing. What a linear "
                   "predriver cannot give is ex2's: three stages that each return less than the "
                   "sum of their steps. io_buf is linear, so superposition works there and gives "
                   "its gates where our schedule was 0.15-0.27 low. What came next - a non-linear "
                   "stage fitted at full swing - is the next deck.")

    s = d.add_slide("What we tried after looking inside, in one table", section="Tried after looking inside")
    code(d, s, """attempt, in order                   on ex2              on inv_chain family     on io_buf           verdict

1  measured C_comp, 1.7 pF          every metric worse  -                       -                   was hiding gate error
2  slow the gate, re-derive curve   +67 -> +49 % best   65 -> 49 %              pedestal fixed      cannot beat the schedule
3  RC cascade, 5 stages             +67 -> -3 %         -98 % (dead)            neutral             a fit, not physics
   delay + one RC stage             -                   +65 -> +22 %            -                   halves it
   residual scaled by depth         no effect           no effect               bump height back    io_buf only
   one gate for both sides          peak < 1.5 %        peak < 1.5 %            -                   not the lever
4  gate from the file, superposed   0.99 vs real 0.76   0.99 vs 0.88            within 0.035        no non-linearity""",
         1.4, size=12, x=0.5, w=12.3)
    d.add_notes(s, "In order from 09-08 to 09-10. Every row has a FINDINGS.md: "
                   "ex2_ccomp_correction_2026-09-08, gate_ramp_prototype_2026-09-09, "
                   "gate_cascade_prototype_2026-09-09 (cascade, hybrid, shared gate), "
                   "residual_depth_rule_2026-09-09, physics_map_gate_2026-09-10. The pattern: "
                   "each attempt fixes the buffer whose mechanism it happens to match and fails "
                   "the others. The inverter-chain family column is inv_base8 for rows 2-3 and "
                   "inv_chain for the rest.")

    # ---------------------------------------------------------------- close
    s = d.add_slide("Where this leaves us", section="Summary")
    d.add_text(s, "Settled:", 0.7, 1.3, 12.0, 0.5, size=17)
    y = points(d, s, [
        "the output stage is a fixed map of one gate node, on all three buffers",
        "driven by the real gate, the model is right - the error is the gate's motion",
        "our gate is a schedule: it turns round late - 29 % of a rise time on ex2, 56 % on inv_chain, 2 % on io_buf",
        "the file can give the gate at full swing; no linear predriver built from it gives a cut pulse on ex2",
    ], y=1.85)
    d.add_text(s, "Open:", 0.7, y + 0.15, 12.0, 0.5, size=17)
    points(d, s, [
        "how to reproduce the gate's motion under a pulse cut short, from information a converter can have",
        "inv_chain also needs one Ku curve going up and coming down; io_buf's pull-down must come back from its state",
    ], y=y + 0.7)
    d.add_notes(s, "This deck stops at the question on purpose. What the study has done "
                   "since - a fitted stage chain from the file plus one stressed pad run, "
                   "and a measured map - is the next deck.")

    s = d.add_slide("Before the gate: entering the reversal by rule", section="Backup: before looking inside")
    y = points(d, s, [
        "Fourteen rules for what Ku/Kd do when the input reverses mid-transition: keep the "
        "clock (t-matching), enter the opposite table where it holds the present Ku/Kd "
        "(value-matching), invert the gate trajectory, match the pad. Scored on thirty stressed "
        "cases, pad RMSE: native 96 mV, shipped 92, legacy 154. Nothing else beats native.",
        "io_buf at 1634 ps, the rules that follow the transistor up to the reversal: legacy "
        "overshoots to 1.16 V; t-matching and value-matching land 29 and 23 % low. "
        "Value-matching halves legacy's pad error, 33 vs 66 mV, with a Kd RMSE of 0.24 against "
        "native's 0.03 - the pad passes, the coefficient does not.",
    ])
    d.add_picture_contain(s, FIG["reversal_rules"], **wide(y))
    d.add_notes(s, "results/stress_method_matrix_2026-08-20 (summary.csv, per_case.csv, the "
                   "waveforms). The overlay uses the gated builds, time_match_hybrid and "
                   "coeff_match, as the 08-25 figures did; the ungated ...ReplayFull builders "
                   "score 818 and 373 mV because they replace the whole waveform path and come "
                   "up on the wrong rail before the pulse - t-matching with Ku = 0.94 in the "
                   "low state - which is a defect of that build, not the rule. All-cases "
                   "scoreboard: delay_cmd 92, native 96, predriver_cmd 99, hybrid 106, "
                   "measured_rate 110, pad_match 111, gate_state 122, coeff_match 138, legacy "
                   "154, gate_match 175. Every rule here is a way of choosing a time on the "
                   "opposite table. The next slide is why no such time exists.")

    s = d.add_slide("Why value-matching cannot enter a reversal", section="Backup: before looking inside")
    y = points(d, s, [
        "The rule: at the reversal, find where the opposite table already holds the present "
        "Ku and Kd, and replay from there. Just before io_buf's 2 ns reversal the state is "
        "Ku 0.264, Kd 0.053 - coherent on the rising table, 1.988 and 2.018 ns.",
        "On the falling tables the same pair sits 1.622 ns apart: Ku at 0.859 ns, Kd at 2.481. "
        "No replay time serves both. The forced midpoint is wrong for each, ngspice's timestep "
        "collapses on the jumps, and the 2 ns run never finished.",
    ])
    d.add_picture_contain(s, FIG["value_match_split"], **wide(y))
    d.add_notes(s, "results/io_buf_value_match_misalignment_demo_2026-06-25 (the numbers) and "
                   "io_buf_value_matched_replay_redo_2026-06-25 (the runs). The false pass on "
                   "the 1 ns pulse, from the redo: pad peak 1.516 V (legacy) to 0.027 V "
                   "against native's 0.062 - and Kd minimum +0.551 against native's -0.063, "
                   "the wrong sign. The pad-matched successor "
                   "(three_buffer_pad_matched_replay_v2_2026-08-04): 96 stressed rows, 12 "
                   "improved on pad, Ku and Kd together, 34 where the same pad voltage maps to "
                   "table times more than 0.5 ns apart. The state at a reversal is not a time "
                   "on either table; that is what the gate node was introduced to carry.")

    s = d.add_slide("Every timing parameter in the command layer", section="Backup: before looking inside")
    y = points(d, s, [
        "Scale every fitted command delay, then separately every gate time constant, down "
        "to 5 % of nominal, and watch io_buf's pedestal:",
    ])
    y = code(d, s, """delays:  x1  +85 ps    x0.5  +125    x0.25  +124    x0.05  +61
taus:    x1  +85 ps    x0.5  +118    x0.25  +116    x0.05  +95""", y, size=15)
    d.add_text(s, "Non-monotonic, and 61 of the 85 ps survives at a twentieth of nominal. "
                  "Every timing parameter is ruled out - and later, for good.",
               0.7, y + 0.1, 12.0, 1.0, size=17)
    d.add_notes(s, "results/delay_pedestal_2026-09-04. Confirmed as permanently ruled out "
                   "in v_indexed_prototype_2026-09-04. It is not a timing knob.")

    s = d.add_slide("pu_off: derived twice, retracted twice", section="Backup: before looking inside")
    y = points(d, s, [
        "PU_OFF_SCALE = 0.70 was derived against an artifact: the transistor's stressed Ku "
        "peak of 1.29 is C_comp·dV/dt differentiation noise and does not converge with the grid.",
        "The re-derived 0.29 destroys the +64.9 mV reversal overshoot. A ten-point sweep: "
        "best for Ku 0.10, best for the overshoot 0.86, strictly monotone and opposite.",
    ])
    d.add_picture_contain(s, FIG["pu_off"], **wide(y))
    d.add_notes(s, "One parameter, two jobs, opposite directions: pu_off sets both the gate "
                   "turn-off and its phase against the residual spike. No interior optimum "
                   "exists, so it is not derivable; both derivations were choosing which "
                   "requirement to satisfy while believing they had measured something. "
                   "Also retracted in the same file: that native re-solves the coefficient "
                   "at runtime, and that our gate decays more slowly than native's.")

    path = d.save(OUT)
    print(f"wrote {path.relative_to(ROOT)}  ({path.stat().st_size / 1024:.0f} KB, "
          f"{len(d.prs.slides._sldIdLst)} slides)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
