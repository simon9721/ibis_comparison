#!/usr/bin/env python3
"""Build a cached-data explainer deck for voltage-matched Ku/Kd replay."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
import sys


ROOT = Path(__file__).resolve().parents[1]
LOCAL_DEPS = ROOT / ".codex_deps" / "presentation" / "python"
for path in (LOCAL_DEPS, ROOT, ROOT / "scripts", ROOT / "tools" / "pybis2spice"):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from tools.presentation_kit import EquationRenderer, GreenDeck


STUDY = ROOT / "results" / "three_buffer_single_true_output_reversal_2026-08-11"
DEFAULT_OUT = ROOT / "results" / "voltage_matching_explainer_deck_2026-08-14"

BLACK = "#111111"
GRAY = "#737373"
GREEN = "#2B7A43"
RED = "#CA3030"
BLUE = "#2467AD"
ORANGE = "#DA8426"
PURPLE = "#7B2CBF"
CYAN = "#56B4E9"
GRID = "#D9DEE5"

PALE_BLUE = RGBColor(232, 242, 252)
PALE_ORANGE = RGBColor(255, 242, 228)
PALE_PURPLE = RGBColor(243, 237, 250)


def ensure(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def write_rows(path: Path, rows: list[dict[str, object]]) -> None:
    if not rows:
        return
    fields: list[str] = []
    for row in rows:
        for key in row:
            if key not in fields:
                fields.append(key)
    ensure(path.parent)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def load_csv_columns(path: Path, names: list[str]) -> dict[str, np.ndarray]:
    header = path.open(encoding="utf-8").readline().strip().split(",")
    indexes = [header.index(name) for name in names]
    values = np.loadtxt(path, delimiter=",", skiprows=1, usecols=indexes)
    if values.ndim == 1:
        values = values.reshape(1, -1)
    return {name: values[:, index] for index, name in enumerate(names)}


def event_row(device: str, direction: str) -> dict[str, str]:
    rows = read_rows(STUDY / "voltage_matching_mapping_events.csv")
    return next(row for row in rows if row["device"] == device and row["direction"] == direction)


def metric_row(device: str, direction: str) -> dict[str, str]:
    rows = read_rows(STUDY / "candidate_metrics.csv")
    return next(
        row
        for row in rows
        if row["device"] == device
        and row["direction"] == direction
        and row["flow"] == "voltage_matching"
    )


def style_axis(axis, *, xlabel: str = "", ylabel: str = "") -> None:
    axis.grid(True, color=GRID, linewidth=0.8)
    axis.spines[["top", "right"]].set_visible(False)
    axis.tick_params(labelsize=10)
    axis.set_xlabel(xlabel, fontsize=11)
    axis.set_ylabel(ylabel, fontsize=11)


def save_figure(figure, path: Path) -> Path:
    ensure(path.parent)
    figure.savefig(path, dpi=210, bbox_inches="tight", facecolor="white")
    plt.close(figure)
    return path


def plot_problem(figures: Path) -> Path:
    path = STUDY / "waveforms" / "inv_chain.csv"
    data = load_csv_columns(
        path,
        ["time_ns", "input_v", "hspice_native_pad_v", "hspice_transistor_pad_v"],
    )
    event = event_row("inv_chain", "short_high")
    reverse = float(event["input_reverse_ns"])
    t = data["time_ns"] - reverse
    mask = (t >= -0.22) & (t <= 0.75)

    figure, axes = plt.subplots(2, 1, figsize=(10.8, 6.4), sharex=True)
    axes[0].plot(t[mask], data["input_v"][mask], color=CYAN, linewidth=2.8)
    axes[0].axvline(0, color=PURPLE, linestyle="--", linewidth=1.6)
    axes[0].text(0.02, 0.82, "second edge", transform=axes[0].transAxes, color=PURPLE, fontsize=10)
    style_axis(axes[0], ylabel="Input (V)")
    axes[0].set_title("Input reverses before the output response is complete", loc="left", fontweight="bold")

    axes[1].plot(
        t[mask], data["hspice_native_pad_v"][mask], color=BLACK, linewidth=3.0,
        label="HSPICE native IBIS",
    )
    axes[1].plot(
        t[mask], data["hspice_transistor_pad_v"][mask], color=GRAY, linewidth=2.6,
        label="HSPICE transistor",
    )
    axes[1].axvline(0, color=PURPLE, linestyle="--", linewidth=1.6)
    style_axis(axes[1], xlabel="Time from second input edge (ns)", ylabel="Pad voltage (V)")
    axes[1].legend(loc="upper right", frameon=False, ncol=2)
    figure.tight_layout()
    return save_figure(figure, figures / "01_short_pulse_problem_inv_chain.png")


def plot_full_swing_calibration(figures: Path) -> Path:
    reference = json.loads(
        (STUDY / "calibration" / "io_buf" / "fast_5ps" / "pad_replay_reference.json").read_text(
            encoding="utf-8"
        )
    )
    figure, axes = plt.subplots(1, 2, figsize=(12.8, 5.2), sharey=True)
    for axis, direction, color in zip(axes, ("rising", "falling"), (BLUE, ORANGE)):
        t = np.asarray(reference[direction]["time_ns"], dtype=float)
        v = np.asarray(reference[direction]["pad_v"], dtype=float)
        axis.plot(t, v, color=color, linewidth=2.8)
        axis.set_title(f"{direction.capitalize()} full-swing pad calibration", fontweight="bold")
        style_axis(axis, xlabel="Time from input threshold (ns)", ylabel="Vpad (V)")
        axis.set_xlim(0, 8.5)
    figure.suptitle(
        "Real offline calibration: legacy pybis, io_buf, 50 ohm || 2 pF",
        fontsize=15,
        fontweight="bold",
    )
    figure.tight_layout(rect=(0, 0, 1, 0.94))
    return save_figure(figure, figures / "02_full_swing_calibration_io_buf.png")


def plot_inverse_mapping(figures: Path, device: str = "io_buf") -> Path:
    reference_path = STUDY / "calibration" / device / "fast_5ps" / "pad_replay_reference.json"
    reference = json.loads(reference_path.read_text(encoding="utf-8"))
    event = event_row(device, "short_high")
    falling = reference["falling"]
    t = np.asarray(falling["time_ns"], dtype=float)
    v = np.asarray(falling["pad_v"], dtype=float)
    sample = float(event["sampled_pad_v"])
    lookup = float(event["opposite_table_lookup_start_ns"])
    matched_v = float(np.interp(lookup, t, v))

    figure, axis = plt.subplots(figsize=(10.8, 5.8))
    axis.plot(t, v, color=BLUE, linewidth=2.8, label="full-swing falling pad trajectory")
    axis.axhline(sample, color=RED, linestyle="--", linewidth=1.5, label=f"sampled Vpad = {sample:.3f} V")
    axis.axvline(lookup, color=PURPLE, linestyle="--", linewidth=1.5, label=f"matched table time = {lookup:.3f} ns")
    axis.scatter([lookup], [matched_v], color=RED, edgecolor="white", linewidth=1.0, s=100, zorder=5)
    axis.annotate(
        "Use this time coordinate\nfor both falling Ku and Kd",
        xy=(lookup, matched_v), xytext=(0.58, 0.62), textcoords="axes fraction",
        arrowprops={"arrowstyle": "->", "color": RED}, fontsize=12,
    )
    axis.set_xlim(0, min(4.0, float(t[-1])))
    style_axis(axis, xlabel="Falling calibration-table time (ns)", ylabel="Calibration pad voltage (V)")
    axis.legend(loc="upper right", frameon=False)
    axis.set_title("Real io_buf mapping: voltage becomes one opposite-table time", loc="left", fontweight="bold")
    figure.tight_layout()
    return save_figure(figure, figures / "03_voltage_to_time_inverse_mapping_io_buf.png")


def plot_runtime_sample(figures: Path) -> Path:
    device = "io_buf"
    event = event_row(device, "short_high")
    data = load_csv_columns(
        STUDY / "waveforms" / f"{device}.csv",
        [
            "time_ns", "input_v", "hspice_native_pad_v", "voltage_matching_pad_v",
            "voltage_matching_padsamp", "voltage_matching_pmsample",
            "voltage_matching_pmlatchpulse", "voltage_matching_padmapactive",
            "voltage_matching_padstart_latch", "voltage_matching_padarg",
        ],
    )
    reverse = float(event["input_reverse_ns"])
    t = data["time_ns"] - reverse
    mask = (t >= -0.35) & (t <= 1.35)

    figure, axes = plt.subplots(3, 1, figsize=(11.2, 7.8), sharex=True)
    axes[0].plot(t[mask], data["hspice_native_pad_v"][mask], color=BLACK, linewidth=2.8, label="HSPICE native IBIS")
    axes[0].plot(t[mask], data["voltage_matching_pad_v"][mask], color=RED, linewidth=2.3, label="Voltage matching V2")
    axes[0].plot(t[mask], data["voltage_matching_padsamp"][mask], color=PURPLE, linewidth=2.0, label="PADSAMP")
    axes[0].scatter(
        [float(event["sample_capture_time_ns"]) - reverse], [float(event["sampled_pad_v"])],
        color=RED, edgecolor="white", s=80, zorder=6,
    )
    axes[0].legend(loc="upper right", frameon=False, ncol=3)
    style_axis(axes[0], ylabel="Voltage (V)")

    axes[1].plot(t[mask], data["voltage_matching_pmsample"][mask], color=RED, linewidth=2.0, label="sample pulse")
    axes[1].plot(t[mask], data["voltage_matching_pmlatchpulse"][mask], color=ORANGE, linewidth=2.0, label="start-time latch")
    axes[1].plot(t[mask], data["voltage_matching_padmapactive"][mask], color=GREEN, linewidth=2.2, label="replay active")
    axes[1].legend(loc="upper right", frameon=False, ncol=3)
    style_axis(axes[1], ylabel="Control")

    axes[2].plot(t[mask], data["voltage_matching_padstart_latch"][mask], color=PURPLE, linewidth=2.2, label="latched start")
    axes[2].plot(t[mask], data["voltage_matching_padarg"][mask], color=BLUE, linewidth=2.2, label="start + elapsed")
    axes[2].legend(loc="upper left", frameon=False, ncol=2)
    style_axis(axes[2], xlabel="Time from second input edge (ns)", ylabel="Table time (ns)")

    for axis in axes:
        axis.axvline(0, color=CYAN, linestyle="--", linewidth=1.3)
    axes[0].set_title("Runtime sequence: sample, latch, then advance the matched table coordinate", loc="left", fontweight="bold")
    figure.tight_layout()
    return save_figure(figure, figures / "04_runtime_sample_latch_replay_io_buf.png")


def plot_kukd_replay(figures: Path) -> Path:
    device = "io_buf"
    event = event_row(device, "short_high")
    data = load_csv_columns(
        STUDY / "waveforms" / f"{device}.csv",
        [
            "time_ns", "hspice_native_pad_v", "voltage_matching_pad_v",
            "hspice_native_ku", "hspice_native_kd", "voltage_matching_ku",
            "voltage_matching_kd", "voltage_matching_kupadmatch",
            "voltage_matching_kdpadmatch",
        ],
    )
    reverse = float(event["input_reverse_ns"])
    t = data["time_ns"] - reverse
    mask = (t >= -0.35) & (t <= 2.1)
    figure, axes = plt.subplots(3, 1, figsize=(11.0, 8.0), sharex=True)
    axes[0].plot(t[mask], data["hspice_native_ku"][mask], color=BLACK, linewidth=2.8, label="HSPICE native IBIS")
    axes[0].plot(t[mask], data["voltage_matching_ku"][mask], color=RED, linewidth=2.2, label="final Ku")
    axes[0].plot(t[mask], data["voltage_matching_kupadmatch"][mask], color=PURPLE, linewidth=1.7, label="opposite-table Ku")
    axes[0].legend(loc="upper right", frameon=False, ncol=3)
    style_axis(axes[0], ylabel="Ku")
    axes[1].plot(t[mask], data["hspice_native_kd"][mask], color=BLACK, linewidth=2.8, label="HSPICE native IBIS")
    axes[1].plot(t[mask], data["voltage_matching_kd"][mask], color=ORANGE, linewidth=2.2, label="final Kd")
    axes[1].plot(t[mask], data["voltage_matching_kdpadmatch"][mask], color=PURPLE, linewidth=1.7, label="opposite-table Kd")
    axes[1].axhline(0, color=GRAY, linewidth=0.7)
    axes[1].legend(loc="upper right", frameon=False, ncol=3)
    style_axis(axes[1], ylabel="Kd")
    axes[2].plot(t[mask], data["hspice_native_pad_v"][mask], color=BLACK, linewidth=2.8, label="HSPICE native IBIS")
    axes[2].plot(t[mask], data["voltage_matching_pad_v"][mask], color=RED, linewidth=2.2, label="Voltage matching V2")
    axes[2].legend(loc="upper right", frameon=False, ncol=2)
    style_axis(axes[2], xlabel="Time from second input edge (ns)", ylabel="Pad voltage (V)")
    for axis in axes:
        axis.axvline(0, color=CYAN, linestyle="--", linewidth=1.3)
    axes[0].set_title("One matched time drives the paired falling Ku and Kd tables", loc="left", fontweight="bold")
    figure.tight_layout()
    return save_figure(figure, figures / "05_runtime_kukd_replay_io_buf.png")


def plot_inv_chain_failure(figures: Path) -> Path:
    device = "inv_chain"
    event = event_row(device, "short_high")
    reference = json.loads(
        (STUDY / "calibration" / device / "fast_5ps" / "pad_replay_reference.json").read_text(encoding="utf-8")
    )
    falling = reference["falling"]
    table_t = np.asarray(falling["time_ns"], dtype=float)
    table_v = np.asarray(falling["pad_v"], dtype=float)
    sample = float(event["sampled_pad_v"])
    lookup = float(event["opposite_table_lookup_start_ns"])
    late = lookup + float(event["opposite_table_start_span_ns"])
    nearest = int(np.argmin(np.abs(table_v - sample)))

    data = load_csv_columns(
        STUDY / "waveforms" / f"{device}.csv",
        ["time_ns", "input_v", "hspice_native_pad_v", "hspice_transistor_pad_v", "voltage_matching_pad_v"],
    )
    reverse = float(event["input_reverse_ns"])
    t = data["time_ns"] - reverse
    mask = (t >= -0.18) & (t <= 0.70)

    figure, axes = plt.subplots(1, 2, figsize=(13.0, 5.4))
    axes[0].plot(t[mask], data["hspice_native_pad_v"][mask], color=BLACK, linewidth=2.8, label="HSPICE native IBIS")
    axes[0].plot(t[mask], data["hspice_transistor_pad_v"][mask], color=GRAY, linewidth=2.5, label="HSPICE transistor")
    axes[0].plot(t[mask], data["voltage_matching_pad_v"][mask], color=RED, linewidth=2.2, label="Voltage matching V2")
    axes[0].scatter([float(event["sample_capture_time_ns"]) - reverse], [sample], color=RED, marker="D", s=80, zorder=5)
    axes[0].axvline(0, color=CYAN, linestyle="--", linewidth=1.3)
    axes[0].set_title("Runtime: pad is still near zero", fontweight="bold")
    style_axis(axes[0], xlabel="Time from second edge (ns)", ylabel="Pad voltage (V)")
    axes[0].legend(loc="upper left", frameon=False)

    axes[1].plot(table_t, table_v, color=BLUE, linewidth=2.5, label="falling calibration pad")
    axes[1].axhline(sample, color=RED, linestyle="--", linewidth=1.4)
    axes[1].axvline(lookup, color=PURPLE, linestyle="--", linewidth=1.4, label=f"implemented lookup = {lookup:.3f} ns")
    axes[1].axvline(late, color=ORANGE, linestyle=":", linewidth=1.5, label=f"latest crossing = {late:.3f} ns")
    axes[1].scatter([table_t[nearest]], [table_v[nearest]], color=GREEN, s=75, zorder=5, label=f"nearest sampled point = {table_t[nearest]:.3f} ns")
    axes[1].set_xlim(0, 14)
    axes[1].set_title("Offline inverse is endpoint-sensitive and ambiguous", fontweight="bold")
    style_axis(axes[1], xlabel="Falling calibration-table time (ns)", ylabel="Calibration pad voltage (V)")
    axes[1].legend(loc="upper right", frameon=False, fontsize=9)
    figure.tight_layout()
    return save_figure(figure, figures / "06_inv_chain_propagation_delay_and_mapping_ambiguity.png")


def plot_three_buffer_summary(figures: Path) -> Path:
    devices = ["io_buf", "inv_chain", "ex2"]
    metrics = [metric_row(device, "short_high") for device in devices]
    events = [event_row(device, "short_high") for device in devices]
    pad = np.asarray([float(row["pad_rmse_mv"]) for row in metrics])
    ku = np.asarray([float(row["ku_rmse"]) for row in metrics])
    kd = np.asarray([float(row["kd_rmse"]) for row in metrics])
    sample = np.asarray([float(row["sampled_pad_v"]) for row in events])
    ambiguous = np.asarray([row["mapping_ambiguous"].lower() == "true" for row in events])
    x = np.arange(len(devices))

    figure, axes = plt.subplots(1, 3, figsize=(13.2, 4.8))
    colors = [GREEN if not value else RED for value in ambiguous]
    axes[0].bar(x, sample, color=colors)
    axes[0].set_xticks(x, devices)
    axes[0].set_title("Sampled pad voltage", fontweight="bold")
    style_axis(axes[0], ylabel="V")
    for index, value in enumerate(ambiguous):
        axes[0].text(index, sample[index], "ambiguous" if value else "unique", ha="center", va="bottom", fontsize=9)
    axes[1].bar(x, pad, color=BLUE)
    axes[1].set_xticks(x, devices)
    axes[1].set_title("Pad error", fontweight="bold")
    style_axis(axes[1], ylabel="RMSE (mV)")
    width = 0.34
    axes[2].bar(x - width / 2, ku, width, color=RED, label="Ku")
    axes[2].bar(x + width / 2, kd, width, color=ORANGE, label="Kd")
    axes[2].set_xticks(x, devices)
    axes[2].set_title("Coefficient error", fontweight="bold")
    style_axis(axes[2], ylabel="RMSE")
    axes[2].legend(frameon=False)
    figure.suptitle("Original Voltage Matching V2: real short-high results", fontsize=15, fontweight="bold")
    figure.tight_layout(rect=(0, 0, 1, 0.94))
    return save_figure(figure, figures / "07_three_buffer_voltage_matching_summary.png")


def delay_rows() -> list[dict[str, object]]:
    # Values are read from the same directional-fit comments in the generated
    # gate-state models. The delayed voltage matcher uses the minimum relevant
    # output-stage action for each reverse direction.
    source = {
        "io_buf": {"pu_on": 0.00165, "pu_off": 0.00165, "pd_off": 0.00371834, "pd_on": 0.00279171},
        "inv_chain": {"pu_on": 0.268425, "pu_off": 0.247282, "pd_off": 0.277529, "pd_on": 0.242304},
        "ex2": {"pu_on": 1.00146, "pu_off": 0.6683, "pd_off": 0.983366, "pd_on": 0.708729},
    }
    rows = []
    for device, values in source.items():
        rows.append({
            "device": device,
            **values,
            "short_high_sample_delay_ns": min(values["pu_off"], values["pd_on"]),
            "short_low_sample_delay_ns": min(values["pu_on"], values["pd_off"]),
        })
    return rows


def plot_delayed_sampling(figures: Path, rows: list[dict[str, object]]) -> Path:
    devices = [str(row["device"]) for row in rows]
    high = np.asarray([float(row["short_high_sample_delay_ns"]) for row in rows])
    low = np.asarray([float(row["short_low_sample_delay_ns"]) for row in rows])
    x = np.arange(len(devices))
    width = 0.34
    figure, axes = plt.subplots(1, 2, figsize=(12.6, 4.9), gridspec_kw={"width_ratios": [1.05, 1.45]})
    axes[0].bar(x - width / 2, 1000 * high, width, color=PURPLE, label="short-high: min(PU off, PD on)")
    axes[0].bar(x + width / 2, 1000 * low, width, color=GREEN, label="short-low: min(PU on, PD off)")
    axes[0].set_xticks(x, devices)
    axes[0].set_yscale("log")
    axes[0].legend(frameon=False, fontsize=9)
    axes[0].set_title("Real IBIS-derived sample delays", fontweight="bold")
    style_axis(axes[0], ylabel="Delay (ps, log scale)")

    axes[1].hlines(0.5, 0, 1.0, color=GRAY, linewidth=2)
    axes[1].scatter([0.05, 0.48, 0.75], [0.5, 0.5, 0.5], s=[90, 90, 90], color=[CYAN, PURPLE, RED], zorder=5)
    axes[1].text(0.05, 0.62, "second edge", ha="center", color=CYAN, fontsize=10)
    axes[1].text(0.48, 0.62, "wait IBIS delay", ha="center", color=PURPLE, fontsize=10)
    axes[1].text(0.75, 0.62, "sample Vpad\nand start mapping", ha="center", color=RED, fontsize=10)
    axes[1].annotate("", xy=(0.73, 0.5), xytext=(0.08, 0.5), arrowprops={"arrowstyle": "->", "color": PURPLE, "linewidth": 2})
    axes[1].set_xlim(0, 1)
    axes[1].set_ylim(0.2, 0.9)
    axes[1].axis("off")
    axes[1].set_title("Delayed extension changes only the sample time", fontweight="bold")
    figure.tight_layout()
    return save_figure(figure, figures / "08_delayed_voltage_sampling_design.png")


def add_small_label(deck: GreenDeck, slide, text: str, x: float, y: float, w: float) -> None:
    deck.add_text(slide, text.upper(), x, y, w, 0.24, size=9.5, color=deck.theme.gray, bold=True, align=PP_ALIGN.RIGHT)


def build_deck(out: Path, figure_paths: dict[str, Path], delays: list[dict[str, object]]) -> Path:
    renderer = EquationRenderer(backend="auto")
    deck = GreenDeck(equation_renderer=renderer)
    deck.set_title_slide(
        "Voltage-Matched Ku/Kd Replay\nfor Interrupted IBIS Transitions",
        "Simon Hwang\n8/14/2026",
        title_size=30,
        notes="""
This deck explains the Voltage Matching V2 experiment using stored waveforms from io_buf, inv_chain, and ex2. It separates the algorithm into its offline calibration stage and runtime replay stage.

The central idea is not to convert pad voltage directly into Ku and Kd. Pad voltage is first converted into a time coordinate on the opposite full-swing calibration trajectory. That one time coordinate is then used to evaluate the paired opposite Ku and Kd tables.

The deck also shows the limitation exposed by inv_chain: a pad that is still near zero because of propagation delay is not equivalent to a truly settled-low internal state. The final slide discusses the delayed-sampling extension, whose code and unit tests exist but whose corrected cross-buffer run was not complete when this deck was generated.
""",
    )

    slide = deck.add_slide("1. Why interrupted transitions are different", section="Problem")
    deck.add_picture_contain(slide, figure_paths["problem"], 0.6, 1.05, 7.55, 5.25)
    deck.add_box(slide, "Normal transition", 8.55, 1.25, 3.6, 0.62, fill=deck.theme.light_green, line=deck.theme.green, bold=True)
    deck.add_text(slide, "The first edge finishes. Starting the opposite full-swing table from its endpoint is reasonable.", 8.65, 2.02, 3.35, 1.05, size=16)
    deck.add_box(slide, "Interrupted transition", 8.55, 3.35, 3.6, 0.62, fill=deck.theme.light_red, line=deck.theme.red, bold=True)
    deck.add_text(slide, "The second edge arrives while internal drive and loaded pad response still contain history from the first edge.", 8.65, 4.12, 3.35, 1.2, size=16)
    deck.add_takeaway(slide, "At the second edge, restarting the opposite Ku/Kd tables from t = 0 discards unfinished switching history.")
    deck.add_source(slide, "Cached inv_chain short-high reference: three_buffer_single_true_output_reversal_2026-08-11/waveforms/inv_chain.csv")
    deck.add_notes(slide, """
Use the upper panel to identify the second input edge at time zero. The lower panel shows that the transistor and native-IBIS pad responses occur after this command edge. The output has propagation delay and memory; the instantaneous input command does not uniquely define the current output state.

Legacy pybis chooses a complete rising or falling coefficient table from the newest edge and resets that table's elapsed-time coordinate. Voltage matching was introduced to avoid assuming that every reverse edge begins from a fully settled endpoint.
""")

    slide = deck.add_slide("2. The voltage-matching idea in one chain", section="Concept")
    boxes = [
        ("Full-swing\ncalibration", 0.55, PALE_BLUE, deck.theme.blue),
        ("Store Vpad(t),\nKu(t), Kd(t)", 3.0, deck.theme.light_green, deck.theme.green),
        ("Sample Vpad at\nthe reverse event", 5.45, PALE_ORANGE, deck.theme.orange),
        ("Find time on the\nopposite Vpad(t)", 7.9, PALE_PURPLE, RGBColor(123, 44, 191)),
        ("Replay opposite\nKu/Kd from that time", 10.35, deck.theme.light_red, deck.theme.red),
    ]
    for text, x, fill, line in boxes:
        deck.add_box(slide, text, x, 2.05, 2.15, 1.05, fill=fill, line=line, size=16, bold=True)
    for index in range(len(boxes) - 1):
        deck.add_arrow(slide, boxes[index][1] + 2.15, 2.58, boxes[index + 1][1], 2.58, color=deck.theme.green)
    deck.add_equation(slide, r"V_s \rightarrow T_{opp}(V_s) \rightarrow \{K_{u,opp}(T),K_{d,opp}(T)\}", 2.25, 3.65, 8.9, 1.0, font_size_pt=30)
    deck.add_text(slide, "Important: voltage is an index into an opposite pad trajectory. It is not directly converted into two independent coefficient values.", 1.6, 5.0, 10.2, 0.85, size=20, bold=True, align=PP_ALIGN.CENTER)
    deck.add_takeaway(slide, "The shared matched time keeps opposite Ku and Kd synchronized as one table pair.")
    deck.add_notes(slide, """
There are two phases. The first two boxes happen offline. The remaining boxes happen during the interrupted transient simulation.

The calibration simulation records the loaded pad voltage and the corresponding legacy Ku and Kd for complete rising and falling transitions. At runtime, V2 samples the pad once at a reverse event. It searches the opposite-direction pad calibration for the same voltage. The resulting time coordinate is then applied to both opposite coefficient tables.

This is deliberately different from coefficient value matching, where Ku and Kd might each imply a different opposite-table time. Voltage matching chooses one coordinate from the loaded output and preserves pair alignment by construction.
""")

    slide = deck.add_slide("3. Offline step: run one complete rise and fall", section="Calibration")
    deck.add_picture_contain(slide, figure_paths["calibration"], 0.55, 1.0, 8.25, 5.55)
    deck.add_code_box(slide, "Vin ... 5.00n 0\n+ 5.05n VDD\n+ 25.00n VDD\n+ 25.05n 0\nRload pad 0 50\nCload pad 0 2p\n.save V(pad) V(xdrv.ku) V(xdrv.kd)", 9.0, 1.3, 3.45, 2.55, size=11)
    deck.add_bullets(slide, [
        "50 ps command edges",
        "50 ohm || 2 pF calibration load",
        "legacy pybis complete transitions",
        "rising and falling sections stored separately",
    ], 9.05, 4.18, 3.2, 1.55, size=15)
    deck.add_takeaway(slide, "The inverse voltage map is setup-specific because the calibration pad waveform depends on the load.")
    deck.add_source(slide, "Real io_buf calibration.raw and pad_replay_reference.json; source=legacy_pybis_ngspice.")
    deck.add_notes(slide, """
This slide shows the actual io_buf calibration run, not an analytical sketch. The left panel is the complete rising pad event and the right panel is the complete falling pad event.

The calibration deck holds each logic state for about twenty nanoseconds, so these are settled full-swing transitions. The recorded waveform is cut into fourteen-nanosecond rising and falling sections and downsampled into pad_replay_reference.json for embedding in the generated model.

The purpose of this extra calibration is to obtain Vpad versus time. It does not re-extract Ku and Kd; the original pybis/IBIS-derived coefficient tables already provide Ku(t) and Kd(t). Voltage matching connects the new Vpad(t) coordinate to those existing tables through matched time.

Because Vpad(t) is affected by source strength, output package, 50-ohm resistance, and 2-pF capacitance, this map is valid only for the declared calibration setup. It is not a general internal-state observer for arbitrary loads.
""")

    slide = deck.add_slide("4. Offline step: invert Vpad(t) into table time", section="Inverse map")
    deck.add_picture_contain(slide, figure_paths["inverse"], 0.65, 1.0, 7.7, 5.45)
    deck.add_equation(slide, r"T_F(V_s)=\operatorname*{arg\,match}_{t}\left[V_F(t),V_s\right]", 8.55, 1.45, 3.75, 0.9, font_size_pt=27)
    deck.add_text(slide, "Real io_buf short-high event", 8.75, 2.75, 3.25, 0.35, size=18, bold=True, color=deck.theme.green)
    deck.add_bullets(slide, [
        "sampled Vpad = 0.357 V",
        "falling lookup time = 0.484 ns",
        "start-span = 0 ns",
        "mapping marked unambiguous",
    ], 8.72, 3.2, 3.35, 1.75, size=16)
    deck.add_takeaway(slide, "For a short-high pulse, the falling pad trajectory determines the starting time of falling Ku and Kd.")
    deck.add_source(slide, "Actual values: voltage_matching_mapping_events.csv, io_buf/short_high.")
    deck.add_notes(slide, """
The horizontal red line is the pad voltage sampled during the interrupted event. The vertical purple line is the time on the stored falling calibration trajectory returned by the inverse lookup.

For this io_buf example the sampled voltage is 0.357 V and the inverse falling lookup returns 0.484 ns. The earliest and latest crossing policies agree, so the mapping span is zero and the event is not labeled ambiguous.

This lookup is generated offline as a PWL voltage-to-time source. At runtime, ngspice evaluates the PWL using PADSAMP. The resulting PADSTARTCMD is then latched so later changes in pad voltage cannot move the replay start.
""")

    slide = deck.add_slide("5. Runtime step: sample once, then latch", section="Runtime")
    deck.add_picture_contain(slide, figure_paths["runtime"], 0.55, 1.0, 8.2, 5.5)
    deck.add_code_box(slide, "BPADSAMPLE PADSAMP 0 I =\n -Csample*PMSAMPLE*\n (V(OUT)-V(PADSAMP))/tau_sample\n\nBPADSTART ...\n (PADSTARTCMD-PADSTART_LATCH)", 9.0, 1.25, 3.45, 2.3, size=10.7)
    deck.add_bullets(slide, [
        "PMSAMPLE opens a short sampling window",
        "PADSAMP stores one voltage",
        "PADSTART_LATCH stores one inverse-map time",
        "No continuous pad feedback after the latch",
    ], 9.0, 3.9, 3.3, 1.85, size=14.5)
    deck.add_takeaway(slide, "After the short latch transaction, the replay is open-loop again; the pad does not continuously control Ku/Kd.")
    deck.add_source(slide, "Cached io_buf runtime diagnostics from waveforms/io_buf.csv.")
    deck.add_notes(slide, """
The first panel shows the physical pad, the candidate pad, and PADSAMP. PADSAMP is a capacitor-backed sample-and-hold node. It changes only while PMSAMPLE is active.

The second panel shows the control sequence: sample the voltage, latch the inverse-map result, and activate the replay. The third panel shows the stored start time and PADARG. PADARG begins at the matched start and then increases with runtime elapsed after activation.

The model is not a continuously feedback-controlled pad model. Continuous feedback would create a nonlinear algebraic loop between output voltage, coefficient choice, current, and output voltage. V2 intentionally takes one snapshot, latches one coordinate, and returns to open-loop table playback.
""")

    slide = deck.add_slide("6. Runtime step: one time coordinate drives both coefficients", section="Runtime")
    deck.add_picture_contain(slide, figure_paths["kukd"], 0.55, 1.0, 8.1, 5.55)
    deck.add_equation(slide, r"A(t)=T_F(V_s)+(t-t_{activate})", 8.85, 1.3, 3.25, 0.7, font_size_pt=25)
    deck.add_equation(slide, r"K_u(t)=K_{u,F}(A(t))", 8.85, 2.35, 3.25, 0.65, font_size_pt=25)
    deck.add_equation(slide, r"K_d(t)=K_{d,F}(A(t))", 8.85, 3.25, 3.25, 0.65, font_size_pt=25)
    deck.add_text(slide, "Both tables receive exactly the same PADARG. This avoids Ku-derived and Kd-derived start-time disagreement.", 8.85, 4.35, 3.2, 1.05, size=16, bold=True)
    deck.add_takeaway(slide, "Voltage matching preserves Ku/Kd table alignment, but alignment alone does not guarantee physical state correctness.")
    deck.add_source(slide, "Real io_buf HSPICE-native and Voltage Matching V2 coefficient waveforms.")
    deck.add_notes(slide, """
PADARG is the independent variable used to evaluate the opposite Ku and Kd tables. It equals the latched matched start time plus elapsed time after replay activation.

For a short-high event the new command is falling, so KURM and KDRM are ignored and KUFM and KDFM are selected. Both falling coefficients use the same PADARG. For a short-low event the rising pair is selected instead.

The figure compares the final candidate coefficients against HSPICE native IBIS. The purple curves are the raw opposite-table replay targets. The red and orange curves are the final Ku and Kd after the brief hold/blend transaction. The bottom panel shows the output consequence.
""")

    slide = deck.add_slide("7. The generated ngspice implementation", section="Implementation")
    deck.add_text(slide, "Offline-generated inverse map", 0.65, 1.0, 3.7, 0.3, size=18, bold=True, color=deck.theme.blue)
    deck.add_code_box(slide, "B30 TR_PAD_EARLY 0 V=pwl(V(PADSAMP), ...)\nB31 TF_PAD_EARLY 0 V=pwl(V(PADSAMP), ...)\nB38 PADSTARTCMD 0 V=\n (NINX>0.5) ? TR_PAD_EARLY : TF_PAD_EARLY", 0.65, 1.42, 5.75, 1.65, size=10.6)
    deck.add_text(slide, "Runtime timer and table argument", 6.8, 1.0, 4.1, 0.3, size=18, bold=True, color=deck.theme.green)
    deck.add_code_box(slide, "B41 PMELAPSED 0 V=\n max(0,time_ns-PMT0-edge_delay)\nB42 PADARG 0 V=\n PADSTART_LATCH+PMELAPSED", 6.8, 1.42, 5.75, 1.65, size=10.6)
    deck.add_text(slide, "Opposite coefficient lookup", 0.65, 3.45, 3.7, 0.3, size=18, bold=True, color=deck.theme.orange)
    deck.add_code_box(slide, "B45 KUFM 0 V=pwl(PADARG, Ku_fall_table)\nB46 KDFM 0 V=pwl(PADARG, Kd_fall_table)\nB47 KUPADMATCH 0 V=(NINX>0.5)?KURM:KUFM\nB48 KDPADMATCH 0 V=(NINX>0.5)?KDRM:KDFM", 0.65, 3.87, 5.75, 1.75, size=10.1)
    deck.add_text(slide, "Continuous handoff", 6.8, 3.45, 3.7, 0.3, size=18, bold=True, color=deck.theme.red)
    deck.add_code_box(slide, "KUTARGET=(1-HPMALPHA)*KULEG\n        +HPMALPHA*KUPADMATCH\nKDTARGET=(1-HPMALPHA)*KDLEG\n        +HPMALPHA*KDPADMATCH\nKu=KUTARGET; Kd=KDTARGET", 6.8, 3.87, 5.75, 1.75, size=10.3)
    deck.add_takeaway(slide, "The generated .subckt contains the calibration map, sample/latch state, fresh replay timer, and final coefficient blend.")
    deck.add_source(slide, "Implementation: tools/pybis2spice/pybis2spice/subcircuit.py, create_ngspice_pad_matched_replay_input_control_netlist().")
    deck.add_notes(slide, """
Walk through the four blocks in reading order. The inverse PWL tables are generated offline and embedded directly into the subcircuit. At runtime PADSTARTCMD chooses the rising or falling inverse map according to the new input direction.

PMT0 stores the replay activation time. PMELAPSED is a fresh timer independent of the legacy HNX edge timer. PADARG is the only independent variable passed to both opposite coefficient tables.

The final block blends the legacy coefficients and matched coefficients with HPMALPHA. During the sample and latch pulses, KUSAMP and KDSAMP temporarily hold the pre-handoff values to reduce coefficient jumps. Outside an active voltage-match transaction, V2 is exactly the legacy path.
""")

    slide = deck.add_slide("8. Real limitation: one voltage does not identify hidden history", section="Limitation")
    deck.add_picture_contain(slide, figure_paths["failure"], 0.55, 1.0, 8.35, 5.5)
    deck.add_box(slide, "inv_chain short-high", 9.2, 1.2, 3.0, 0.58, fill=deck.theme.light_red, line=deck.theme.red, bold=True)
    deck.add_bullets(slide, [
        "second edge: 5.129 ns",
        "sampled pad: 0.000191 V",
        "implemented falling start: 6.475 ns",
        "start disagreement span: 7.222 ns",
        "mapping class: ambiguous",
    ], 9.15, 2.05, 3.1, 2.2, size=15.5)
    deck.add_text(slide, "Near-zero pad can mean either settled low or a rising transition that has not propagated to the pad yet.", 9.2, 4.65, 3.0, 1.0, size=16, bold=True, color=deck.theme.red)
    deck.add_takeaway(slide, "Pad voltage alone is not a unique state variable when propagation delay, slope, and ringing history matter.")
    deck.add_source(slide, "Real inv_chain event and calibration data; no new simulation used for this figure.")
    deck.add_notes(slide, """
The left panel shows the key physical ambiguity. At the second input edge, the loaded pad is still essentially zero because the first edge has not propagated through the buffer. Internally, however, a rising response is pending.

The right panel shows what the inverse falling map sees. A near-zero voltage also occurs near the settled end of the falling calibration. The implemented uniformly sampled inverse map returns 6.475 ns, while the nearest stored trajectory point with a similar voltage appears around 0.458 ns. The earliest-versus-latest start span is 7.222 ns, so the algorithm correctly flags this event as ambiguous.

This is not merely an ngspice numerical issue. It is an observability limitation: Vpad without slope or command history cannot distinguish different hidden states that share the same instantaneous voltage.
""")

    slide = deck.add_slide("9. Real three-buffer result: useful mechanism, inconsistent trust", section="Evidence")
    deck.add_picture_contain(slide, figure_paths["summary"], 0.65, 1.0, 8.25, 5.45)
    deck.add_text(slide, "What improved", 9.2, 1.25, 2.8, 0.3, size=18, bold=True, color=deck.theme.green)
    deck.add_bullets(slide, [
        "one shared Ku/Kd replay time",
        "explicit ambiguity metric",
        "no continuous output feedback",
        "ordinary path stays legacy when inactive",
    ], 9.15, 1.65, 3.05, 1.8, size=15)
    deck.add_text(slide, "What did not generalize", 9.2, 3.8, 3.0, 0.3, size=18, bold=True, color=deck.theme.red)
    deck.add_bullets(slide, [
        "inv_chain and ex2 sample near an endpoint",
        "two of three short-high maps are ambiguous",
        "pad and coefficient errors remain large",
    ], 9.15, 4.2, 3.05, 1.3, size=15)
    deck.add_takeaway(slide, "Voltage Matching V2 is a valuable baseline and diagnostic, but the current instantaneous-pad version is not a general replacement.")
    deck.add_source(slide, "Actual candidate_metrics.csv and voltage_matching_mapping_events.csv, short-high rows.")
    deck.add_notes(slide, """
The first chart shows the sampled voltage. Green labels indicate a unique inverse-map result and red labels indicate ambiguity. io_buf produces an unambiguous mapping, while inv_chain and ex2 are ambiguous near the low endpoint.

The pad and coefficient error charts are all measured against HSPICE native IBIS over the study's active window. The important result is not that one model wins every bar. The important result is that a clean voltage match and a shared coefficient coordinate do not by themselves prove the hidden switching state is correct.

This method remains useful because it gives us a clear baseline, produces diagnostics such as start-span and mapping ambiguity, and exposes exactly where voltage-only state observation breaks down.
""")

    slide = deck.add_slide("10. Delayed-sampling extension: wait for output-stage onset", section="Extension")
    deck.add_picture_contain(slide, figure_paths["delay"], 0.6, 1.0, 7.55, 5.35)
    deck.add_equation(slide, r"d_{high}=\min(d_{PU,off},d_{PD,on})", 8.45, 1.25, 3.65, 0.65, font_size_pt=23)
    deck.add_equation(slide, r"d_{low}=\min(d_{PU,on},d_{PD,off})", 8.45, 2.05, 3.65, 0.65, font_size_pt=23)
    deck.add_bullets(slide, [
        "same directional delays as gate-state fitting",
        "io_buf: 1.65 ps",
        "inv_chain short-high: 242.3 ps",
        "ex2 short-high: 668.3 ps",
    ], 8.55, 3.05, 3.4, 1.65, size=15)
    deck.add_takeaway(slide, "The extension changes when Vpad is sampled; the voltage-to-time-to-Ku/Kd mapping is otherwise unchanged.")
    deck.add_source(slide, "Delay values from generated gate-state model fit comments; saved in data/delayed_sampling_parameters.csv.")
    deck.add_notes(slide, """
The delayed extension addresses the specific case where the second input edge occurs before any output response is visible. For a short-high reversal, the sample delay is the earlier of pullup-off and pulldown-on onset. For short-low it is the earlier of pullup-on and pulldown-off onset.

These are not HSPICE-tuned values. They come from the same IBIS-derived coefficient-table onset extraction used by the directional gate-state model. The three buffers span nearly three orders of magnitude, which is why one fixed delay would not be appropriate.

The extension still has a conceptual limitation. The voltage sampled after the second edge may contain responses from both the first and second commands. Delaying the sample can make the pad observable, but it does not guarantee that instantaneous pad voltage uniquely identifies the hidden coefficient history.
""")

    slide = deck.add_slide("11. Delayed implementation status: bugs found before acceptance", section="Status")
    deck.add_box(slide, "Attempt 1: transport-line delay", 0.75, 1.25, 3.5, 0.72, fill=deck.theme.light_red, line=deck.theme.red, bold=True)
    deck.add_text(slide, "The 1.65 ps io_buf delay forced approximately 20 fs global timesteps and produced more than 500 MB for one partial run.", 0.85, 2.15, 3.25, 1.25, size=16)
    deck.add_arrow(slide, 4.35, 2.55, 4.95, 2.55, color=deck.theme.green)
    deck.add_box(slide, "Attempt 2: latched timestamp", 4.95, 1.25, 3.5, 0.72, fill=PALE_ORANGE, line=deck.theme.orange, bold=True)
    deck.add_text(slide, "The first timestamp node initialized at zero during operating-point solve, creating a false sample near simulation start.", 5.05, 2.15, 3.25, 1.25, size=16)
    deck.add_arrow(slide, 8.55, 2.55, 9.15, 2.55, color=deck.theme.green)
    deck.add_box(slide, "Current fix: explicit event arm", 9.15, 1.25, 3.35, 0.72, fill=deck.theme.light_green, line=deck.theme.green, bold=True)
    deck.add_text(slide, "A real reverse-edge pulse must arm the timer before a delayed sample can fire. Generation and 46 pybis tests pass.", 9.25, 2.15, 3.05, 1.25, size=16)
    deck.add_code_box(slide, "on reverse edge:\n  arm = 1\n  t0 = simulation_time\n\nsample = arm &&\n  (time-t0 >= d_sample) &&\n  (time-t0 < d_sample + 20 ps)", 3.65, 4.05, 6.0, 1.45, size=12)
    deck.add_text(slide, "Corrected three-buffer correlation run: not completed yet", 3.35, 5.82, 6.6, 0.38, size=18, bold=True, color=deck.theme.red, align=PP_ALIGN.CENTER)
    deck.add_takeaway(slide, "No delayed waveform is claimed as evidence until the armed-timer rerun completes and its diagnostic nodes confirm the intended sample time.")
    deck.add_notes(slide, """
This slide distinguishes numerical implementation problems from modeling results. The transport-line run was not stuck; it was advancing with an impractically small timestep imposed by the shortest transmission-line delay.

The timestamp replacement completed quickly, but diagnostic nodes showed HREVERSE_SAMPLE firing at 1.65 ps after simulation start instead of at the real reverse edge around 6.53 ns. The capacitor initial condition was not a safe event sentinel during the operating-point solve.

The current code adds a separate armed state. The sample timing expression is ignored until a real reverse-edge pulse sets this latch. Static compilation and all 46 pybis unit tests pass, with one unrelated test skipped. The cross-buffer transient campaign was interrupted before the corrected run completed, so the deck intentionally shows no delayed correlation waveform.
""")

    slide = deck.add_slide("12. What we know and what must be tested next", section="Conclusion")
    deck.add_box(slide, "Established", 0.75, 1.15, 3.55, 0.65, fill=deck.theme.light_green, line=deck.theme.green, size=19, bold=True)
    deck.add_bullets(slide, [
        "full-swing calibration and inverse map are reproducible",
        "one shared replay time keeps Ku/Kd paired",
        "sample/latch implementation avoids continuous pad feedback",
        "ambiguity is measurable rather than hidden",
    ], 0.85, 2.0, 3.35, 2.4, size=15.5)
    deck.add_box(slide, "Not established", 4.88, 1.15, 3.55, 0.65, fill=deck.theme.light_red, line=deck.theme.red, size=19, bold=True)
    deck.add_bullets(slide, [
        "instantaneous Vpad is not a unique hidden-state coordinate",
        "delayed Vpad may mix first- and second-edge responses",
        "current map is load-specific",
        "delayed cross-buffer correlation is unfinished",
    ], 4.98, 2.0, 3.35, 2.4, size=15.5)
    deck.add_box(slide, "Next validation", 9.0, 1.15, 3.55, 0.65, fill=PALE_BLUE, line=deck.theme.blue, size=19, bold=True)
    deck.add_bullets(slide, [
        "complete armed-timer rerun from cached HSPICE references",
        "verify actual sample time and sampled voltage",
        "compare original versus delayed Ku/Kd and pad",
        "reject endpoint/ambiguous mappings explicitly",
    ], 9.1, 2.0, 3.25, 2.4, size=15.5)
    deck.add_equation(slide, r"\text{success}=\text{pad agreement}+\text{Ku agreement}+\text{Kd agreement}", 2.25, 5.2, 8.8, 0.72, font_size_pt=27, color="#2B7A43")
    deck.add_takeaway(slide, "Voltage matching is a strong diagnostic baseline; production trust requires coefficient-correct results across buffers and directions.")
    deck.add_notes(slide, """
Close by separating what is technically complete from what remains experimental. The calibration, inverse mapping, latching, and paired replay are all implemented and reproducible. Original V2 results are available for all three buffers and both reversal directions.

The negative result is equally important: instantaneous loaded pad voltage does not uniquely reveal pending internal drive history. The delayed extension is a reasonable experiment for propagation-delay-dominated cases, but it must be judged by Ku, Kd, and pad together.

The next action is narrow: complete the armed-timestamp delayed runs using existing HSPICE references, inspect the diagnostic sample times, and compare original versus delayed mapping. Do not promote the method based only on a smaller pad error.
""")

    output = out / "Voltage_Matching_V2_Explainer_2026-08-14.pptx"
    return deck.save(output)


def write_readme(out: Path, deck_path: Path, figures: dict[str, Path]) -> None:
    lines = [
        "# Voltage Matching V2 Explainer Deck",
        "",
        "This package explains the cached-data Voltage Matching V2 experiment and the delayed-sampling extension.",
        "",
        "## Main Deliverable",
        "",
        f"- `{deck_path.name}`",
        "- `Voltage_Matching_V2_Explainer_2026-08-14.pdf`: PowerPoint-rendered review copy.",
        "- `Voltage_Matching_V2_Explainer_contact_sheet.png`: all slides at a glance.",
        "- `rendered_slides/`: one 1600 x 900 PNG per slide.",
        "",
        "## Evidence Policy",
        "",
        "- Original V2 figures use cached full-swing calibration and interrupted-event waveforms.",
        "- HSPICE was not rerun for this deck.",
        "- Delayed-sampling slides use real IBIS-derived delay parameters, but do not claim an unfinished delayed waveform result.",
        "",
        "## Figures",
        "",
    ]
    lines.extend(f"- `{path.relative_to(out).as_posix()}`" for path in figures.values())
    lines.extend([
        "",
        "## Rebuild",
        "",
        "```powershell",
        "$env:PYTHONPATH = \".codex_deps/presentation/python;.\"",
        "py -3.14 scripts/build_voltage_matching_explainer_deck.py",
        "```",
        "",
    ])
    (out / "README.md").write_text("\n".join(lines), encoding="ascii")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    out = args.output_dir.resolve()
    figures = out / "figures"
    ensure(figures)

    figure_paths = {
        "problem": plot_problem(figures),
        "calibration": plot_full_swing_calibration(figures),
        "inverse": plot_inverse_mapping(figures),
        "runtime": plot_runtime_sample(figures),
        "kukd": plot_kukd_replay(figures),
        "failure": plot_inv_chain_failure(figures),
        "summary": plot_three_buffer_summary(figures),
    }
    delays = delay_rows()
    write_rows(out / "data" / "delayed_sampling_parameters.csv", delays)
    figure_paths["delay"] = plot_delayed_sampling(figures, delays)
    write_rows(out / "data" / "mapping_events_used.csv", [
        event_row("io_buf", "short_high"),
        event_row("inv_chain", "short_high"),
        event_row("ex2", "short_high"),
    ])
    write_rows(out / "data" / "metrics_used.csv", [
        metric_row("io_buf", "short_high"),
        metric_row("inv_chain", "short_high"),
        metric_row("ex2", "short_high"),
    ])
    deck_path = build_deck(out, figure_paths, delays)
    write_readme(out, deck_path, figure_paths)
    print(deck_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
