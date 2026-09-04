from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .equations import EquationAsset, EquationRenderer
from .theme import DEFAULT_TEMPLATE, LAB_GREEN_THEME, PresentationTheme


class GreenDeck:
    """Small, project-neutral API for scripted lab presentations."""

    def __init__(
        self,
        template: Path | str = DEFAULT_TEMPLATE,
        *,
        theme: PresentationTheme = LAB_GREEN_THEME,
        remove_example_slides: bool = True,
        equation_renderer: EquationRenderer | None = None,
    ) -> None:
        self.template = Path(template)
        if not self.template.exists():
            raise FileNotFoundError(f"Presentation template not found: {self.template}")
        self.theme = theme
        self.prs = Presentation(str(self.template))
        self.equations = equation_renderer or EquationRenderer()
        if remove_example_slides:
            self.remove_slides_after_title()

    def remove_slides_after_title(self) -> None:
        """Retain the template title slide and remove its example content."""
        slide_ids = self.prs.slides._sldIdLst
        for slide_id in list(slide_ids)[1:]:
            self.prs.part.drop_rel(slide_id.rId)
            slide_ids.remove(slide_id)

    def set_title_slide(
        self,
        title: str,
        subtitle: str,
        *,
        title_size: float = 35,
        subtitle_size: float = 16,
        notes: str | None = None,
    ):
        slide = self.prs.slides[0]
        title_shape = slide.shapes.title
        title_shape.text = title
        p = title_shape.text_frame.paragraphs[0]
        p.font.name = self.theme.title_font
        p.font.size = Pt(title_size)
        p.font.bold = True
        p.font.color.rgb = self.theme.black
        if len(slide.placeholders) > 1:
            subtitle_shape = slide.placeholders[1]
            subtitle_shape.text = subtitle
            for paragraph in subtitle_shape.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.font.name = self.theme.title_font
                paragraph.font.size = Pt(subtitle_size)
                paragraph.font.color.rgb = self.theme.black
        if notes:
            self.add_notes(slide, notes)
        return slide

    def add_slide(self, title: str, *, section: str | None = None, notes: str | None = None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[self.theme.content_layout_index])
        title_shape = slide.shapes.title
        title_shape.text = title
        p = title_shape.text_frame.paragraphs[0]
        p.font.name = self.theme.title_font
        p.font.size = Pt(self.theme.title_size_pt)
        p.font.bold = False
        p.font.color.rgb = self.theme.black
        p.alignment = PP_ALIGN.LEFT
        if section:
            self.add_text(
                slide,
                section.upper(),
                10.25,
                0.12,
                2.35,
                0.25,
                size=9.5,
                color=self.theme.gray,
                bold=True,
                align=PP_ALIGN.RIGHT,
            )
        if notes:
            self.add_notes(slide, notes)
        return slide

    def add_text(
        self,
        slide,
        text: str,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        size: float | None = None,
        color: RGBColor | None = None,
        bold: bool = False,
        font: str | None = None,
        align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.TOP,
        margin: float = 0.05,
    ):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(margin)
        tf.margin_right = Inches(margin)
        tf.margin_top = Inches(margin)
        tf.margin_bottom = Inches(margin)
        tf.vertical_anchor = valign
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = align
        p.font.name = font or self.theme.body_font
        p.font.size = Pt(size or self.theme.body_size_pt)
        p.font.bold = bold
        p.font.color.rgb = color or self.theme.dark
        return box

    def add_bullets(
        self,
        slide,
        items: list[str],
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        size: float | None = None,
        color: RGBColor | None = None,
        spacing: float = 7,
    ):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.08)
        tf.margin_right = Inches(0.04)
        tf.margin_top = Inches(0.04)
        tf.margin_bottom = Inches(0.04)
        for idx, item in enumerate(items):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = f"•  {item}"
            p.level = 0
            p.font.name = self.theme.body_font
            p.font.size = Pt(size or self.theme.body_size_pt)
            p.font.color.rgb = color or self.theme.dark
            p.space_after = Pt(spacing)
        return box

    def add_box(
        self,
        slide,
        text: str,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        fill: RGBColor | None = None,
        line: RGBColor | None = None,
        size: float = 16,
        bold: bool = False,
        color: RGBColor | None = None,
        rounded: bool = True,
    ):
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill or self.theme.light_green
        shape.line.color.rgb = line or self.theme.green
        shape.line.width = Pt(1.25)
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.08)
        tf.margin_right = Inches(0.08)
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = self.theme.body_font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color or self.theme.dark
        return shape

    def add_arrow(
        self,
        slide,
        x1: float,
        y1: float,
        x2: float,
        y2: float,
        *,
        color: RGBColor | None = None,
        width: float = 2.2,
    ):
        conn = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(x1),
            Inches(y1),
            Inches(x2),
            Inches(y2),
        )
        conn.line.color.rgb = color or self.theme.green
        conn.line.width = Pt(width)
        conn.line.end_arrowhead = True
        return conn

    def add_takeaway(self, slide, text: str, *, y: float = 6.63):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.42),
            Inches(y),
            Inches(12.15),
            Inches(0.43),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.theme.pale_green
        shape.line.color.rgb = self.theme.pale_green
        tf = shape.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = self.theme.body_font
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = self.theme.green
        return shape

    def add_source(self, slide, text: str):
        return self.add_text(slide, text, 0.52, 7.16, 11.8, 0.18, size=7.5, color=self.theme.gray)

    def add_code_box(self, slide, code: str, x: float, y: float, w: float, h: float, *, size: float = 11.5):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(246, 248, 247)
        shape.line.color.rgb = RGBColor(164, 174, 168)
        shape.line.width = Pt(1)
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = False
        tf.margin_left = Inches(0.13)
        tf.margin_right = Inches(0.08)
        tf.margin_top = Inches(0.1)
        tf.margin_bottom = Inches(0.08)
        # An autoshape defaults to centred, middle-anchored text. For a code box
        # that silently destroys column alignment -- a monospaced table comes out
        # with every row centred on its own width, so nothing lines up. One
        # paragraph per line, left aligned, anchored top.
        tf.vertical_anchor = MSO_ANCHOR.TOP
        lines = code.splitlines() or [""]
        for idx, line in enumerate(lines):
            para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            para.text = line
            para.alignment = PP_ALIGN.LEFT
            para.font.name = self.theme.code_font
            para.font.size = Pt(size)
            para.font.color.rgb = RGBColor(30, 48, 39)
        return shape

    def add_picture_contain(
        self,
        slide,
        path: Path | str,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        border: bool = True,
    ):
        path = Path(path)
        with Image.open(path) as image:
            aspect = image.width / image.height
        box_aspect = w / h
        if aspect >= box_aspect:
            pic_w = w
            pic_h = w / aspect
            pic_x = x
            pic_y = y + (h - pic_h) / 2
        else:
            pic_h = h
            pic_w = h * aspect
            pic_x = x + (w - pic_w) / 2
            pic_y = y
        if border:
            frame = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(pic_x - 0.03),
                Inches(pic_y - 0.03),
                Inches(pic_w + 0.06),
                Inches(pic_h + 0.06),
            )
            frame.fill.solid()
            frame.fill.fore_color.rgb = self.theme.white
            frame.line.color.rgb = self.theme.mid_gray
            frame.line.width = Pt(0.8)
        return slide.shapes.add_picture(
            str(path),
            Inches(pic_x),
            Inches(pic_y),
            width=Inches(pic_w),
            height=Inches(pic_h),
        )

    def add_equation(
        self,
        slide,
        tex: str,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        font_size_pt: float = 30,
        color: str = "#000000",
        display: bool = True,
    ) -> tuple[object, EquationAsset]:
        asset = self.equations.render(
            tex,
            font_size_pt=font_size_pt,
            color=color,
            display=display,
        )
        picture = self.add_picture_contain(slide, asset.path, x, y, w, h, border=False)
        # Add useful alternative text without relying on editable PowerPoint math.
        picture._element.nvPicPr.cNvPr.set("descr", f"Equation: {asset.tex}")
        return picture, asset

    @staticmethod
    def add_notes(slide, text: str) -> None:
        frame = slide.notes_slide.notes_text_frame
        frame.clear()
        frame.paragraphs[0].text = text.strip()

    def save(self, output: Path | str) -> Path:
        output = Path(output)
        output.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(output))
        return output
